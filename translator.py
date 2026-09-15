"""Document extraction, chunking, translation and output writers."""

from __future__ import annotations

import asyncio
import html
import json
import re
import time
import uuid
from pathlib import Path
from typing import Awaitable, Callable, List, Optional

import aiohttp
import docx
import ebooklib
from bs4 import BeautifulSoup
from ebooklib import epub

from config import CHUNK_SIZE, CONCURRENCY_LIMIT, LANGUAGES, MAX_RETRIES, REQUEST_TIMEOUT

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None

try:
    from deep_translator import GoogleTranslator
except ImportError:  # pragma: no cover
    GoogleTranslator = None


class JobCancelled(Exception):
    pass


def extract_text_epub(path: str) -> str:
    try:
        try:
            book = epub.read_epub(path, options={"ignore_ncx": True})
        except TypeError:
            book = epub.read_epub(path)
    except Exception:
        return ""
    sections, seen = [], set()
    items = [i for i in book.get_items() if i.get_type() == ebooklib.ITEM_DOCUMENT]
    for item in items:
        try:
            soup = BeautifulSoup(item.get_body_content(), "lxml")
        except Exception:
            soup = BeautifulSoup(item.get_body_content(), "html.parser")
        for junk in soup(["script", "style", "nav", "svg"]):
            junk.decompose()
        text = soup.get_text("\n", strip=True)
        key = hash(text[:400])
        if len(text) > 30 and key not in seen:
            seen.add(key)
            sections.append(text)
    return "\n\n".join(sections)


def extract_text_txt(path: str) -> str:
    raw = Path(path).read_bytes()
    for encoding in (("utf-16",) if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "utf-8", "cp1252", "latin-1")):
        try:
            return raw.decode(encoding).replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text_docx(path: str) -> str:
    try:
        document = docx.Document(path)
        parts = [paragraph.text.rstrip() for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_text_pdf(path: str) -> str:
    if PdfReader is None:
        return ""
    try:
        reader = PdfReader(path)
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages if (page.extract_text() or "").strip())
    except Exception:
        return ""


EXTRACTORS = {".epub": extract_text_epub, ".txt": extract_text_txt, ".docx": extract_text_docx, ".pdf": extract_text_pdf}


def normalise_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\u00ad", "")
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t\u00a0]+\n", "\n", text)).strip()


def split_text_by_size(text: str, max_kb: int, factor: float = 1.0) -> List[str]:
    if max_kb <= 0:
        return [text]
    max_bytes = max(1024, int(max_kb * 1024 / max(factor, 0.1)))
    parts, current, size = [], [], 0
    for line in text.split("\n"):
        line_size = len(line.encode("utf-8")) + 1
        if current and size + line_size > max_bytes:
            parts.append("\n".join(current)); current, size = [], 0
        current.append(line); size += line_size
    return parts + (["\n".join(current)] if current else []) or [text]


_SENTENCE_RE = re.compile(r"(?<=[.!?।。！？])\s+")


def build_chunks(text: str, limit: int = CHUNK_SIZE) -> List[List[str]]:
    chunks, buffer, size = [], [], 0
    for line in text.split("\n"):
        line = line.rstrip()
        pieces = _SENTENCE_RE.split(line) if len(line) > limit else [line]
        for piece in pieces:
            if len(piece) > limit:
                pieces = [piece[i:i + limit] for i in range(0, len(piece), limit)]
            if size + len(piece) + 1 > limit and buffer:
                chunks.append(buffer); buffer, size = [], 0
            buffer.append(piece); size += len(piece) + 1
    if buffer:
        chunks.append(buffer)
    return chunks


_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _xml_safe(value: str) -> str:
    return _XML_ILLEGAL.sub("", value)


def write_txt(text: str, path: str, title: str, lang: str) -> None:
    Path(path).write_text(text, encoding="utf-8")


def write_docx(text: str, path: str, title: str, lang: str) -> None:
    document = docx.Document(); document.add_heading(_xml_safe(title), level=1)
    for line in text.split("\n"):
        if line.strip(): document.add_paragraph(_xml_safe(line.strip()))
    document.save(path)


def write_epub(text: str, path: str, title: str, lang: str) -> None:
    book = epub.EpubBook(); book.set_identifier(uuid.uuid4().hex); book.set_title(title); book.set_language(lang.split("-")[0])
    chapters = []
    lines = text.split("\n")
    for index in range(0, len(lines), 400):
        number = index // 400 + 1
        body = "".join(f"<p>{html.escape(_xml_safe(line.strip()))}</p>" for line in lines[index:index + 400] if line.strip())
        chapter = epub.EpubHtml(title=f"Section {number}", file_name=f"section_{number:03d}.xhtml", lang=lang.split("-")[0])
        chapter.content = f"<html><body><h2>{html.escape(title)}</h2>{body}</body></html>"
        book.add_item(chapter); chapters.append(chapter)
    book.toc = tuple(chapters); book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav()); book.spine = ["nav"] + chapters
    epub.write_epub(path, book)


WRITERS = {"txt": write_txt, "docx": write_docx, "epub": write_epub}


class TranslationEngine:
    def __init__(self, target: str, concurrency: int = CONCURRENCY_LIMIT):
        self.target, self.sem = target, asyncio.Semaphore(concurrency)
        self.total = self.done = self.failed = 0; self.started = time.time()
        self._session: Optional[aiohttp.ClientSession] = None

    async def _session_for_translate(self) -> aiohttp.ClientSession:
        if self._session is None or self._session.closed:
            self._session = aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=REQUEST_TIMEOUT))
        return self._session

    async def _one(self, index: int, lines: List[str], cancel: asyncio.Event) -> tuple[int, List[str]]:
        if not any(line.strip() for line in lines): self.done += 1; return index, lines
        async with self.sem:
            for attempt in range(MAX_RETRIES):
                if cancel.is_set(): raise JobCancelled()
                try:
                    session = await self._session_for_translate()
                    payload = [("q", line) for line in lines if line.strip()]
                    async with session.post("https://translate.googleapis.com/translate_a/t", params={"client": "gtx", "sl": "auto", "tl": self.target}, data=payload) as response:
                        if response.status != 200: raise RuntimeError(f"HTTP {response.status}")
                        result = json.loads(await response.text())
                    translated = list(lines); pos = 0
                    for line_no, line in enumerate(lines):
                        if line.strip(): translated[line_no] = html.unescape(result[pos][0] if isinstance(result[pos], list) else result[pos]); pos += 1
                    self.done += 1; return index, translated
                except Exception:
                    await asyncio.sleep(min(2 ** attempt, 8))
            if GoogleTranslator is not None:
                try:
                    translated = await asyncio.to_thread(GoogleTranslator(source="auto", target=self.target).translate, "\n".join(lines))
                    self.done += 1; return index, translated.split("\n")
                except Exception:
                    pass
            self.failed += 1; self.done += 1; return index, lines

    async def run(self, chunks: List[List[str]], cancel: asyncio.Event, on_progress: Optional[Callable] = None) -> str:
        self.total = len(chunks)
        results = await asyncio.gather(*(self._one(i, chunk, cancel) for i, chunk in enumerate(chunks)))
        if on_progress:
            await on_progress(self)
        if self._session and not self._session.closed: await self._session.close()
        return "\n".join("\n".join(lines) for _, lines in sorted(results))

    @property
    def speed(self) -> float:
        return self.done / max(time.time() - self.started, 0.001)

    @property
    def eta(self) -> float:
        return (self.total - self.done) / self.speed if self.done and self.speed else 0