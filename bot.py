#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
╔══════════════════════════════════════════════════════════════════════════╗
║        📚 NovelTranslator PRO  —  Telegram Document Translation Bot       ║
║      v6.4  ·  Master + embedded/extra Workers · GridFS · Mini App        ║
╠══════════════════════════════════════════════════════════════════════════╣
║  • Approval-based access (no password) → users request, owner/admins    ║
║    approve for 1 week / month / year / lifetime / custom; auto-expiry,  ║
║    reminders, re-request, reject, ban, audit log, admin roles           ║
║  • MongoDB persistence (motor)        → users / settings / stats / jobs  ║
║    survive restarts. Falls back to RAM when MONGO_URI is not set.        ║
║  • Telegram Mini App (/app)           → premium dashboard: settings,     ║
║    live job progress, history, owner panel — served by the same process  ║
║  • Hierarchical reply keyboard        → tap a menu → sub-menu appears    ║
║  • Bot session kept in memory         → no .session files                ║
║  • Health server on $PORT + keep-alive ping + graceful SIGTERM handling  ║
╚══════════════════════════════════════════════════════════════════════════╝
"""

import asyncio
import calendar
import hashlib
import hmac
import html
import json
import logging
import os
import random
import re
import shutil
import sys
import time
import uuid
from collections import deque
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any, Deque, Dict, List, Optional, Tuple
from urllib.parse import parse_qsl

import aiohttp
import docx
import ebooklib
from aiohttp import web
from bs4 import BeautifulSoup
from ebooklib import epub
from database import job_repo_from_store

# Pyrogram still calls asyncio.get_event_loop() at import/Client-init time.
# On Python ≥ 3.12 that is deprecated (and fails on 3.14) when no loop exists,
# so we create one explicitly *before* importing/instantiating the client.
try:
    asyncio.get_running_loop()
except RuntimeError:
    asyncio.set_event_loop(asyncio.new_event_loop())

from pyrogram import Client, enums, filters, idle  # noqa: E402
from pyrogram.errors import (FloodWait, MessageIdInvalid, MessageNotModified,  # noqa: E402
                             RPCError)
from pyrogram.raw import functions  # noqa: E402
from pyrogram.types import (BotCommand, BotCommandScopeAllPrivateChats,  # noqa: E402
                            BotCommandScopeChat, BotCommandScopeDefault, CallbackQuery,
                            ChatMemberUpdated, InlineKeyboardButton, InlineKeyboardMarkup,
                            KeyboardButton, MenuButtonWebApp, Message, ReplyKeyboardMarkup,
                            ReplyKeyboardRemove, WebAppInfo)

try:  # optional: MongoDB persistence
    from motor.motor_asyncio import AsyncIOMotorClient
    HAS_MOTOR = True
except ImportError:  # pragma: no cover
    HAS_MOTOR = False

try:
    from pypdf import PdfReader
    HAS_PDF = True
except ImportError:  # pragma: no cover
    HAS_PDF = False

try:  # optional: .env for local development only
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:  # pragma: no cover
    pass

try:  # optional fallback translator (thread based)
    from deep_translator import GoogleTranslator as _FallbackTranslator
    HAS_FALLBACK = True
except ImportError:  # pragma: no cover
    HAS_FALLBACK = False


# ═══════════════════════════════════════════════════════════════════════════
# 🔑 CONFIGURATION — single source of truth lives in config.py
#     (environment variables / .env only; nothing is parsed here any more)
# ═══════════════════════════════════════════════════════════════════════════
from config import (  # noqa: E402
    ADMIN_USERS, AUDIT_LIMIT, AUTHORIZED_USERS, API_HASH, API_ID, BACKUP_GROUP_ID, BASE_DIR,
    BOT_NAME, BOT_TOKEN, CHUNK_SIZE, CONCURRENCY_LIMIT, DB_BUDGET_MB, DB_JOB_TTL_DAYS,
    DB_MAX_JOB_DOCS, DEFAULT_APPROVAL, DEFAULT_FORMAT, DEFAULT_LANG, DEFAULT_SPLIT_KB, EDIT_INTERVAL,
    EMBEDDED_WORKER, EXPANSION, EXPIRY_REMINDER_DAYS, HISTORY_LIMIT, INBOX_DIR, INIT_DATA_MAX_AGE,
    KEEP_ALIVE, KEEP_ALIVE_INTERVAL, LANGUAGES, LEGACY_USERS, MAX_INPUT_MB, MAX_JOBS_PER_USER,
    MAX_RETRIES, MAX_SPLIT_KB, MINIAPP_DEV_USER, MINI_APP_DIR, MINI_APP_URL, MIN_SPLIT_KB, MONGO_DB,
    MONGO_URI, OUTPUT_FORMATS, OWNER_ID, PENDING_TTL, PORT, PUBLIC_MODE,
    REJECT_COOLDOWN_H, RENDER_EXTERNAL_URL, REQUEST_TIMEOUT, SPLIT_PRESETS,
    STORAGE_DIR, VERSION, env as _env,
    validate_config as _validate_config,
)

# PDF input only when pypdf is importable (config lists it unconditionally)
INPUT_EXTS = {".epub", ".txt", ".docx"} | ({".pdf"} if HAS_PDF else set())

# ═══════════════════════════════════════════════════════════════════════════
# 🔧 LOGGING  (stdout only — Render captures it; no log files on disk)
# ═══════════════════════════════════════════════════════════════════════════
log = logging.getLogger("bot")

def validate_config() -> None:
    try:
        _validate_config(require_telegram=True, role="master")
    except RuntimeError as e:
        log.critical("%s", e)
        log.critical("Set the missing values in the Render dashboard → Environment (or .env), then redeploy.")
        sys.exit(1)
    if _env("SECURITY_CODE"):
        log.warning("SECURITY_CODE is no longer used — access is approval based now (v5). You can remove the variable.")

validate_config()


# ═══════════════════════════════════════════════════════════════════════════
# 🎫 ACCESS MODEL  —  approval states, plans & duration parsing
# ═══════════════════════════════════════════════════════════════════════════
ACCESS_NONE, ACCESS_PENDING, ACCESS_APPROVED = "none", "pending", "approved"
ACCESS_EXPIRED, ACCESS_REJECTED, ACCESS_BANNED = "expired", "rejected", "banned"
ACCESS_LABELS = {
    ACCESS_NONE: ("No access", "🔒"), ACCESS_PENDING: ("Pending approval", "⏳"),
    ACCESS_APPROVED: ("Approved", "✅"), ACCESS_EXPIRED: ("Expired", "⌛"),
    ACCESS_REJECTED: ("Rejected", "❌"), ACCESS_BANNED: ("Banned", "🚫"),
}
# Quick-approve presets shown to admins (code → label, duration spec)
PLAN_PRESETS: Dict[str, Tuple[str, str]] = {
    "1w": ("1 Week", "1w"), "1m": ("1 Month", "1m"), "3m": ("3 Months", "3m"),
    "6m": ("6 Months", "6m"), "1y": ("1 Year", "1y"), "forever": ("Lifetime", "forever"),
}
FOREVER_WORDS = {"forever", "lifetime", "permanent", "unlimited", "never", "always", "∞", "inf", "0"}

def _add_months(dt: datetime, months: float) -> datetime:
    whole = int(months)
    frac = months - whole
    month0 = dt.month - 1 + whole
    year = dt.year + month0 // 12
    month = month0 % 12 + 1
    day = min(dt.day, calendar.monthrange(year, month)[1])
    return dt.replace(year=year, month=month, day=day) + timedelta(days=frac * 30)

def parse_duration(text: str, start: Optional[float] = None) -> Optional[int]:
    """Turn a human duration into an absolute expiry timestamp.
    Returns 0 for lifetime, None when the text is not understood.

      forever · lifetime · 30 · 30d · 2w · 1m (month) · 6mo · 1y · 12h · 2026-12-31
    """
    t = (text or "").strip().lower().replace(" ", "")
    if not t:
        return None
    if t in FOREVER_WORDS:
        return 0
    start = start or time.time()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        try:
            dt = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)), 23, 59, 59)
        except ValueError:
            return None
        ts = int(dt.timestamp())
        return ts if ts > start else None
    m = re.fullmatch(r"(\d+(?:\.\d+)?)(h|hr|hrs|hour|hours|d|day|days|w|wk|wks|week|weeks|"
                     r"m|mo|mon|month|months|y|yr|yrs|year|years)?", t)
    if not m:
        return None
    n = float(m.group(1))
    unit = (m.group(2) or "d")[0]
    if n <= 0 or n > 100000:
        return None
    if unit == "h":
        return int(start + n * 3600)
    if unit == "d":
        return int(start + n * 86400)
    if unit == "w":
        return int(start + n * 7 * 86400)
    base = datetime.fromtimestamp(start)
    if unit == "m":
        return int(_add_months(base, n).timestamp())
    if unit == "y":
        return int(_add_months(base, n * 12).timestamp())
    return None

def plan_label(spec: str) -> str:
    """Human label for a duration spec / preset code."""
    if spec in PLAN_PRESETS:
        return PLAN_PRESETS[spec][0]
    if spec.lower() in FOREVER_WORDS:
        return "Lifetime"
    return spec or "—"

def fmt_date(ts: float) -> str:
    return time.strftime("%d %b %Y", time.localtime(ts)) if ts else "—"

def fmt_datetime(ts: float) -> str:
    return time.strftime("%d %b %Y · %H:%M", time.localtime(ts)) if ts else "—"

def days_left(expires: int) -> Optional[int]:
    """None → lifetime, otherwise whole days remaining (can be 0 = today)."""
    if not expires:
        return None
    return max(0, int((expires - time.time()) // 86400))

def expiry_label(expires: int) -> str:
    if not expires:
        return "♾ Lifetime"
    left = expires - time.time()
    if left <= 0:
        return f"expired {fmt_date(expires)}"
    d = int(left // 86400)
    if d >= 1:
        return f"until {fmt_date(expires)} ({d} day{'s' if d != 1 else ''} left)"
    h = max(1, int(left // 3600))
    return f"until {fmt_datetime(expires)} ({h} h left)"


# ═══════════════════════════════════════════════════════════════════════════
# 🧠 STORE  —  write-through in-memory cache backed by MongoDB (optional)
# ═══════════════════════════════════════════════════════════════════════════
class Store:
    """All state is kept in RAM for fast synchronous access; every mutation is
    additionally persisted to MongoDB (when MONGO_URI is configured) through a
    background writer queue, so handlers never block on the database.

    Collections: users · stats (single doc) · jobs (history) · chats · meta · audit

    Access model (v5): every person who talks to the bot gets a profile with an
    ``access`` block → {status, expires, plan, approved_by, …}. Only users whose
    status is *approved* and whose expiry is in the future (or 0 = lifetime) can
    translate. Owner + admins approve / extend / reject / ban.

    Bootstrapping on every start:
      • OWNER_ID           → owner (lifetime access, can promote admins)
      • ADMIN_USERS        → admins (lifetime access, can approve users)
      • AUTHORIZED_USERS   → lifetime approved users
      • MongoDB            → previously saved users / stats / history / audit
    """

    def __init__(self):
        self.owner_id: int = OWNER_ID
        self.users: Dict[int, dict] = {}
        self.chats: Dict[int, dict] = {}
        self.stats = {"jobs": 0, "parts": 0, "chars": 0, "failed": 0, "cancelled": 0}
        self.history: Dict[int, Deque[dict]] = {}          # uid → recent jobs (newest first)
        self.audit_log: Deque[dict] = deque(maxlen=AUDIT_LIMIT)
        self.booted = time.time()
        self.db = None                                      # motor database (or None)
        self.connected = False
        self._writes: Optional[asyncio.Queue] = None
        self._writer_task: Optional[asyncio.Task] = None
        # storage budget bookkeeping (free-tier guard)
        self.db_size_mb: float = 0.0                        # dataSize + indexSize, refreshed by check_budget()
        self.db_job_docs: int = 0
        self.db_checked: float = 0.0
        self.db_pruned_total: int = 0
        self._history_writes = 0                            # inserts since last budget check
        self._apply_env_roles()

    def _apply_env_roles(self) -> None:
        """Env-configured people always win: owner → admins → lifetime users."""
        if OWNER_ID:
            u = self.ensure_user(OWNER_ID, "Owner", save=False)
            u["role"] = "owner"
            self._grant(u, 0, 0, "env")
        for uid in ADMIN_USERS:
            if uid == OWNER_ID:
                continue
            u = self.ensure_user(uid, "Admin", save=False)
            u["role"] = "admin"
            self._grant(u, 0, 0, "env")
        for uid in AUTHORIZED_USERS:
            u = self.ensure_user(uid, "Pre-authorized", save=False)
            if u["access"]["status"] != ACCESS_APPROVED or u["access"].get("expires"):
                self._grant(u, 0, 0, "env")

    # ── MongoDB lifecycle ──────────────────────────────────────────────
    async def connect(self) -> bool:
        if not MONGO_URI:
            log.info("MONGO_URI not set → running with in-memory storage only")
            return False
        if not HAS_MOTOR:
            log.error("MONGO_URI is set but 'motor' is not installed → in-memory storage only")
            return False
        try:
            client = AsyncIOMotorClient(MONGO_URI, serverSelectionTimeoutMS=8000, appname="noveltranslator")
            await client.admin.command("ping")
            self.db = client[MONGO_DB]
            await self._ensure_indexes()
            await self._load()
            self._writes = asyncio.Queue()
            self._writer_task = asyncio.create_task(self._writer(), name="mongo_writer")
            self.connected = True
            log.info("MongoDB connected (db=%s, users=%d)", MONGO_DB, len(self.users))
            try:
                await self.check_budget(force=True)
            except Exception as e:
                log.debug("initial budget check skipped: %s", e)
            return True
        except Exception as e:
            log.error("MongoDB connection failed (%s) → falling back to in-memory storage", e)
            self.db = None
            return False

    async def _ensure_indexes(self) -> None:
        """Create the `jobs` indexes. The TTL index is (re)created whenever the
        configured DB_JOB_TTL_DAYS differs from the one already on the server —
        MongoDB refuses to change expireAfterSeconds via createIndex."""
        want = DB_JOB_TTL_DAYS * 24 * 3600
        try:
            await self.db.jobs.create_index([("uid", 1), ("ts", -1)], name="uid_ts")
        except Exception as e:
            log.debug("index uid_ts: %s", e)
        try:
            existing = await self.db.jobs.index_information()
        except Exception:
            existing = {}
        for name, info in existing.items():
            keys = info.get("key") or []
            if keys and keys[0][0] == "ts" and len(keys) == 1 and "expireAfterSeconds" in info:
                if int(info["expireAfterSeconds"]) == want:
                    return
                try:
                    await self.db.jobs.drop_index(name)
                    log.info("jobs TTL index changed %ss → %ss", info["expireAfterSeconds"], want)
                except Exception as e:
                    log.debug("drop_index %s: %s", name, e)
        try:
            await self.db.jobs.create_index("ts", name="ts_ttl", expireAfterSeconds=want)
        except Exception as e:
            log.warning("jobs TTL index: %s", e)

    async def _load(self) -> None:
        """Merge persisted state into the RAM cache. Env-configured users win."""
        migrated = 0
        async for doc in self.db.users.find({}):
            try:
                uid = int(doc["_id"])
            except (KeyError, TypeError, ValueError):
                continue
            doc.pop("_id", None)
            merged = {**self._default_user(doc.get("name", "User")), **doc}
            merged["stats"] = {**{"jobs": 0, "parts": 0, "chars": 0}, **(doc.get("stats") or {})}
            legacy = "access" not in doc
            if legacy:
                # v4 user (unlocked with the old security code) → migrate
                merged["access"] = self._default_access()
                if LEGACY_USERS == "keep":
                    self._grant(merged, 0, 0, "env", plan="legacy")
                migrated += 1
            else:
                merged["access"] = {**self._default_access(), **(doc.get("access") or {})}
            if merged.get("role") not in ("owner", "admin", "user"):
                merged["role"] = "user"
            self.users[uid] = merged
            if legacy:
                self._save_user(uid)
        if migrated:
            log.info("Migrated %d legacy users to the approval system (LEGACY_USERS=%s)", migrated, LEGACY_USERS)
        # env roles are re-applied on top of whatever was stored
        self._apply_env_roles()
        s = await self.db.stats.find_one({"_id": "global"})
        if s:
            for k in self.stats:
                self.stats[k] = int(s.get(k, 0) or 0)
        meta = await self.db.meta.find_one({"_id": "meta"})
        if meta and not self.owner_id and meta.get("owner_id"):
            self.owner_id = int(meta["owner_id"])
            if self.owner_id in self.users:
                self.users[self.owner_id]["role"] = "owner"
        try:
            cursor = self.db.audit.find({}, {"_id": 0}).sort("ts", -1).limit(AUDIT_LIMIT)
            rows = [d async for d in cursor]
            for d in reversed(rows):
                self.audit_log.append(d)
        except Exception as e:
            log.debug("audit load: %s", e)
        async for doc in self.db.chats.find({}):
            try:
                self.chats[int(doc["_id"])] = {"title": doc.get("title", "Chat"), "type": doc.get("type", "group"),
                                               "added": doc.get("added", 0)}
            except (KeyError, TypeError, ValueError):
                pass
        cursor = self.db.jobs.find({}, {"_id": 0}).sort("ts", -1).limit(HISTORY_LIMIT * 20)
        async for doc in cursor:
            dq = self.history.setdefault(int(doc.get("uid", 0)), deque(maxlen=HISTORY_LIMIT))
            if len(dq) < HISTORY_LIMIT:
                dq.append(doc)

    async def _writer(self) -> None:
        while True:
            op = await self._writes.get()
            try:
                await op()
            except Exception as e:
                log.warning("mongo write failed: %s", e)
            finally:
                self._writes.task_done()

    async def flush(self, timeout: float = 5.0) -> None:
        if self._writes:
            try:
                await asyncio.wait_for(self._writes.join(), timeout)
            except asyncio.TimeoutError:
                pass
        if self._writer_task:
            self._writer_task.cancel()

    def _enqueue(self, coro_factory) -> None:
        if self.db is not None and self._writes is not None:
            self._writes.put_nowait(coro_factory)

    # ── free-tier storage budget ───────────────────────────────────────
    async def db_usage(self) -> Tuple[float, int]:
        """Return (size_mb, job_docs). size = data + indexes of the whole database."""
        if self.db is None:
            return 0.0, 0
        st = await self.db.command("dbStats")                          # byte-precise, then convert
        size_mb = (float(st.get("dataSize", 0) or 0) + float(st.get("indexSize", 0) or 0)) / (1024 * 1024)
        docs = await self.db.jobs.estimated_document_count()
        return round(size_mb, 2), int(docs)

    async def prune_jobs(self, keep_docs: Optional[int] = None) -> int:
        """Trim the `jobs` collection: per-user cap (HISTORY_LIMIT) first, then a
        global cap so the collection can never outgrow the free tier."""
        if self.db is None:
            return 0
        removed = 0
        keep_docs = keep_docs if keep_docs is not None else DB_MAX_JOB_DOCS
        # 1) per-user: anything older than the user's HISTORY_LIMIT-th newest job
        pipeline = [{"$group": {"_id": "$uid", "n": {"$sum": 1}}}, {"$match": {"n": {"$gt": HISTORY_LIMIT}}}]
        async for grp in self.db.jobs.aggregate(pipeline):
            uid = grp["_id"]
            cutoff = await self.db.jobs.find({"uid": uid}, {"ts": 1}).sort("ts", -1) \
                .skip(HISTORY_LIMIT - 1).limit(1).to_list(1)
            if cutoff:
                r = await self.db.jobs.delete_many({"uid": uid, "ts": {"$lt": cutoff[0]["ts"]}})
                removed += r.deleted_count
        # 2) global cap
        total = await self.db.jobs.estimated_document_count()
        if total > keep_docs:
            cutoff = await self.db.jobs.find({}, {"ts": 1}).sort("ts", -1).skip(keep_docs - 1).limit(1).to_list(1)
            if cutoff:
                r = await self.db.jobs.delete_many({"ts": {"$lt": cutoff[0]["ts"]}})
                removed += r.deleted_count
        # 3) visitors never seen for a year, with zero jobs and no active access are just noise
        stale = int(time.time()) - 365 * 24 * 3600
        try:
            q = {"last_seen": {"$lt": stale}, "stats.jobs": 0, "role": "user",
                 "access.status": {"$nin": [ACCESS_APPROVED, ACCESS_BANNED]}}
            r = await self.db.users.delete_many(q)
            for uid in [u for u, d in list(self.users.items())
                        if d.get("last_seen", 0) < stale and not d.get("stats", {}).get("jobs")
                        and d.get("role") == "user" and u not in AUTHORIZED_USERS
                        and d.get("access", {}).get("status") not in (ACCESS_APPROVED, ACCESS_BANNED)]:
                self.users.pop(uid, None)
            removed += r.deleted_count
            # audit log: keep the collection tiny
            n_audit = await self.db.audit.estimated_document_count()
            if n_audit > AUDIT_LIMIT * 2:
                cutoff = await self.db.audit.find({}, {"ts": 1}).sort("ts", -1).skip(AUDIT_LIMIT).limit(1).to_list(1)
                if cutoff:
                    r = await self.db.audit.delete_many({"ts": {"$lt": cutoff[0]["ts"]}})
                    removed += r.deleted_count
        except Exception as e:
            log.debug("stale user prune: %s", e)
        if removed:
            self.db_pruned_total += removed
            log.info("DB prune: removed %d documents", removed)
        return removed

    async def check_budget(self, force: bool = False) -> None:
        """Keep the database inside DB_BUDGET_MB. Runs from the janitor (and
        after bursts of history inserts). Never raises."""
        if self.db is None:
            return
        if not force and time.time() - self.db_checked < 600 and self._history_writes < 50:
            return
        self._history_writes = 0
        try:
            size, docs = await self.db_usage()
            self.db_size_mb, self.db_job_docs, self.db_checked = size, docs, time.time()
            if docs > DB_MAX_JOB_DOCS or size > DB_BUDGET_MB * 0.8:
                await self.prune_jobs()
            # still over budget → shrink aggressively, halving the job cap each pass
            keep = DB_MAX_JOB_DOCS
            while size > DB_BUDGET_MB and keep > 100:
                keep //= 2
                await self.prune_jobs(keep_docs=keep)
                size, docs = await self.db_usage()
            if size > DB_BUDGET_MB:
                log.warning("MongoDB is %.1f MB (> budget %d MB) even after pruning — check the Atlas dashboard",
                            size, DB_BUDGET_MB)
            self.db_size_mb, self.db_job_docs = size, docs
        except Exception as e:
            log.warning("budget check failed: %s", e)

    def budget_info(self) -> dict:
        pct = round(100 * self.db_size_mb / DB_BUDGET_MB, 1) if DB_BUDGET_MB else 0
        return {"size_mb": self.db_size_mb, "budget_mb": DB_BUDGET_MB, "percent": pct,
                "job_docs": self.db_job_docs, "max_job_docs": DB_MAX_JOB_DOCS,
                "ttl_days": DB_JOB_TTL_DAYS, "pruned": self.db_pruned_total,
                "checked": int(self.db_checked)}

    # ── persistence helpers (fire-and-forget) ──────────────────────────
    def _save_user(self, uid: int) -> None:
        u = self.users.get(uid)
        if u is None:
            return
        snapshot = json.loads(json.dumps(u))
        self._enqueue(lambda: self.db.users.replace_one({"_id": uid}, snapshot, upsert=True))

    def _delete_user(self, uid: int) -> None:
        self._enqueue(lambda: self.db.users.delete_one({"_id": uid}))

    def _save_stats(self) -> None:
        snapshot = dict(self.stats)
        self._enqueue(lambda: self.db.stats.replace_one({"_id": "global"}, snapshot, upsert=True))

    def _save_meta(self) -> None:
        snapshot = {"owner_id": self.owner_id, "updated": int(time.time())}
        self._enqueue(lambda: self.db.meta.update_one({"_id": "meta"}, {"$set": snapshot}, upsert=True))

    @staticmethod
    def _default_access() -> dict:
        return {"status": ACCESS_NONE, "expires": 0, "plan": "", "approved_by": 0, "approved_at": 0,
                "requested_at": 0, "requests": 0, "note": "", "reason": "", "updated": 0,
                "reminded": 0, "history": []}

    @classmethod
    def _default_user(cls, name: str) -> dict:
        return {"name": name, "username": "", "joined": int(time.time()), "role": "user",
                "lang": DEFAULT_LANG, "fmt": DEFAULT_FORMAT, "split": DEFAULT_SPLIT_KB,
                "last_seen": int(time.time()),
                "stats": {"jobs": 0, "parts": 0, "chars": 0},
                "access": cls._default_access()}

    # ── users ──────────────────────────────────────────────────────────
    def user(self, uid: int) -> Optional[dict]:
        return self.users.get(uid)

    def ensure_user(self, uid: int, name: str, username: str = "", save: bool = True) -> dict:
        """Create a profile on first contact / refresh the display name."""
        u = self.users.get(uid)
        changed = False
        if u is None:
            u = self._default_user(name or "User")
            self.users[uid] = u
            changed = True
        else:
            if name and u.get("name") != name and name not in ("Owner", "Admin", "Pre-authorized", "User"):
                u["name"] = name
                changed = True
            if "access" not in u:
                u["access"] = self._default_access()
                changed = True
        if username and u.get("username") != username:
            u["username"] = username
            changed = True
        if uid == self.owner_id and u.get("role") != "owner":
            u["role"] = "owner"
            changed = True
        if changed and save:
            self._save_user(uid)
        return u

    def touch(self, uid: int) -> None:
        u = self.users.get(uid)
        if u and time.time() - u.get("last_seen", 0) > 600:
            u["last_seen"] = int(time.time())
            self._save_user(uid)

    # ── roles ──────────────────────────────────────────────────────────
    def is_owner(self, uid: int) -> bool:
        return bool(self.owner_id) and uid == self.owner_id

    def is_admin(self, uid: int) -> bool:
        """Owner or promoted admin — may approve / reject / extend users."""
        if self.is_owner(uid):
            return True
        u = self.users.get(uid)
        return bool(u and u.get("role") == "admin")

    def role(self, uid: int) -> str:
        if self.is_owner(uid):
            return "owner"
        u = self.users.get(uid)
        return u.get("role", "user") if u else "user"

    def set_admin(self, uid: int, make_admin: bool, by: int) -> bool:
        """Promote / demote. Admins automatically get lifetime access."""
        if self.is_owner(uid):
            return False
        u = self.ensure_user(uid, "Admin" if make_admin else "User")
        if make_admin:
            if u.get("role") == "admin":
                return False
            u["role"] = "admin"
            self._grant(u, 0, by, "admin", plan="forever")
            self.audit("promote", by, uid)
        else:
            if u.get("role") != "admin":
                return False
            u["role"] = "user"
            self.audit("demote", by, uid)
        self._save_user(uid)
        return True

    def admins(self) -> List[int]:
        ids = [uid for uid, u in self.users.items() if u.get("role") == "admin"]
        if self.owner_id:
            ids.insert(0, self.owner_id)
        return ids

    # ── access / approval ──────────────────────────────────────────────
    def access(self, uid: int) -> dict:
        """Current access block with lazy expiry (approved + past expiry → expired)."""
        u = self.users.get(uid)
        if not u:
            return self._default_access()
        a = u.setdefault("access", self._default_access())
        if a["status"] == ACCESS_APPROVED and a.get("expires") and a["expires"] <= time.time() \
                and not self.is_owner(uid):
            a["status"] = ACCESS_EXPIRED
            a["updated"] = int(time.time())
            self._push_hist(a, {"action": "expired"})
            self._save_user(uid)
            self.audit("expired", 0, uid)
        return a

    def status(self, uid: int) -> str:
        if self.is_owner(uid):
            return ACCESS_APPROVED
        return self.access(uid)["status"]

    def is_authorized(self, uid: int) -> bool:
        """Can this person use the translator right now?"""
        if self.is_owner(uid):
            return True
        if uid not in self.users:
            return PUBLIC_MODE
        a = self.access(uid)
        if a["status"] == ACCESS_BANNED:
            return False
        if PUBLIC_MODE:
            return True
        return a["status"] == ACCESS_APPROVED

    def is_banned(self, uid: int) -> bool:
        u = self.users.get(uid)
        return bool(u and u.get("access", {}).get("status") == ACCESS_BANNED)

    @staticmethod
    def _push_hist(a: dict, entry: dict) -> None:
        h = a.setdefault("history", [])
        h.append({"ts": int(time.time()), **entry})
        if len(h) > 12:
            del h[:-12]

    def _grant(self, u: dict, expires: int, by: int, source: str, plan: str = "") -> None:
        a = u.setdefault("access", self._default_access())
        a.update({"status": ACCESS_APPROVED, "expires": int(expires or 0), "approved_by": by,
                  "approved_at": int(time.time()), "plan": plan or ("forever" if not expires else a.get("plan", "")),
                  "updated": int(time.time()), "reminded": 0, "reason": ""})
        if source != "env":
            self._push_hist(a, {"action": "approve", "by": by, "expires": int(expires or 0), "plan": a["plan"]})

    def request_access(self, uid: int, name: str, username: str = "", note: str = "") -> Tuple[bool, str]:
        """User asks for approval. Returns (ok, reason)."""
        self.ensure_user(uid, name, username)
        a = self.access(uid)
        now = int(time.time())
        if a["status"] == ACCESS_BANNED:
            return False, "banned"
        if a["status"] == ACCESS_APPROVED or self.is_owner(uid):
            return False, "approved"
        if a["status"] == ACCESS_PENDING:
            return False, "pending"
        if a["status"] == ACCESS_REJECTED and REJECT_COOLDOWN_H and a.get("updated", 0) + REJECT_COOLDOWN_H * 3600 > now:
            return False, "cooldown"
        a.update({"status": ACCESS_PENDING, "requested_at": now, "requests": int(a.get("requests", 0)) + 1,
                  "note": (note or "")[:200], "updated": now})
        self._push_hist(a, {"action": "request"})
        self._save_user(uid)
        self.audit("request", uid, uid, note=(note or "")[:80])
        return True, "ok"

    def approve(self, uid: int, spec: str, by: int, name: str = "", extend: bool = False) -> Optional[dict]:
        """Approve (or extend) access. `spec` is a duration (1m, 1y, forever, 45d, 2026-12-31).
        extend=True adds the duration on top of a still-valid expiry."""
        u = self.ensure_user(uid, name or "User")
        a = self.access(uid)
        if a["status"] == ACCESS_BANNED:
            return None
        start = time.time()
        if extend and a["status"] == ACCESS_APPROVED:
            if not a.get("expires"):
                return a                               # already lifetime — nothing to extend
            start = max(start, a["expires"])
        expires = parse_duration(spec, start)
        if expires is None:
            return None
        plan = spec if spec in PLAN_PRESETS else ("forever" if expires == 0 else spec)
        self._grant(u, expires, by, "admin", plan=plan)
        self._save_user(uid)
        self.audit("extend" if extend else "approve", by, uid, plan=plan, expires=expires)
        return a

    def reject(self, uid: int, by: int, reason: str = "") -> bool:
        u = self.users.get(uid)
        if not u or self.is_owner(uid):
            return False
        a = self.access(uid)
        a.update({"status": ACCESS_REJECTED, "reason": (reason or "")[:200], "updated": int(time.time()),
                  "expires": 0, "plan": ""})
        self._push_hist(a, {"action": "reject", "by": by, "reason": (reason or "")[:80]})
        self._save_user(uid)
        self.audit("reject", by, uid, reason=(reason or "")[:80])
        return True

    def revoke(self, uid: int, by: int = 0, reason: str = "") -> bool:
        """Remove access but keep the profile (stats / history stay)."""
        u = self.users.get(uid)
        if not u or self.is_owner(uid):
            return False
        a = self.access(uid)
        if a["status"] not in (ACCESS_APPROVED, ACCESS_EXPIRED, ACCESS_PENDING):
            return False
        if u.get("role") == "admin":
            u["role"] = "user"
        a.update({"status": ACCESS_NONE, "expires": 0, "plan": "", "reason": (reason or "")[:200],
                  "updated": int(time.time())})
        self._push_hist(a, {"action": "revoke", "by": by})
        self._save_user(uid)
        self.audit("revoke", by, uid, reason=(reason or "")[:80])
        return True

    def ban(self, uid: int, by: int, reason: str = "") -> bool:
        if self.is_owner(uid):
            return False
        u = self.ensure_user(uid, "User")
        if u.get("role") == "admin":
            u["role"] = "user"
        a = self.access(uid)
        a.update({"status": ACCESS_BANNED, "expires": 0, "plan": "", "reason": (reason or "")[:200],
                  "updated": int(time.time())})
        self._push_hist(a, {"action": "ban", "by": by, "reason": (reason or "")[:80]})
        self._save_user(uid)
        self.audit("ban", by, uid, reason=(reason or "")[:80])
        return True

    def unban(self, uid: int, by: int) -> bool:
        u = self.users.get(uid)
        if not u or u.get("access", {}).get("status") != ACCESS_BANNED:
            return False
        u["access"].update({"status": ACCESS_NONE, "reason": "", "updated": int(time.time())})
        self._push_hist(u["access"], {"action": "unban", "by": by})
        self._save_user(uid)
        self.audit("unban", by, uid)
        return True

    def approved_users(self) -> List[int]:
        return [uid for uid in list(self.users) if self.is_authorized(uid)]

    def pending_users(self) -> List[Tuple[int, dict]]:
        rows = [(uid, u) for uid, u in self.users.items() if u.get("access", {}).get("status") == ACCESS_PENDING]
        rows.sort(key=lambda kv: kv[1]["access"].get("requested_at", 0))
        return rows

    def count_by_status(self) -> Dict[str, int]:
        out = {k: 0 for k in ACCESS_LABELS}
        for uid in list(self.users):
            out[self.status(uid)] = out.get(self.status(uid), 0) + 1
        return out

    def sweep_expired(self) -> List[int]:
        """Mark approved users whose time is over as expired. Returns their ids."""
        out = []
        now = time.time()
        for uid, u in list(self.users.items()):
            a = u.get("access") or {}
            if a.get("status") == ACCESS_APPROVED and a.get("expires") and a["expires"] <= now and not self.is_owner(uid):
                self.access(uid)                        # lazy expiry does the bookkeeping
                out.append(uid)
        return out

    def due_reminders(self) -> List[Tuple[int, int]]:
        """Approved users expiring within EXPIRY_REMINDER_DAYS that were not reminded yet."""
        if not EXPIRY_REMINDER_DAYS:
            return []
        now = time.time()
        out = []
        for uid, u in self.users.items():
            a = u.get("access") or {}
            exp = a.get("expires") or 0
            if a.get("status") == ACCESS_APPROVED and exp and now < exp <= now + EXPIRY_REMINDER_DAYS * 86400 \
                    and not a.get("reminded"):
                a["reminded"] = int(now)
                self._save_user(uid)
                out.append((uid, exp))
        return out

    # ── audit log ──────────────────────────────────────────────────────
    def audit(self, action: str, actor: int, target: int, **extra) -> None:
        entry = {"ts": int(time.time()), "action": action, "actor": int(actor or 0), "target": int(target or 0),
                 "actor_name": ((self.users.get(actor) or {}).get("name", "") if actor else "system"),
                 "target_name": (self.users.get(target) or {}).get("name", "")}
        for k, v in extra.items():
            if v not in (None, "") and (v != 0 or k == "expires"):
                entry[k] = v
        self.audit_log.append(entry)
        snap = dict(entry)
        self._enqueue(lambda: self.db.audit.insert_one(dict(snap)))

    def recent_audit(self, limit: int = 40) -> List[dict]:
        return list(self.audit_log)[-limit:][::-1]

    # ── prefs & stats ──────────────────────────────────────────────────
    def pref(self, uid: int, key: str, default=None):
        u = self.users.get(uid)
        return u.get(key, default) if u else default

    def set_pref(self, uid: int, key: str, value):
        u = self.users.get(uid) or self.ensure_user(uid, "User")
        u[key] = value
        self._save_user(uid)

    def bump(self, uid: int, parts: int, chars: int, failed=False, cancelled=False):
        u = self.users.get(uid)
        if u and not failed and not cancelled:
            u["stats"]["jobs"] += 1
            u["stats"]["parts"] += parts
            u["stats"]["chars"] += chars
            self._save_user(uid)
        if failed:
            self.stats["failed"] += 1
        elif cancelled:
            self.stats["cancelled"] += 1
        else:
            self.stats["jobs"] += 1
            self.stats["parts"] += parts
            self.stats["chars"] += chars
        self._save_stats()

    # ── job history ────────────────────────────────────────────────────
    def add_history(self, uid: int, entry: dict) -> None:
        entry = {"uid": uid, "ts": int(time.time()), **entry}
        dq = self.history.setdefault(uid, deque(maxlen=HISTORY_LIMIT))
        dq.appendleft(entry)
        snapshot = {k: v for k, v in entry.items() if k != "error" or v}
        if isinstance(snapshot.get("error"), str):
            snapshot["error"] = snapshot["error"][:200]      # keep job docs tiny
        self._history_writes += 1
        self._enqueue(lambda: self.db.jobs.insert_one(dict(snapshot)))

    def user_history(self, uid: int, limit: int = 20) -> List[dict]:
        return list(self.history.get(uid, ()))[:limit]

    def recent_history(self, limit: int = 20) -> List[dict]:
        allj = [j for dq in self.history.values() for j in dq]
        allj.sort(key=lambda j: j.get("ts", 0), reverse=True)
        return allj[:limit]

    # ── chats ──────────────────────────────────────────────────────────
    def track_chat(self, chat_id: int, title: str, ctype: str):
        doc = {"title": title, "type": ctype, "added": int(time.time())}
        self.chats[chat_id] = doc
        self._enqueue(lambda: self.db.chats.replace_one({"_id": chat_id}, doc, upsert=True))

    def untrack_chat(self, chat_id: int):
        self.chats.pop(chat_id, None)
        self._enqueue(lambda: self.db.chats.delete_one({"_id": chat_id}))


MemoryStore = Store          # backwards-compatible alias
store = Store()
QUEUE_REPO = None            # database.MongoDatabase view over store.db (shared jobs_queue + GridFS)
EMBEDDED = None              # worker.TranslationWorker running inside the master (EMBEDDED_WORKER=1)


# ═══════════════════════════════════════════════════════════════════════════
# 🎨 UI HELPERS
# ═══════════════════════════════════════════════════════════════════════════
DIV = "━━━━━━━━━━━━━━━━━━━━"

def esc(s) -> str:
    return html.escape(str(s if s is not None else ""), quote=False)

def b(s) -> str:
    return f"<b>{esc(s)}</b>"

def code(s) -> str:
    return f"<code>{esc(s)}</code>"

def header(title: str, icon: str = "📚") -> str:
    return f"{icon} <b>{esc(title)}</b>\n{DIV}\n"

def progress_bar(ratio: float, width: int = 12) -> str:
    ratio = max(0.0, min(1.0, ratio))
    filled = int(round(ratio * width))
    return "▰" * filled + "▱" * (width - filled)

def fmt_time(seconds: float) -> str:
    seconds = int(max(0, seconds))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"

def fmt_size(num: float) -> str:
    num = float(num or 0)
    for unit in ("B", "KB", "MB", "GB"):
        if num < 1024 or unit == "GB":
            return f"{num:.0f} {unit}" if unit == "B" else f"{num:.1f} {unit}"
        num /= 1024
    return f"{num:.1f} GB"

def fmt_int(n) -> str:
    try:
        return f"{int(n):,}"
    except (TypeError, ValueError):
        return str(n)

def lang_label(code_: str) -> str:
    name, flag = LANGUAGES.get(code_, (code_.upper(), "🌐"))
    return f"{flag} {name}"

def split_label(kb: int) -> str:
    if not kb:
        return "No split"
    return f"{kb / 1024:g} MB" if kb >= 1024 else f"{kb} KB"

def kb_rows(buttons: List[InlineKeyboardButton], per_row: int = 2) -> List[List[InlineKeyboardButton]]:
    return [buttons[i:i + per_row] for i in range(0, len(buttons), per_row)]

def app_inline_button() -> Optional[InlineKeyboardButton]:
    if MINI_APP_URL.startswith("https://"):
        return InlineKeyboardButton("📱 Open Mini App", web_app=WebAppInfo(url=MINI_APP_URL))
    return None

def home_keyboard(uid: int) -> InlineKeyboardMarkup:
    rows = []
    if (btn := app_inline_button()):
        rows.append([btn])
    rows += [
        [InlineKeyboardButton("⚙️ Settings", callback_data="nav:settings"),
         InlineKeyboardButton("📊 My Stats", callback_data="nav:mystats")],
        [InlineKeyboardButton("📋 Queue", callback_data="nav:queue"),
         InlineKeyboardButton("ℹ️ Help", callback_data="nav:help")],
    ]
    if store.is_owner(uid):
        rows.append([InlineKeyboardButton("👑 Owner Panel", callback_data="nav:owner")])
    return InlineKeyboardMarkup(rows)

def back_home_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([[InlineKeyboardButton("🏠 Home", callback_data="nav:home")]])


# ── Reply keyboard (persistent bottom menu) — hierarchical & role based ────
# The bottom keyboard is organised as small *pages* instead of one big grid:
#
#   MAIN      →  📱 Mini App | 🛠 Tools | ⚙️ Settings | 👑 Admin (owner/admin) | ℹ️ Help
#   TOOLS     →  📋 Queue | 📊 My Stats | 🛑 Cancel Job | 🆔 My ID | ◀️ Back
#   SETTINGS  →  🌐 Language | 📄 Format | ✂️ Split | 📱 Mini App | ◀️ Back
#   ADMIN     →  ⏳ Requests | 👥 Users | 👑 Owner Panel | 📣 Broadcast | 🔗 Links | ◀️ Back
#   LOCKED    →  🙋 Request Access | 🆔 My ID          (people without approval)
#
# Tapping a menu button *replaces* the keyboard with that sub-page, so only
# 3–5 buttons are ever visible at once. Every leaf maps to a command.
BTN_HOME      = "🏠 Home"
BTN_APP       = "📱 Mini App"
BTN_TOOLS     = "🛠 Tools"
BTN_SETTINGS  = "⚙️ Settings"
BTN_ADMIN     = "👑 Admin"
BTN_HELP      = "ℹ️ Help"
BTN_BACK      = "◀️ Back"

BTN_QUEUE     = "📋 Queue"
BTN_MYSTATS   = "📊 My Stats"
BTN_CANCEL    = "🛑 Cancel Job"
BTN_MYID      = "🆔 My ID"

BTN_LANG      = "🌐 Language"
BTN_FORMAT    = "📄 Format"
BTN_SPLIT     = "✂️ Split size"

BTN_OWNER     = "👑 Owner Panel"
BTN_USERS     = "👥 Users"
BTN_PENDING   = "⏳ Requests"
BTN_BROADCAST = "📣 Broadcast"
BTN_LINKS     = "🔗 Links"
BTN_REQUEST   = "🙋 Request Access"
BTN_MYACCESS  = "🎫 My Access"

# Leaf buttons → command they trigger
USER_BUTTONS: Dict[str, str] = {
    BTN_HOME: "start", BTN_HELP: "help", BTN_QUEUE: "queue", BTN_MYSTATS: "mystats",
    BTN_CANCEL: "cancel", BTN_MYID: "id", BTN_SETTINGS: "settings", BTN_MYACCESS: "access",
    BTN_LANG: "setlang", BTN_FORMAT: "setformat", BTN_SPLIT: "setsplit", BTN_APP: "app",
}
# Admin buttons (owner + promoted admins)
ADMIN_BUTTONS: Dict[str, str] = {
    BTN_PENDING: "pending", BTN_USERS: "users",
}
# Owner-only buttons
OWNER_BUTTONS: Dict[str, str] = {
    BTN_OWNER: "stats", BTN_BROADCAST: "broadcast", BTN_LINKS: "links",
}
ALL_BUTTONS: Dict[str, str] = {**USER_BUTTONS, **ADMIN_BUTTONS, **OWNER_BUTTONS}
# Buttons that only switch the keyboard page (no command)
MENU_BUTTONS = {BTN_TOOLS: "tools", BTN_ADMIN: "admin", BTN_BACK: "main"}
KB_PAGE: Dict[int, str] = {}      # chat_id → current keyboard page (RAM only)

def _app_button() -> KeyboardButton:
    """Bottom-keyboard *Mini App* button.

    Deliberately a plain text button (no ``web_app=``): tapping it is handled
    like the ``/app`` command and the bot replies with a message carrying an
    inline **📱 Open Mini App** button. This keeps the flow identical to the
    command, always shows a fresh link and avoids Telegram's confusing
    behaviour where a web_app reply-button opens the app without any trace
    in the chat.
    """
    return KeyboardButton(BTN_APP)

def reply_keyboard(uid: int, page: str = "main") -> ReplyKeyboardMarkup:
    """Bottom keyboard for the given page. Owner/admins get the extra Admin page."""
    owner = store.is_owner(uid)
    admin = store.is_admin(uid)
    if page == "tools":
        rows = [[KeyboardButton(BTN_QUEUE), KeyboardButton(BTN_MYSTATS)],
                [KeyboardButton(BTN_CANCEL), KeyboardButton(BTN_MYACCESS)],
                [KeyboardButton(BTN_MYID), KeyboardButton(BTN_BACK)]]
        ph = "🛠 Tools — pick an action…"
    elif page == "settings":
        rows = [[KeyboardButton(BTN_LANG), KeyboardButton(BTN_FORMAT)],
                [KeyboardButton(BTN_SPLIT), _app_button()],
                [KeyboardButton(BTN_BACK)]]
        ph = "⚙️ Settings — what to change?"
    elif page == "admin" and admin:
        rows = [[KeyboardButton(BTN_PENDING), KeyboardButton(BTN_USERS)]]
        if owner:
            rows.append([KeyboardButton(BTN_OWNER), KeyboardButton(BTN_BROADCAST)])
            rows.append([KeyboardButton(BTN_LINKS), KeyboardButton(BTN_BACK)])
        else:
            rows.append([KeyboardButton(BTN_BACK)])
        ph = "👑 Admin — approve requests, manage users…"
    else:
        page = "main"
        rows = [[_app_button()],
                [KeyboardButton(BTN_TOOLS), KeyboardButton(BTN_SETTINGS)],
                [KeyboardButton(BTN_ADMIN), KeyboardButton(BTN_HELP)] if admin else [KeyboardButton(BTN_HELP)]]
        ph = "📎 Send a document or open a menu…"
    return ReplyKeyboardMarkup(rows, resize_keyboard=True, is_persistent=True, placeholder=ph)

def locked_keyboard(status: str = ACCESS_NONE) -> ReplyKeyboardMarkup:
    """Minimal keyboard for people who are not approved (yet)."""
    if status in (ACCESS_PENDING, ACCESS_BANNED):
        rows = [[KeyboardButton(BTN_MYACCESS), KeyboardButton(BTN_MYID)]]
        ph = "⏳ Waiting for approval…" if status == ACCESS_PENDING else "🚫 Access blocked"
    else:
        rows = [[KeyboardButton(BTN_REQUEST)], [KeyboardButton(BTN_MYACCESS), KeyboardButton(BTN_MYID)]]
        ph = "🔒 Tap “Request Access” to get started…"
    return ReplyKeyboardMarkup(rows, resize_keyboard=True, is_persistent=True, placeholder=ph)

PAGE_TITLES = {
    "main": ("Main Menu", "🏠", "Choose a section below.\n📱 Mini App sends you a link to the full dashboard."),
    "tools": ("Tools", "🛠", "Queue · stats · cancel · access · your ID"),
    "settings": ("Settings", "⚙️", "Change your defaults for ⚡ Quick Start."),
    "admin": ("Admin", "👑", "Approve requests, manage users, broadcast."),
}

async def show_page(m: Message, page: str, note: str = "") -> None:
    """Swap the bottom keyboard to another page (short confirmation message)."""
    uid = m.from_user.id
    if page == "admin" and not store.is_admin(uid):
        page = "main"
    KB_PAGE[m.chat.id] = page
    title, icon, hint = PAGE_TITLES[page]
    await m.reply(f"{icon} {b(title)}\n<i>{esc(note or hint)}</i>", reply_markup=reply_keyboard(uid, page))


# ── Bot command menu (the “/” list) — role based via BotCommandScope ───────
USER_COMMANDS = [
    BotCommand("start",    "🏠 Home screen & menu"),
    BotCommand("app",      "📱 Open the Mini App dashboard"),
    BotCommand("settings", "⚙️ Default language, format, split size"),
    BotCommand("queue",    "📋 Current queue status"),
    BotCommand("cancel",   "🛑 Cancel your active / queued jobs"),
    BotCommand("mystats",  "📊 Your usage statistics"),
    BotCommand("access",   "🎫 Your access plan & expiry"),
    BotCommand("help",     "ℹ️ How to use the bot"),
    BotCommand("id",       "🆔 Show your Telegram ID"),
]
ADMIN_COMMANDS = USER_COMMANDS + [
    BotCommand("pending",   "⏳ Access requests waiting for approval"),
    BotCommand("users",     "👥 List users & their access"),
    BotCommand("approve",   "✅ /approve <id> [1w|1m|1y|forever|45d]"),
    BotCommand("extend",    "➕ /extend <id> <duration>"),
    BotCommand("reject",    "❌ /reject <id> [reason]"),
    BotCommand("revoke",    "🔒 /revoke <id> — remove access"),
    BotCommand("ban",       "🚫 /ban <id> [reason]  ·  /unban <id>"),
    BotCommand("userinfo",  "🔎 /userinfo <id> — profile & access"),
]
OWNER_COMMANDS = ADMIN_COMMANDS + [
    BotCommand("stats",     "👑 Owner panel / global stats"),
    BotCommand("admins",    "🛡 List admins · /addadmin <id> · /deladmin <id>"),
    BotCommand("audit",     "📜 Recent access actions"),
    BotCommand("broadcast", "📣 Message all users (text or reply)"),
    BotCommand("links",     "🔗 Invite links of admin chats"),
]
LOCKED_COMMANDS = [
    BotCommand("start",   "🔒 Request access to the bot"),
    BotCommand("request", "🙋 Ask the owner for approval"),
    BotCommand("access",  "🎫 Check your request status"),
    BotCommand("id",      "🆔 Show your Telegram ID"),
]
_COMMANDS_SET: set = set()   # chat ids that already have per-chat commands

async def set_global_commands() -> None:
    """Default scope: what a stranger / normal user sees in the “/” menu."""
    base = USER_COMMANDS if PUBLIC_MODE else LOCKED_COMMANDS
    for scope in (BotCommandScopeDefault(), BotCommandScopeAllPrivateChats()):
        try:
            await app.set_bot_commands(base, scope=scope)
        except Exception as e:
            log.debug("set_bot_commands(%s): %s", scope, e)

async def set_user_commands(uid: int, force: bool = False) -> None:
    """Per-chat scope: approved users → user commands, admins → + approval tools, owner → everything."""
    if not uid or (uid in _COMMANDS_SET and not force):
        return
    if store.is_owner(uid):
        cmds = OWNER_COMMANDS
    elif store.is_admin(uid):
        cmds = ADMIN_COMMANDS
    else:
        cmds = USER_COMMANDS
    try:
        await app.set_bot_commands(cmds, scope=BotCommandScopeChat(uid))
        _COMMANDS_SET.add(uid)
    except Exception as e:
        log.debug("set_bot_commands(chat %s): %s", uid, e)

async def clear_user_commands(uid: int) -> None:
    """Revoked user → falls back to the locked default menu."""
    _COMMANDS_SET.discard(uid)
    try:
        await app.delete_bot_commands(scope=BotCommandScopeChat(uid))
    except Exception as e:
        log.debug("delete_bot_commands(chat %s): %s", uid, e)

def user_prefs(uid: int) -> Tuple[str, str, int]:
    lang = store.pref(uid, "lang", DEFAULT_LANG)
    fmt = store.pref(uid, "fmt", DEFAULT_FORMAT)
    split = store.pref(uid, "split", DEFAULT_SPLIT_KB)
    if lang not in LANGUAGES: lang = DEFAULT_LANG
    if fmt not in OUTPUT_FORMATS: fmt = DEFAULT_FORMAT
    try: split = int(split)
    except (TypeError, ValueError): split = DEFAULT_SPLIT_KB
    return lang, fmt, split

def access_line(uid: int) -> str:
    """One-line access summary used on the home screen."""
    if store.is_owner(uid):
        return "👑 Owner · ♾ Lifetime"
    a = store.access(uid)
    role = store.role(uid)
    if a["status"] == ACCESS_APPROVED:
        prefix = "🛡 Admin · " if role == "admin" else ""
        return prefix + ("♾ Lifetime" if not a.get("expires") else expiry_label(a["expires"]))
    label, icon = ACCESS_LABELS.get(a["status"], ("Unknown", "❓"))
    return f"{icon} {label}"

def text_home(uid: int, name: str) -> str:
    lang, fmt, split = user_prefs(uid)
    return (
        header(BOT_NAME) +
        f"👋 Hello, {b(name)}!\n\n"
        f"Send me a document and I will\n"
        f"translate it into your language.\n\n"
        f"📥 {b('Input')}: {', '.join(sorted(e.lstrip('.').upper() for e in INPUT_EXTS))} (≤ {MAX_INPUT_MB} MB)\n"
        f"📤 {b('Output')}: TXT · DOCX · EPUB\n\n"
        f"{DIV}\n"
        f"🌐 Language: {b(lang_label(lang))}\n"
        f"📄 Format: {b(fmt.upper())}\n"
        f"✂️ Split: {b(split_label(split))}\n"
        f"🎫 Access: {b(access_line(uid))}\n"
        f"{DIV}\n"
        f"<i>Tip: change defaults in ⚙️ Settings\nand use ⚡ Quick Start next time.</i>"
        + ("\n<i>📱 Tap <b>Mini App</b> for the full dashboard.</i>" if MINI_APP_URL.startswith("https://") else "")
    )

def text_help() -> str:
    return (
        header("How it works", "ℹ️") +
        f"1️⃣ Send a {' / '.join(sorted(INPUT_EXTS))} file\n"
        "2️⃣ Choose language, format & split\n"
        "3️⃣ Watch the live dashboard\n"
        "4️⃣ Receive translated parts\n\n"
        f"{b('Commands')}\n"
        "/start – Home screen\n"
        "/app – Open the Mini App dashboard\n"
        "/settings – Default language, format, split\n"
        "/queue – Current queue status\n"
        "/cancel – Cancel your active job\n"
        "/mystats – Your usage statistics\n"
        "/access – Your access plan & expiry\n"
        "/help – This message\n"
        "/id – Your Telegram ID\n\n"
        f"{b('Menu')}\n"
        "<i>Use the ⌨️ buttons at the bottom of the chat.\n"
        "🛠 Tools / ⚙️ Settings open a sub-menu;\n"
        "◀️ Back returns to the main menu.</i>\n\n"
        f"{b('Split size')}\n"
        "Large books are delivered in parts.\n"
        f"Custom size: {MIN_SPLIT_KB} KB – {MAX_SPLIT_KB // 1024} MB\n"
        "e.g. <code>750</code>, <code>2 MB</code>, <code>900kb</code>\n\n"
        f"{b('Storage')}\n"
        + ("<i>Settings, statistics and job history are saved in MongoDB.</i>"
           if store.connected else
           "<i>No database connected — settings and statistics\nreset whenever the server restarts.</i>")
    )

def text_settings(uid: int) -> str:
    lang, fmt, split = user_prefs(uid)
    return (
        header("Settings", "⚙️") +
        f"🌐 Language: {b(lang_label(lang))}\n"
        f"📄 Format: {b(fmt.upper())}\n"
        f"✂️ Split size: {b(split_label(split))}\n"
        f"{DIV}\n"
        "<i>These are used by ⚡ Quick Start\nand pre-selected in the wizard.</i>"
        + ("" if store.connected else "\n<i>⚠️ No database — settings reset on restart.</i>")
    )

def settings_keyboard() -> InlineKeyboardMarkup:
    rows = [
        [InlineKeyboardButton("🌐 Language", callback_data="st:lang"),
         InlineKeyboardButton("📄 Format", callback_data="st:fmt")],
        [InlineKeyboardButton("✂️ Split size", callback_data="st:spl")],
    ]
    if (btn := app_inline_button()):
        rows.append([btn])
    rows.append([InlineKeyboardButton("🏠 Home", callback_data="nav:home")])
    return InlineKeyboardMarkup(rows)

def language_keyboard(prefix: str, selected: str, back_cb: str) -> InlineKeyboardMarkup:
    btns = [InlineKeyboardButton(f"{'✅ ' if c == selected else ''}{flag} {name}", callback_data=f"{prefix}{c}")
            for c, (name, flag) in LANGUAGES.items()]
    rows = kb_rows(btns, 3)
    rows.append([InlineKeyboardButton("◀️ Back", callback_data=back_cb)])
    return InlineKeyboardMarkup(rows)

def format_keyboard(prefix: str, selected: str, back_cb: str) -> InlineKeyboardMarkup:
    btns = [InlineKeyboardButton(f"{'✅ ' if f == selected else ''}{label}", callback_data=f"{prefix}{f}")
            for f, label in OUTPUT_FORMATS.items()]
    return InlineKeyboardMarkup(kb_rows(btns, 3) + [[InlineKeyboardButton("◀️ Back", callback_data=back_cb)]])

def split_keyboard(prefix: str, selected: int, back_cb: str, custom_cb: Optional[str]) -> InlineKeyboardMarkup:
    btns = [InlineKeyboardButton(f"{'✅ ' if kb == selected else ''}{label}", callback_data=f"{prefix}{kb}")
            for kb, label in SPLIT_PRESETS]
    if selected and selected not in dict(SPLIT_PRESETS):
        btns.append(InlineKeyboardButton(f"✅ {split_label(selected)}", callback_data=f"{prefix}{selected}"))
    rows = kb_rows(btns, 2)
    if custom_cb:
        rows.append([InlineKeyboardButton("✏️ Custom size", callback_data=custom_cb)])
    rows.append([InlineKeyboardButton("◀️ Back", callback_data=back_cb)])
    return InlineKeyboardMarkup(rows)


# ═══════════════════════════════════════════════════════════════════════════
# 📬 FLOOD-SAFE MESSAGE UTILITIES
# ═══════════════════════════════════════════════════════════════════════════
class LiveMessage:
    """Rate-limited, flood-safe editor for a single Telegram message."""

    def __init__(self, message: Optional[Message], min_interval: float = EDIT_INTERVAL):
        self.msg = message
        self.min_interval = min_interval
        self._last_edit = 0.0
        self._last_text = ""
        self._lock = asyncio.Lock()
        self.alive = message is not None

    async def update(self, text: str, keyboard: Optional[InlineKeyboardMarkup] = None, force: bool = False):
        if not self.alive:
            return
        now = time.time()
        if not force and (now - self._last_edit) < self.min_interval:
            return
        if text == self._last_text and not force:
            return
        async with self._lock:
            try:
                await self.msg.edit_text(text, reply_markup=keyboard, disable_web_page_preview=True)
                self._last_edit = time.time()
                self._last_text = text
            except MessageNotModified:
                self._last_edit = time.time()
                self._last_text = text
            except FloodWait as e:
                # Back off: do not attempt another edit until the wait expires
                self._last_edit = time.time() + float(e.value)
            except MessageIdInvalid:
                self.alive = False
            except RPCError as e:
                if "MESSAGE_ID_INVALID" in str(e) or "not modified" in str(e).lower():
                    self.alive = False
                else:
                    log.warning("edit failed: %s", e)
            except Exception as e:
                log.warning("edit failed: %s", e)

async def safe_send(chat_id: int, text: str, **kwargs) -> Optional[Message]:
    if not chat_id:
        return None
    for _ in range(3):
        try:
            return await app.send_message(chat_id, text, disable_web_page_preview=True, **kwargs)
        except FloodWait as e:
            await asyncio.sleep(min(float(e.value) + 1, 120))
        except Exception as e:
            log.debug("safe_send(%s) failed: %s", chat_id, e)
            return None
    return None

async def safe_send_document(chat_id: int, path: str, caption: str, **kwargs) -> Optional[Message]:
    if not chat_id:
        return None
    for attempt in range(4):
        try:
            return await app.send_document(chat_id, path, caption=caption[:1000], **kwargs)
        except FloodWait as e:
            await asyncio.sleep(min(float(e.value) + 2, 120))
        except RPCError as e:
            # reply_to_message_id pointing at a deleted topic etc. → retry without it
            if kwargs and attempt == 0:
                kwargs = {}
                continue
            log.warning("send_document(%s) failed: %s", chat_id, e)
            await asyncio.sleep(2)
        except Exception as e:
            log.warning("send_document(%s) failed: %s", chat_id, e)
            await asyncio.sleep(2)
    return None


# ═══════════════════════════════════════════════════════════════════════════
# 📖 TEXT EXTRACTION
# ═══════════════════════════════════════════════════════════════════════════
_BLOCK_TAGS = ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "blockquote", "div", "td", "pre"]

def extract_text_epub(path: str) -> str:
    try:
        try:
            book = epub.read_epub(path, options={"ignore_ncx": True})
        except TypeError:
            book = epub.read_epub(path)
    except Exception as e:
        log.warning("epub read failed: %s", e)
        return ""

    items = []
    for entry in getattr(book, "spine", []) or []:
        idref = entry[0] if isinstance(entry, (tuple, list)) else entry
        item = book.get_item_with_id(idref)
        if item is not None and item.get_type() == ebooklib.ITEM_DOCUMENT:
            items.append(item)
    if not items:
        items = [i for i in book.get_items() if i.get_type() == ebooklib.ITEM_DOCUMENT]

    sections, seen = [], set()
    for item in items:
        try:
            soup = BeautifulSoup(item.get_body_content(), "lxml")
        except Exception:
            try:
                soup = BeautifulSoup(item.get_body_content(), "html.parser")
            except Exception:
                continue
        for junk in soup(["script", "style", "nav", "svg"]):
            junk.decompose()
        blocks = []
        for el in soup.find_all(_BLOCK_TAGS):
            if el.find(_BLOCK_TAGS):
                continue
            txt = el.get_text(" ", strip=True)
            if txt:
                if el.name.startswith("h"):
                    blocks.append("")
                blocks.append(txt)
        text = "\n".join(blocks).strip() if blocks else soup.get_text("\n", strip=True)
        key = hash(text[:400])
        if len(text) > 30 and key not in seen:
            seen.add(key)
            sections.append(text)
    return "\n\n".join(sections)

def extract_text_txt(path: str) -> str:
    with open(path, "rb") as fh:
        raw = fh.read()
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        encodings = ("utf-16",)
    else:
        encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
    text = None
    for enc in encodings:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")

def extract_text_docx(path: str) -> str:
    try:
        d = docx.Document(path)
        parts = [p.text.rstrip() for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        log.warning("docx read failed: %s", e)
        return ""

def extract_text_pdf(path: str) -> str:
    if not HAS_PDF:
        return ""
    try:
        reader = PdfReader(path)
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                return ""
        pages = []
        for page in reader.pages:
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            if t.strip():
                pages.append(t.strip())
        return "\n\n".join(pages)
    except Exception as e:
        log.warning("pdf read failed: %s", e)
        return ""

EXTRACTORS = {".epub": extract_text_epub, ".txt": extract_text_txt,
              ".docx": extract_text_docx, ".pdf": extract_text_pdf}

def normalise_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\u00ad", "")
    text = re.sub(r"[ \t\u00a0]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ═══════════════════════════════════════════════════════════════════════════
# ✂️ SPLITTING & SMART CHUNKING
# ═══════════════════════════════════════════════════════════════════════════
def split_text_by_size(text: str, max_kb: int, factor: float = 1.0) -> List[str]:
    """Split on line boundaries so that each part's *estimated output size*
    (len_utf8 × factor) stays below max_kb. factor accounts for translated
    scripts (Devanagari etc.) taking ~2-3× more bytes than Latin source."""
    if max_kb <= 0:
        return [text]
    max_bytes = max(1024, int(max_kb * 1024 / max(factor, 0.1)))
    parts, current, size = [], [], 0
    for line in text.split("\n"):
        lb = len(line.encode("utf-8")) + 1
        if size + lb > max_bytes and current:
            parts.append("\n".join(current))
            current, size = [], 0
        current.append(line)
        size += lb
    if current:
        parts.append("\n".join(current))
    return parts or [text]

_SENTENCE_RE = re.compile(r"(?<=[.!?।。！？])\s+")

def _split_long_line(line: str, limit: int) -> List[str]:
    pieces, buf = [], ""
    for sent in _SENTENCE_RE.split(line):
        if not sent:
            continue
        if len(sent) > limit:
            if buf:
                pieces.append(buf); buf = ""
            pieces.extend(sent[i:i + limit] for i in range(0, len(sent), limit))
            continue
        if len(buf) + len(sent) + 1 > limit and buf:
            pieces.append(buf); buf = sent
        else:
            buf = f"{buf} {sent}".strip()
    if buf:
        pieces.append(buf)
    return pieces

def build_chunks(text: str, limit: int = CHUNK_SIZE) -> List[List[str]]:
    """Group lines into chunks of ≤ limit chars. Each chunk is a *list of
    lines*; they are sent to the API as separate segments so paragraph
    structure survives translation exactly."""
    chunks: List[List[str]] = []
    buf: List[str] = []
    size = 0

    def flush():
        nonlocal buf, size
        if buf:
            chunks.append(buf)
            buf, size = [], 0

    for line in text.split("\n"):
        line = line.rstrip()
        if len(line) > limit:
            flush()
            for piece in _split_long_line(line, limit):
                chunks.append([piece])
            continue
        if size + len(line) + 1 > limit:
            flush()
        buf.append(line)
        size += len(line) + 1
    flush()
    return chunks


# ═══════════════════════════════════════════════════════════════════════════
# 📝 OUTPUT WRITERS
# ═══════════════════════════════════════════════════════════════════════════
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

def _xml_safe(s: str) -> str:
    return _XML_ILLEGAL.sub("", s)

def write_txt(text: str, path: str, title: str, lang: str):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)

def write_docx(text: str, path: str, title: str, lang: str):
    d = docx.Document()
    d.add_heading(_xml_safe(title), level=1)
    for line in text.split("\n"):
        line = line.strip()
        if line:
            d.add_paragraph(_xml_safe(line))
    d.save(path)

def write_epub(text: str, path: str, title: str, lang: str):
    book = epub.EpubBook()
    book.set_identifier(uuid.uuid4().hex)
    book.set_title(title)
    book.set_language(lang.split("-")[0])
    book.add_author(BOT_NAME)
    lines = [l for l in text.split("\n")]
    per_section = 400
    chapters = []
    groups = [lines[i:i + per_section] for i in range(0, len(lines), per_section)] or [[]]
    for n, group in enumerate(groups, 1):
        body = "".join(f"<p>{html.escape(_xml_safe(l.strip()))}</p>" for l in group if l.strip())
        ch = epub.EpubHtml(title=f"Section {n}", file_name=f"section_{n:03d}.xhtml", lang=lang.split("-")[0])
        ch.content = (f"<html xmlns=\"http://www.w3.org/1999/xhtml\"><head><title>{html.escape(title)}</title></head>"
                      f"<body><h2>{html.escape(title)}{' — Section ' + str(n) if len(groups) > 1 else ''}</h2>{body}</body></html>")
        book.add_item(ch)
        chapters.append(ch)
    book.toc = tuple(chapters)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav"] + chapters
    epub.write_epub(path, book)

WRITERS = {"txt": write_txt, "docx": write_docx, "epub": write_epub}


# ═══════════════════════════════════════════════════════════════════════════
# 🧩 JOB MODEL & QUEUE
# ═══════════════════════════════════════════════════════════════════════════
class JobCancelled(Exception):
    pass

@dataclass
class Job:
    job_id: str
    chat_id: int
    user_id: int
    user_name: str
    file_path: str
    ext: str
    novel_name: str
    file_size: int
    lang: str = DEFAULT_LANG
    out_format: str = DEFAULT_FORMAT
    split_kb: int = DEFAULT_SPLIT_KB
    status: str = "pending"
    created: float = field(default_factory=time.time)
    msg: Optional[Message] = None
    cancel: asyncio.Event = field(default_factory=asyncio.Event)
    progress: dict = field(default_factory=dict)      # live info for the Mini App
    gridfs_id: str = ""                                  # source file stored in MongoDB GridFS

    @property
    def lang_name(self) -> str:
        return LANGUAGES.get(self.lang, (self.lang, ""))[0]

    @property
    def out_dir(self) -> str:
        return os.path.join(STORAGE_DIR, f"{self.chat_id}_{self.job_id}")

    def cleanup(self):
        try:
            if self.file_path and os.path.exists(self.file_path):
                os.remove(self.file_path)
        except Exception:
            pass
        shutil.rmtree(self.out_dir, ignore_errors=True)
        # A wizard that never reached the queue leaves its upload in GridFS
        # (the source is uploaded *before* the options are chosen). Drop it now
        # instead of waiting for the orphan sweep. Enqueued jobs are owned by
        # the worker, which deletes the file when it finishes.
        if self.status == "pending" and self.gridfs_id and QUEUE_REPO:
            fid, self.gridfs_id = self.gridfs_id, ""
            try:
                asyncio.get_running_loop().create_task(QUEUE_REPO.delete_file(fid))
            except RuntimeError:            # no running loop (shutdown) → orphan sweep handles it
                pass


PENDING: Dict[str, Job] = {}          # wizard not finished yet
QUEUE: List[Job] = []                 # waiting for the worker
ACTIVE: Optional[Job] = None          # currently translating
QUEUE_WAKE: Optional[asyncio.Event] = None
USER_STATE: Dict[int, dict] = {}      # chat_id → {"state": ..., "job_id": ...}
SHUTTING_DOWN = False

def queue_position(job: Job) -> int:
    try:
        return QUEUE.index(job) + 1
    except ValueError:
        return 0

def user_job_count(uid: int) -> int:
    n = sum(1 for j in QUEUE if j.user_id == uid)
    if ACTIVE and ACTIVE.user_id == uid:
        n += 1
    return n

async def user_job_count_all(uid: int) -> int:
    """RAM queue + persisted jobs_queue (queued/running) for the per-user cap."""
    n = user_job_count(uid)
    if QUEUE_REPO:
        try:
            n += await QUEUE_REPO.queued_count(uid)
        except Exception as e:
            log.debug("queued_count: %s", e)
    return n

def summary_block(job: Job) -> str:
    return (
        f"📘 File: {b(job.novel_name)}\n"
        f"💾 Size: {b(fmt_size(job.file_size))}\n"
        f"🌐 Language: {b(lang_label(job.lang))}\n"
        f"📄 Format: {b(job.out_format.upper())}\n"
        f"✂️ Split: {b(split_label(job.split_kb))}\n"
    )

def queued_text(job: Job, pos: int) -> str:
    return (
        header("Queued", "⏳") + summary_block(job) + f"{DIV}\n"
        f"📍 Position: {b('#' + str(pos))}\n"
        f"{'▶️ Next up!' if pos == 1 else '🕐 Please wait…'}"
    )

def cancel_kb(job: Job) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([[InlineKeyboardButton("❌ Cancel job", callback_data=f"qx:{job.job_id}")]])

async def enqueue(job: Job, message: Message):
    job.msg = message
    job.status = "queued"
    if QUEUE_REPO and job.gridfs_id:
        await QUEUE_REPO.enqueue_job({
            "job_id": job.job_id, "file_id": job.gridfs_id, "chat_id": job.chat_id,
            "user_id": job.user_id, "user_name": job.user_name, "name": job.novel_name,
            "size": job.file_size, "ext": job.ext, "lang": job.lang, "fmt": job.out_format,
            "split_kb": job.split_kb,
        })
        pos = await QUEUE_REPO.queue_position(job.job_id) or 1
        store.audit("enqueue", job.user_id, job.user_id, job=job.job_id, name=job.novel_name[:60])
    else:
        # RAM-only mode is kept for the offline Mini App harness. Production
        # master/worker deployments require MongoDB so jobs survive restarts.
        QUEUE.append(job)
        pos = queue_position(job)
    await LiveMessage(message).update(queued_text(job, pos), cancel_kb(job), force=True)
    if QUEUE_WAKE:
        QUEUE_WAKE.set()

async def cancel_persisted(job_id: str, uid: int) -> bool:
    """Cancel a queued/running job in the shared MongoDB queue (owner may cancel anyone's)."""
    if not QUEUE_REPO or not job_id:
        return False
    try:
        ok = await QUEUE_REPO.cancel_job(job_id, uid, owner=store.is_owner(uid))
    except Exception as e:
        log.debug("cancel_job: %s", e)
        return False
    if ok:
        store.audit("cancel", uid, uid, job=job_id)
        doc = await QUEUE_REPO.get_job(job_id)
        # a *queued* job is never picked up again → drop its GridFS source right away;
        # running jobs are cleaned by the worker that notices the cancellation.
        if doc and not doc.get("worker_id"):
            await QUEUE_REPO.delete_file(str(doc.get("file_id", "")))
    return ok

async def cancel_all_persisted(uid: int) -> int:
    if not QUEUE_REPO:
        return 0
    n = 0
    try:
        for doc in await QUEUE_REPO.list_user_jobs(uid, owner=False, limit=50):
            if doc.get("status") in ("queued", "running"):
                n += int(await cancel_persisted(str(doc.get("_id")), uid))
    except Exception as e:
        log.debug("cancel_all_persisted: %s", e)
    return n

async def notify_positions():
    for i, j in enumerate(list(QUEUE), 1):
        await LiveMessage(j.msg).update(queued_text(j, i), cancel_kb(j), force=True)

async def queue_worker():
    global ACTIVE
    log.info("Queue worker started")
    while not SHUTTING_DOWN:
        if not QUEUE:
            QUEUE_WAKE.clear()
            await QUEUE_WAKE.wait()
            continue
        job = QUEUE.pop(0)
        ACTIVE = job
        job.status = "running"
        asyncio.create_task(notify_positions())
        try:
            await process_job(job)
        except asyncio.CancelledError:
            raise
        except Exception as e:
            log.exception("Worker error: %s", e)
        finally:
            ACTIVE = None
            job.cleanup()

async def access_sweep() -> None:
    """Expire finished plans, notify the user, remind people whose plan ends soon."""
    for uid in store.sweep_expired():
        await clear_user_commands(uid)
        # drop their queued / pending work — they can no longer use the bot
        for j in [j for j in QUEUE if j.user_id == uid]:
            QUEUE.remove(j); j.status = "cancelled"; j.cleanup()
        for jid, j in [(k, v) for k, v in PENDING.items() if v.user_id == uid]:
            PENDING.pop(jid, None); j.cleanup()
        await safe_send(uid, header("Access Expired", "⌛") +
                        "Your access plan has ended.\n"
                        "Tap 🙋 <b>Request Access</b> to ask for a renewal.",
                        reply_markup=locked_keyboard(ACCESS_EXPIRED))
        if store.admins():
            await notify_admins(f"⌛ Access of {user_line(uid)} expired.",
                                InlineKeyboardMarkup([[InlineKeyboardButton("➕ 1 Month", callback_data=f"ap:{uid}:1m"),
                                                       InlineKeyboardButton("➕ 1 Year", callback_data=f"ap:{uid}:1y"),
                                                       InlineKeyboardButton("🔎 Profile", callback_data=f"ui:{uid}")]]))
    for uid, exp in store.due_reminders():
        left = days_left(exp)
        when = "today" if not left else f"in {left} day{'s' if left != 1 else ''}"
        await safe_send(uid, header("Plan Ending Soon", "🔔") +
                        f"Your access expires {b(when)} ({fmt_datetime(exp)}).\n"
                        "Ask the owner for an extension if you want to keep translating.",
                        reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("🎫 My Access", callback_data="nav:access")]]))
        store.audit("reminder", 0, uid, expires=exp)

async def janitor():
    while not SHUTTING_DOWN:
        await asyncio.sleep(300)
        now = time.time()
        if store.connected:
            await store.check_budget()
        try:
            await access_sweep()
        except Exception as e:
            log.warning("access sweep failed: %s", e)
        if QUEUE_REPO and not EMBEDDED:      # the embedded worker already runs this loop
            try:
                await QUEUE_REPO.requeue_stale()
                await QUEUE_REPO.prune_finished()
                await QUEUE_REPO.prune_orphan_files()
            except Exception as e:
                log.debug("queue maintenance: %s", e)
        for jid, job in list(PENDING.items()):
            if now - job.created > PENDING_TTL:
                job.cleanup()
                PENDING.pop(jid, None)
                USER_STATE.pop(job.chat_id, None)
                await LiveMessage(job.msg).update(
                    header("Session expired", "⌛") + f"📘 {b(job.novel_name)}\nPlease send the file again.",
                    back_home_kb(), force=True)
        # orphaned temp files (e.g. after a crash mid-download)
        active_paths = {j.file_path for j in list(PENDING.values()) + QUEUE} | ({ACTIVE.file_path} if ACTIVE else set())
        for root in (INBOX_DIR, STORAGE_DIR):
            try:
                for name in os.listdir(root):
                    p = os.path.join(root, name)
                    if p in active_paths:
                        continue
                    try:
                        if now - os.path.getmtime(p) > 2 * 3600:
                            if os.path.isdir(p):
                                shutil.rmtree(p, ignore_errors=True)
                            else:
                                os.remove(p)
                    except Exception:
                        pass
            except Exception:
                pass


# ═══════════════════════════════════════════════════════════════════════════
# 🌐 TRANSLATION ENGINE  (async aiohttp → Google web endpoint, thread fallback)
# ═══════════════════════════════════════════════════════════════════════════
HTTP: Optional[aiohttp.ClientSession] = None
_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
       "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
_GT_URL = "https://translate.googleapis.com/translate_a/t"

async def get_http() -> aiohttp.ClientSession:
    global HTTP
    if HTTP is None or HTTP.closed:
        HTTP = aiohttp.ClientSession(
            headers={"User-Agent": _UA},
            timeout=aiohttp.ClientTimeout(total=REQUEST_TIMEOUT),
            connector=aiohttp.TCPConnector(limit=CONCURRENCY_LIMIT + 4, ttl_dns_cache=300),
        )
    return HTTP

class TranslateError(Exception):
    pass

async def gt_translate_lines(lines: List[str], target: str) -> List[str]:
    """Translate a list of segments in a single request. Empty segments are
    passed through untouched. Returns list of the same length."""
    idx = [i for i, l in enumerate(lines) if l.strip()]
    if not idx:
        return list(lines)
    sess = await get_http()
    data = [("q", lines[i]) for i in idx]
    async with sess.post(_GT_URL, params={"client": "gtx", "sl": "auto", "tl": target}, data=data) as resp:
        if resp.status == 429 or resp.status >= 500:
            raise TranslateError(f"HTTP {resp.status}")
        if resp.status != 200:
            raise TranslateError(f"HTTP {resp.status}")
        try:
            payload = json.loads(await resp.text())
        except Exception as e:
            raise TranslateError(f"bad json: {e}")
    if not isinstance(payload, list) or len(payload) != len(idx):
        raise TranslateError("unexpected payload shape")
    out = list(lines)
    for pos, item in zip(idx, payload):
        # sl=auto → [text, detected_lang]; fixed sl → "text"
        txt = item[0] if isinstance(item, list) else item
        if not isinstance(txt, str):
            raise TranslateError("non-string result")
        out[pos] = html.unescape(txt)
    return out

def _fallback_translate(text: str, target: str) -> str:
    if not HAS_FALLBACK:
        raise TranslateError("no fallback")
    res = _FallbackTranslator(source="auto", target=target).translate(text)
    if not res or "Error 500" in res or "Server Error" in res:
        raise TranslateError("fallback bad result")
    return res

class TranslationEngine:
    def __init__(self, target: str, concurrency: int = CONCURRENCY_LIMIT):
        self.target = target
        self.sem = asyncio.Semaphore(concurrency)
        self.total = 0
        self.done = 0
        self.failed = 0
        self.started = time.time()

    async def _translate_one(self, idx: int, lines: List[str], cancel: asyncio.Event) -> Tuple[int, List[str]]:
        if not any(l.strip() for l in lines):
            self.done += 1
            return idx, lines
        async with self.sem:
            if cancel.is_set():
                raise JobCancelled()
            delay = 1.5
            for attempt in range(MAX_RETRIES):
                if cancel.is_set():
                    raise JobCancelled()
                try:
                    res = await gt_translate_lines(lines, self.target)
                    self.done += 1
                    return idx, res
                except (TranslateError, aiohttp.ClientError, asyncio.TimeoutError, json.JSONDecodeError) as e:
                    log.debug("chunk %d attempt %d failed: %s", idx, attempt + 1, e)
                    if attempt == MAX_RETRIES - 2:            # penultimate: try the thread fallback
                        try:
                            joined = "\n".join(lines)
                            res_text = await asyncio.wait_for(
                                asyncio.to_thread(_fallback_translate, joined, self.target), REQUEST_TIMEOUT + 20)
                            self.done += 1
                            return idx, [res_text]
                        except Exception as e2:
                            log.debug("fallback failed: %s", e2)
                    await asyncio.sleep(delay + random.uniform(0, 1.0))
                    delay = min(delay * 2, 20)
            self.done += 1
            self.failed += 1
            return idx, lines                                 # keep original text

    async def run(self, chunks: List[List[str]], cancel: asyncio.Event, on_progress) -> str:
        self.total = len(chunks)
        self.started = time.time()
        results: List[Optional[List[str]]] = [None] * self.total
        tasks = [asyncio.create_task(self._translate_one(i, c, cancel)) for i, c in enumerate(chunks)]
        try:
            for fut in asyncio.as_completed(tasks):
                idx, lines = await fut
                results[idx] = lines
                await on_progress(self)
        except BaseException:
            for t in tasks:
                t.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)
            raise
        return "\n".join("\n".join(results[i] if results[i] is not None else chunks[i]) for i in range(self.total))

    @property
    def speed(self) -> float:
        return self.done / max(time.time() - self.started, 0.001)

    @property
    def eta(self) -> float:
        return (self.total - self.done) / self.speed if self.speed > 0 and self.done else 0.0

def dashboard_text(job: Job, part: int, total_parts: int, eng: TranslationEngine, phase: str) -> str:
    ratio = eng.done / eng.total if eng.total else 0
    return (
        header("Live Dashboard", "🖥") +
        f"📘 {b(job.novel_name)}\n"
        f"🌐 {lang_label(job.lang)}  ·  📄 {job.out_format.upper()}\n"
        f"📦 Part {b(f'{part} / {total_parts}')}\n"
        f"{DIV}\n"
        f"{progress_bar(ratio)} {b(f'{ratio * 100:.1f}%')}\n"
        f"🧩 Chunks: {eng.done} / {eng.total}\n"
        f"⚡ Speed: {eng.speed:.2f} chunk/s\n"
        f"⏱ Elapsed: {fmt_time(time.time() - eng.started)}  ·  ETA: {fmt_time(eng.eta)}\n"
        + (f"⚠️ Retries exhausted: {eng.failed}\n" if eng.failed else "") +
        f"{DIV}\n"
        f"<i>{esc(phase)}</i>"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 🚀 MAIN PIPELINE
# ═══════════════════════════════════════════════════════════════════════════
async def create_backup_topic(title: str) -> Optional[int]:
    if not BACKUP_GROUP_ID:
        return None
    try:
        peer = await app.resolve_peer(BACKUP_GROUP_ID)
        result = await app.invoke(functions.channels.CreateForumTopic(
            channel=peer, title=title[:128], random_id=random.randint(1, 2 ** 62)))
        for upd in result.updates:
            msg = getattr(upd, "message", None)
            if msg is not None and hasattr(msg, "id"):
                return msg.id
    except Exception as e:
        log.debug("create_backup_topic failed (forum topics disabled?): %s", e)
    return None

def safe_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name).strip(" .")
    return name[:90] or "Document"

async def process_job(job: Job):
    live = LiveMessage(job.msg)
    out_dir = job.out_dir
    os.makedirs(out_dir, exist_ok=True)
    started = time.time()
    total_chars = 0
    total_parts = 0
    failed_chunks = 0

    try:
        await live.update(header("Processing", "🔄") + summary_block(job) + f"{DIV}\n📖 Extracting text…",
                          cancel_kb(job), force=True)
        if not os.path.exists(job.file_path):
            raise FileNotFoundError("source file vanished (server restarted?)")
        extractor = EXTRACTORS.get(job.ext, extract_text_txt)
        text = normalise_text(await asyncio.to_thread(extractor, job.file_path))
        if len(text) < 20:
            job.status = "failed"
            store.bump(job.user_id, 0, 0, failed=True)
            store.add_history(job.user_id, {"job_id": job.job_id, "name": job.novel_name, "lang": job.lang,
                                            "fmt": job.out_format, "size": job.file_size, "status": "failed",
                                            "error": "NoText", "user": job.user_name})
            await live.update(header("Failed", "❌") + summary_block(job) + f"{DIV}\n"
                              "No readable text found.\n<i>Scanned PDFs / image-only files are not supported.</i>",
                              back_home_kb(), force=True)
            return
        if job.cancel.is_set():
            raise JobCancelled()

        parts = split_text_by_size(text, job.split_kb, EXPANSION.get(job.lang, 1.15))
        total_parts = len(parts)
        total_chars = len(text)

        topic_id = await create_backup_topic(job.novel_name)
        backup_kwargs = {"reply_to_message_id": topic_id} if topic_id else {}
        if BACKUP_GROUP_ID:
            await safe_send(BACKUP_GROUP_ID,
                            header("New Job", "📥") + summary_block(job) +
                            f"👤 User: {b(job.user_name)} ({code(job.user_id)})\n"
                            f"🧩 Parts: {b(total_parts)}  ·  🔤 Chars: {b(fmt_int(total_chars))}",
                            **backup_kwargs)

        job.progress.update({"parts": total_parts, "chars": total_chars, "phase": "translating"})
        for i, part_text in enumerate(parts, 1):
            if job.cancel.is_set():
                raise JobCancelled()
            chunks = build_chunks(part_text)
            eng = TranslationEngine(job.lang)
            job.progress.update({"part": i, "engine": eng})

            async def on_progress(e: TranslationEngine, _i=i):
                if job.cancel.is_set():
                    raise JobCancelled()
                await live.update(dashboard_text(job, _i, total_parts, e, "Translating…"), cancel_kb(job))

            translated = await eng.run(chunks, job.cancel, on_progress)
            failed_chunks += eng.failed
            await live.update(dashboard_text(job, i, total_parts, eng, f"Building .{job.out_format} file…"),
                              cancel_kb(job), force=True)

            part_title = f"{job.novel_name} — Part {i} of {total_parts}" if total_parts > 1 else job.novel_name
            base = f"Part {i} of {total_parts} | {job.novel_name} [{job.lang_name}]" if total_parts > 1 \
                else f"{job.novel_name} [{job.lang_name}]"
            out_path = os.path.join(out_dir, f"{safe_filename(base)}.{job.out_format}")
            await asyncio.to_thread(WRITERS[job.out_format], translated, out_path, part_title, job.lang)

            caption = (f"📘 {b(job.novel_name)}\n"
                       f"📦 Part {i} / {total_parts}  ·  🌐 {lang_label(job.lang)}\n"
                       f"💾 {fmt_size(os.path.getsize(out_path))}")
            sent = await safe_send_document(job.chat_id, out_path, caption)
            if sent is None:
                log.error("Could not deliver part %d of job %s", i, job.job_id)
                await safe_send(job.chat_id, f"⚠️ Part {i} could not be delivered. Please retry later.")
            if BACKUP_GROUP_ID:
                await safe_send_document(BACKUP_GROUP_ID, out_path,
                                         caption + f"\n👤 {b(job.user_name)} ({code(job.user_id)})", **backup_kwargs)
            try:
                os.remove(out_path)          # free disk as we go
            except Exception:
                pass

        job.status = "done"
        store.bump(job.user_id, total_parts, total_chars)
        elapsed = time.time() - started
        store.add_history(job.user_id, {"job_id": job.job_id, "name": job.novel_name, "lang": job.lang,
                                        "fmt": job.out_format, "split": job.split_kb, "size": job.file_size,
                                        "parts": total_parts, "chars": total_chars, "secs": int(elapsed),
                                        "status": "done", "user": job.user_name})
        await live.update(
            header("Completed", "✅") + summary_block(job) + f"{DIV}\n"
            f"🧩 Parts: {b(total_parts)}\n"
            f"🔤 Characters: {b(fmt_int(total_chars))}\n"
            f"⏱ Time: {b(fmt_time(elapsed))}\n"
            + (f"⚠️ {failed_chunks} chunk(s) kept original text\n" if failed_chunks else "") +
            f"{DIV}\n🚀 Thank you for using {esc(BOT_NAME)}!",
            back_home_kb(), force=True)

    except JobCancelled:
        job.status = "cancelled"
        store.bump(job.user_id, 0, 0, cancelled=True)
        store.add_history(job.user_id, {"job_id": job.job_id, "name": job.novel_name, "lang": job.lang,
                                        "fmt": job.out_format, "size": job.file_size, "status": "cancelled",
                                        "secs": int(time.time() - started), "user": job.user_name})
        reason = "Server is restarting (new deploy). Please resend the file in a minute." if SHUTTING_DOWN \
            else "Job stopped by user."
        await live.update(header("Cancelled", "🛑") + summary_block(job) + f"{DIV}\n{reason}", back_home_kb(), force=True)
        if BACKUP_GROUP_ID and not SHUTTING_DOWN:
            await safe_send(BACKUP_GROUP_ID, f"🛑 Job cancelled: {b(job.novel_name)} by {code(job.user_id)}")
    except asyncio.CancelledError:
        raise
    except Exception as e:
        job.status = "failed"
        log.exception("Job %s failed: %s", job.job_id, e)
        store.bump(job.user_id, 0, 0, failed=True)
        store.add_history(job.user_id, {"job_id": job.job_id, "name": job.novel_name, "lang": job.lang,
                                        "fmt": job.out_format, "size": job.file_size, "status": "failed",
                                        "error": type(e).__name__, "secs": int(time.time() - started),
                                        "user": job.user_name})
        await live.update(header("Error", "⚠️") + summary_block(job) + f"{DIV}\n"
                          f"Something went wrong: {code(type(e).__name__)}\nPlease try again.",
                          back_home_kb(), force=True)
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)


# ═══════════════════════════════════════════════════════════════════════════
# 🤖 CLIENT & FILTERS
# ═══════════════════════════════════════════════════════════════════════════
app = Client(
    "novel_translator_pro",
    api_id=API_ID, api_hash=API_HASH, bot_token=BOT_TOKEN,
    in_memory=True,                    # ← no .session file on disk
    workdir=BASE_DIR,
    parse_mode=enums.ParseMode.HTML,
    sleep_threshold=30,                # auto-sleep small FloodWaits
    max_concurrent_transmissions=2,
)

async def _auth_msg(_, __, m: Message):
    return bool(m.from_user and store.is_authorized(m.from_user.id))

async def _auth_cb(_, __, q: CallbackQuery):
    return bool(q.from_user and store.is_authorized(q.from_user.id))

async def _owner_msg(_, __, m: Message):
    return bool(m.from_user and store.is_owner(m.from_user.id))

async def _admin_msg(_, __, m: Message):
    return bool(m.from_user and store.is_admin(m.from_user.id))

authorized = filters.create(_auth_msg)
authorized_cb = filters.create(_auth_cb)
owner_only = filters.create(_owner_msg)
admin_only = filters.create(_admin_msg)
PRIVATE = filters.private & filters.incoming
OUTPUT_NAME_RE = re.compile(r"^Part\s*\d+\s*(of\s*\d+)?\s*\|", re.IGNORECASE)

def touch_user(m: Message) -> None:
    """Make sure a profile exists / the display name is fresh so prefs & stats work."""
    if m.from_user:
        store.ensure_user(m.from_user.id, m.from_user.first_name or "User", m.from_user.username or "")


# ═══════════════════════════════════════════════════════════════════════════
# 🎫 ACCESS REQUESTS  —  texts, keyboards & admin notifications
# ═══════════════════════════════════════════════════════════════════════════
def user_link(uid: int, name: str = "") -> str:
    u = store.user(uid) or {}
    name = name or u.get("name") or "User"
    return f'<a href="tg://user?id={uid}">{esc(name)}</a>'

def user_line(uid: int) -> str:
    """`Name (@username) · 123456` for admin lists."""
    u = store.user(uid) or {}
    un = f" (@{esc(u['username'])})" if u.get("username") else ""
    return f"{user_link(uid)}{un} · {code(uid)}"

def text_access(uid: int) -> str:
    """The user's own access card (/access)."""
    a = store.access(uid)
    status = store.status(uid)
    label, icon = ACCESS_LABELS.get(status, ("Unknown", "❓"))
    lines = [header("My Access", "🎫"),
             f"{icon} Status: {b(label)}",
             f"🎖 Role: {b(store.role(uid).title())}"]
    if status == ACCESS_APPROVED:
        if store.is_owner(uid) or not a.get("expires"):
            lines.append("⏳ Valid: " + b("♾ Lifetime"))
        else:
            lines.append(f"📦 Plan: {b(plan_label(a.get('plan', '')))}")
            lines.append(f"⏳ Valid: {b(expiry_label(a['expires']))}")
            lines.append(f"📅 Expires: {b(fmt_datetime(a['expires']))}")
        if a.get("approved_at") and not store.is_owner(uid):
            lines.append(f"✅ Approved: {b(fmt_date(a['approved_at']))}")
    elif status == ACCESS_PENDING:
        lines.append(f"📨 Requested: {b(fmt_datetime(a.get('requested_at', 0)))}")
        lines.append("<i>An admin will review your request soon.\nYou will get a message here once it is decided.</i>")
    elif status == ACCESS_EXPIRED:
        lines.append(f"📅 Expired: {b(fmt_date(a.get('expires', 0)))}")
        lines.append("<i>Tap 🙋 Request Access to ask for a renewal.</i>")
    elif status == ACCESS_REJECTED:
        if a.get("reason"):
            lines.append(f"💬 Reason: {esc(a['reason'])}")
        wait = a.get("updated", 0) + REJECT_COOLDOWN_H * 3600 - time.time()
        if REJECT_COOLDOWN_H and wait > 0:
            lines.append(f"<i>You can send a new request in {fmt_time(wait)}.</i>")
        else:
            lines.append("<i>You may send a new request now.</i>")
    elif status == ACCESS_BANNED:
        if a.get("reason"):
            lines.append(f"💬 Reason: {esc(a['reason'])}")
        lines.append("<i>Access to this bot has been blocked.</i>")
    else:
        lines.append("<i>This bot is private. Tap 🙋 Request Access\nand the owner will review your request.</i>")
    lines.append(f"{DIV}\n<i>Your ID: {code(uid)}</i>")
    return "\n".join(lines)

def request_kb(status: str) -> Optional[InlineKeyboardMarkup]:
    if status in (ACCESS_NONE, ACCESS_EXPIRED, ACCESS_REJECTED):
        return InlineKeyboardMarkup([[InlineKeyboardButton("🙋 Request Access", callback_data="req:send")]])
    if status == ACCESS_PENDING:
        return InlineKeyboardMarkup([[InlineKeyboardButton("🔄 Check status", callback_data="req:status"),
                                      InlineKeyboardButton("↩️ Withdraw", callback_data="req:cancel")]])
    return None

def approve_kb(uid: int, compact: bool = False) -> InlineKeyboardMarkup:
    """Quick-approve buttons shown to admins next to a request / user card."""
    rows = [[InlineKeyboardButton("1 Week", callback_data=f"ap:{uid}:1w"),
             InlineKeyboardButton("1 Month", callback_data=f"ap:{uid}:1m"),
             InlineKeyboardButton("3 Months", callback_data=f"ap:{uid}:3m")],
            [InlineKeyboardButton("6 Months", callback_data=f"ap:{uid}:6m"),
             InlineKeyboardButton("1 Year", callback_data=f"ap:{uid}:1y"),
             InlineKeyboardButton("♾ Lifetime", callback_data=f"ap:{uid}:forever")],
            [InlineKeyboardButton("✏️ Custom", callback_data=f"apc:{uid}"),
             InlineKeyboardButton("❌ Reject", callback_data=f"rj:{uid}"),
             InlineKeyboardButton("🚫 Ban", callback_data=f"bn:{uid}")]]
    if not compact:
        rows.append([InlineKeyboardButton("🔎 Profile", callback_data=f"ui:{uid}"),
                     InlineKeyboardButton("⏳ All requests", callback_data="nav:pending")])
    return InlineKeyboardMarkup(rows)

def user_card_kb(uid: int) -> InlineKeyboardMarkup:
    """Management buttons for an existing user (from /userinfo, /users)."""
    status = store.status(uid)
    rows = []
    if status == ACCESS_BANNED:
        rows.append([InlineKeyboardButton("♻️ Unban", callback_data=f"ub:{uid}")])
    else:
        if status == ACCESS_APPROVED:
            rows.append([InlineKeyboardButton("+1 Week", callback_data=f"ex:{uid}:1w"),
                         InlineKeyboardButton("+1 Month", callback_data=f"ex:{uid}:1m"),
                         InlineKeyboardButton("+1 Year", callback_data=f"ex:{uid}:1y")])
            rows.append([InlineKeyboardButton("♾ Lifetime", callback_data=f"ap:{uid}:forever"),
                         InlineKeyboardButton("✏️ Custom", callback_data=f"apc:{uid}"),
                         InlineKeyboardButton("🔒 Revoke", callback_data=f"rv:{uid}")])
        else:
            rows.append([InlineKeyboardButton("✅ 1 Month", callback_data=f"ap:{uid}:1m"),
                         InlineKeyboardButton("✅ 1 Year", callback_data=f"ap:{uid}:1y"),
                         InlineKeyboardButton("♾ Lifetime", callback_data=f"ap:{uid}:forever")])
            rows.append([InlineKeyboardButton("✏️ Custom", callback_data=f"apc:{uid}"),
                         InlineKeyboardButton("❌ Reject", callback_data=f"rj:{uid}")])
        rows.append([InlineKeyboardButton("🚫 Ban", callback_data=f"bn:{uid}")])
    rows.append([InlineKeyboardButton("🔄 Refresh", callback_data=f"ui:{uid}"),
                 InlineKeyboardButton("👥 Users", callback_data="nav:users")])
    return InlineKeyboardMarkup(rows)

def text_user_card(uid: int) -> str:
    u = store.user(uid)
    if not u:
        return header("User", "🔎") + f"{code(uid)} — <i>never talked to the bot.</i>"
    a = store.access(uid)
    status = store.status(uid)
    label, icon = ACCESS_LABELS.get(status, ("Unknown", "❓"))
    s = u.get("stats", {})
    lines = [header("User Profile", "🔎"),
             f"👤 {user_line(uid)}",
             f"🎖 Role: {b(store.role(uid).title())}",
             f"{icon} Access: {b(label)}"]
    if status == ACCESS_APPROVED:
        lines.append("⏳ Valid: " + b("♾ Lifetime" if store.is_owner(uid) or not a.get("expires")
                                       else expiry_label(a["expires"])))
        if a.get("plan"):
            lines.append(f"📦 Plan: {b(plan_label(a['plan']))}")
        if a.get("approved_by"):
            lines.append(f"✅ By: {user_link(a['approved_by'])} · {fmt_date(a.get('approved_at', 0))}")
    if a.get("note"):
        lines.append(f"📝 Note: <i>{esc(a['note'])}</i>")
    if a.get("reason") and status in (ACCESS_REJECTED, ACCESS_BANNED, ACCESS_NONE):
        lines.append(f"💬 Reason: <i>{esc(a['reason'])}</i>")
    lines += [f"📨 Requests: {b(a.get('requests', 0))}",
              f"{DIV}",
              f"📚 Jobs: {b(s.get('jobs', 0))} · 🧩 Parts: {b(s.get('parts', 0))} · 🔤 {b(fmt_int(s.get('chars', 0)))}",
              f"📅 Joined: {b(fmt_date(u.get('joined', 0)))} · 👀 Seen: {b(fmt_date(u.get('last_seen', 0)))}"]
    hist = a.get("history") or []
    if hist:
        lines.append(f"{DIV}\n{b('Recent')}")
        for h in hist[-5:][::-1]:
            extra = ""
            if h.get("expires"):
                extra = f" → {fmt_date(h['expires'])}"
            elif h.get("action") in ("approve", "extend"):
                extra = " → ♾"
            lines.append(f"• {fmt_date(h.get('ts', 0))} · {esc(h.get('action', ''))}{extra}")
    return "\n".join(lines)

async def notify_admins(text: str, kb: Optional[InlineKeyboardMarkup] = None, exclude: int = 0) -> None:
    for aid in store.admins():
        if aid and aid != exclude:
            await safe_send(aid, text, reply_markup=kb)

async def send_access_request(uid: int, name: str, username: str = "", note: str = "") -> Tuple[bool, str]:
    """Create the request and ping every admin with quick-approve buttons."""
    ok, why = store.request_access(uid, name, username, note)
    if not ok:
        return ok, why
    a = store.access(uid)
    text = (header("Access Request", "🙋") +
            f"👤 {user_line(uid)}\n"
            f"📨 Request #{a.get('requests', 1)} · {fmt_datetime(a.get('requested_at', 0))}\n"
            + (f"📝 <i>{esc(note)}</i>\n" if note else "")
            + f"⏳ Waiting: {b(len(store.pending_users()))}\n"
            f"{DIV}\nChoose a plan to approve:")
    await notify_admins(text, approve_kb(uid))
    return True, "ok"

async def grant_and_notify(uid: int, spec: str, by: int, extend: bool = False) -> Optional[dict]:
    """Approve/extend + tell the user + refresh their command menu."""
    a = store.approve(uid, spec, by, extend=extend)
    if a is None:
        return None
    await set_user_commands(uid, force=True)
    what = "Access Extended" if extend else "Access Granted"
    await safe_send(uid, header(what, "✅") +
                    f"🎉 Welcome{'' if extend else ' aboard'}!\n"
                    f"📦 Plan: {b(plan_label(a.get('plan', spec)))}\n"
                    f"⏳ Valid: {b('♾ Lifetime' if not a.get('expires') else expiry_label(a['expires']))}\n\n"
                    "Send /start to open the menu or just drop a document.",
                    reply_markup=reply_keyboard(uid))
    return a

async def _access_denied_reply(m: Message) -> None:
    """Shown to anyone who is not approved — with the right call-to-action."""
    uid = m.from_user.id
    status = store.status(uid)
    if status == ACCESS_BANNED:
        return  # stay silent for banned users (avoid spam loops)
    await m.reply(text_access(uid), reply_markup=locked_keyboard(status))
    kb = request_kb(status)
    if kb:
        hint = {"pending": "⏳ Your request is in the queue.",
                "expired": "⌛ Your plan has ended — request a renewal.",
                "rejected": "❌ Your last request was declined."}.get(status, "🔒 This bot is private.")
        await m.reply(hint, reply_markup=kb)


# ═══════════════════════════════════════════════════════════════════════════
# 🔒 NOT APPROVED  (visitors · pending · expired · rejected · banned)
# ═══════════════════════════════════════════════════════════════════════════
async def do_request(m: Message, note: str = "") -> None:
    uid = m.from_user.id
    ok, why = await send_access_request(uid, m.from_user.first_name or "User", m.from_user.username or "", note)
    if ok:
        await m.reply(header("Request Sent", "📨") +
                      "Your access request has been sent to the owner.\n"
                      "You will be notified here as soon as it is approved.\n\n"
                      f"<i>Your ID: {code(uid)}</i>",
                      reply_markup=locked_keyboard(ACCESS_PENDING))
        await m.reply("⏳ Status: <b>Pending approval</b>", reply_markup=request_kb(ACCESS_PENDING))
    elif why == "pending":
        await m.reply("⏳ Your request is already pending — please wait for an admin.",
                      reply_markup=request_kb(ACCESS_PENDING))
    elif why == "approved":
        await m.reply("✅ You already have access! Send /start.", reply_markup=reply_keyboard(uid))
    elif why == "cooldown":
        a = store.access(uid)
        wait = a.get("updated", 0) + REJECT_COOLDOWN_H * 3600 - time.time()
        await m.reply(f"⏱ Your last request was declined. You can try again in {b(fmt_time(wait))}.")
    # banned → silent

@app.on_message(filters.command(["request", "start"]) & PRIVATE & ~authorized)
async def locked_start(_, m: Message):
    if not m.from_user:
        return
    touch_user(m)
    uid = m.from_user.id
    if store.is_banned(uid):
        return
    if m.command[0].lower() == "request":
        note = m.text.split(None, 1)[1].strip() if len(m.command) > 1 else ""
        return await do_request(m, note)
    status = store.status(uid)
    await m.reply(header(BOT_NAME) +
                  f"👋 Hello, {b(m.from_user.first_name or 'there')}!\n\n"
                  "I translate <b>EPUB / TXT / DOCX</b> books into your language.\n"
                  "This bot is <b>private</b> — access is granted by the owner.\n\n"
                  + {"pending": "⏳ Your request is <b>pending</b>. Hang tight!",
                     "expired": "⌛ Your plan has <b>expired</b>. Request a renewal below.",
                     "rejected": "❌ Your last request was <b>declined</b>."}.get(
                        status, "Tap 🙋 <b>Request Access</b> and I will notify the owner.")
                  + f"\n\n<i>Your ID: {code(uid)}</i>",
                  reply_markup=locked_keyboard(status))
    kb = request_kb(status)
    if kb:
        await m.reply("👇", reply_markup=kb)

@app.on_message(filters.command("access") & PRIVATE)
async def cmd_access(_, m: Message):
    if not m.from_user:
        return
    touch_user(m)
    uid = m.from_user.id
    if store.is_banned(uid) and not store.is_authorized(uid):
        return
    kb = request_kb(store.status(uid)) if not store.is_authorized(uid) else back_home_kb()
    await m.reply(text_access(uid), reply_markup=kb)

@app.on_message(PRIVATE & ~authorized)
async def unauthorized_message(_, m: Message):
    if not m.from_user:
        return
    touch_user(m)
    uid = m.from_user.id
    if store.is_banned(uid):
        return
    txt = (m.text or "").strip()
    if txt == BTN_MYID or txt.startswith("/id"):
        return await m.reply(f"🆔 Your Telegram ID: {code(uid)}")
    if txt == BTN_REQUEST:
        return await do_request(m)
    if txt == BTN_MYACCESS:
        return await cmd_access(_, m)
    # pending users may attach a short note to their request by simply typing
    a = store.access(uid)
    if a["status"] == ACCESS_PENDING and txt and not txt.startswith("/") and len(txt) <= 200:
        a["note"] = txt
        store._save_user(uid)
        return await m.reply("📝 Noted — your message was attached to the request.")
    await _access_denied_reply(m)

@app.on_callback_query(filters.regex(r"^req:") & ~authorized_cb)
async def request_callbacks(_, q: CallbackQuery):
    uid = q.from_user.id
    store.ensure_user(uid, q.from_user.first_name or "User", q.from_user.username or "")
    if store.is_banned(uid):
        return await _answer(q, "🚫 Access blocked.", alert=True)
    action = (q.data or "").split(":", 1)[1]
    if action == "send":
        ok, why = await send_access_request(uid, q.from_user.first_name or "User", q.from_user.username or "")
        if ok:
            await _edit(q, header("Request Sent", "📨") + "The owner has been notified.\n"
                        "You will receive a message here once it is approved.", request_kb(ACCESS_PENDING))
            try:
                await q.message.reply("⏳ Waiting for approval…", reply_markup=locked_keyboard(ACCESS_PENDING))
            except Exception:
                pass
            return await _answer(q, "📨 Request sent")
        if why == "pending":
            await _edit(q, text_access(uid), request_kb(ACCESS_PENDING))
            return await _answer(q, "Already pending")
        if why == "cooldown":
            a = store.access(uid)
            wait = a.get("updated", 0) + REJECT_COOLDOWN_H * 3600 - time.time()
            return await _answer(q, f"⏱ Try again in {fmt_time(wait)}", alert=True)
        return await _answer(q)
    if action == "status":
        await _edit(q, text_access(uid), request_kb(store.status(uid)))
        return await _answer(q, "🔄 Updated")
    if action == "cancel":
        a = store.access(uid)
        if a["status"] == ACCESS_PENDING:
            a.update({"status": ACCESS_NONE, "updated": int(time.time())})
            store._save_user(uid)
            store.audit("withdraw", uid, uid)
        await _edit(q, text_access(uid), request_kb(store.status(uid)))
        return await _answer(q, "Request withdrawn")
    await _answer(q)

@app.on_callback_query(~authorized_cb)
async def unauthorized_callback(_, q: CallbackQuery):
    await q.answer("🔒 Not approved yet. Use 🙋 Request Access first.", show_alert=True)


# ═══════════════════════════════════════════════════════════════════════════
# 💬 COMMANDS
# ═══════════════════════════════════════════════════════════════════════════
@app.on_message(filters.command("start") & PRIVATE & authorized)
async def cmd_start(_, m: Message):
    touch_user(m)
    USER_STATE.pop(m.chat.id, None)
    uid = m.from_user.id
    await set_user_commands(uid)
    store.touch(uid)
    # 1) persistent bottom reply keyboard (main page, role based)  2) inline home menu
    KB_PAGE[m.chat.id] = "main"
    await m.reply(text_home(uid, m.from_user.first_name or "there"),
                  reply_markup=reply_keyboard(uid, "main"))
    await m.reply("⌨️ " + b("Menu ready") + "\n<i>" +
                  ("👑 Owner mode — the Admin page is enabled."
                   if store.is_owner(uid) else "Use the buttons below — each menu opens a sub-menu.") +
                  "</i>", reply_markup=home_keyboard(uid))

@app.on_message(filters.command("help") & PRIVATE & authorized)
async def cmd_help(_, m: Message):
    await m.reply(text_help(), reply_markup=back_home_kb())

@app.on_message(filters.command("settings") & PRIVATE & authorized)
async def cmd_settings(_, m: Message):
    touch_user(m)
    await m.reply(text_settings(m.from_user.id), reply_markup=settings_keyboard())

@app.on_message(filters.command(["setlang", "setformat", "setsplit"]) & PRIVATE & authorized)
async def cmd_set_pref(_, m: Message):
    """Direct jump into one settings sub-page (used by the ⚙️ Settings keyboard page)."""
    touch_user(m)
    uid = m.from_user.id
    lang, fmt, split = user_prefs(uid)
    what = m.command[0].lower()
    if what == "setlang":
        await m.reply(header("Default language", "🌐") + f"Current: {b(lang_label(lang))}\nPick a new default:",
                      reply_markup=language_keyboard("sl:", lang, "nav:settings"))
    elif what == "setformat":
        await m.reply(header("Default format", "📄") + f"Current: {b(fmt.upper())}\nPick a new default:",
                      reply_markup=format_keyboard("sf:", fmt, "nav:settings"))
    else:
        await m.reply(header("Default split size", "✂️") + f"Current: {b(split_label(split))}\nPick a new default:",
                      reply_markup=split_keyboard("ss:", split, "nav:settings", None))

@app.on_message(filters.command("app") & PRIVATE & authorized)
async def cmd_app(_, m: Message):
    touch_user(m)
    if not MINI_APP_URL.startswith("https://"):
        return await m.reply(header("Mini App", "📱") +
                             "The Mini App is not configured yet.\n"
                             "<i>Set <code>PUBLIC_URL</code> (or <code>RENDER_EXTERNAL_URL</code>) to a public "
                             "HTTPS address and redeploy.</i>")
    await m.reply(header("Mini App", "📱") +
                  "Your dashboard: settings, live progress,\nhistory and stats — all in one place.\n\n"
                  "<i>Tap the button below to open it.</i>",
                  reply_markup=InlineKeyboardMarkup([[app_inline_button()]]))

def text_queue(uid: int) -> str:
    lines = [header("Queue Status", "📋")]
    if ACTIVE:
        who = "you" if ACTIVE.user_id == uid else esc(ACTIVE.user_name)
        lines.append(f"▶️ Running: {b(ACTIVE.novel_name)}\n   <i>by {who}</i>\n")
    else:
        lines.append("▶️ Running: <i>nothing</i>\n")
    if QUEUE:
        lines.append(f"⏳ Waiting: {b(len(QUEUE))}")
        for i, j in enumerate(QUEUE[:10], 1):
            mark = "🟢" if j.user_id == uid else "⚪️"
            lines.append(f"{mark} #{i} {esc(j.novel_name[:32])}")
        if len(QUEUE) > 10:
            lines.append(f"… and {len(QUEUE) - 10} more")
    else:
        lines.append("⏳ Waiting: <i>empty</i>")
    return "\n".join(lines)

@app.on_message(filters.command("queue") & PRIVATE & authorized)
async def cmd_queue(_, m: Message):
    await m.reply(text_queue(m.from_user.id), reply_markup=back_home_kb())

@app.on_message(filters.command("cancel") & PRIVATE & authorized)
async def cmd_cancel(_, m: Message):
    uid = m.from_user.id
    n = await cancel_all_persisted(uid)
    if ACTIVE and ACTIVE.user_id == uid:
        ACTIVE.cancel.set(); n += 1
    for j in [j for j in QUEUE if j.user_id == uid]:
        QUEUE.remove(j); j.status = "cancelled"; j.cleanup(); n += 1
        await LiveMessage(j.msg).update(header("Cancelled", "🛑") + summary_block(j), back_home_kb(), force=True)
    for jid, j in [(k, v) for k, v in PENDING.items() if v.user_id == uid]:
        PENDING.pop(jid, None); j.cleanup(); n += 1
    USER_STATE.pop(m.chat.id, None)
    if n:
        await notify_positions()
        await m.reply(f"🛑 Cancelled {b(n)} job(s).")
    else:
        await m.reply("ℹ️ You have no active or queued jobs.")

def text_mystats(uid: int) -> str:
    u = store.user(uid) or {}
    s = u.get("stats", {})
    since = time.strftime("%d %b %Y", time.localtime(u.get("joined", time.time())))
    return (
        header("My Statistics", "📊") +
        f"👤 {b(u.get('name', 'User'))}  ·  {code(uid)}\n"
        f"🎖 Role: {b(store.role(uid).title())}\n"
        f"🎫 Access: {b(access_line(uid))}\n"
        f"📅 Joined: {b(since)}\n"
        f"{DIV}\n"
        f"📚 Files translated: {b(s.get('jobs', 0))}\n"
        f"🧩 Parts delivered: {b(s.get('parts', 0))}\n"
        f"🔤 Characters: {b(fmt_int(s.get('chars', 0)))}\n"
        + ("" if store.connected else f"{DIV}\n<i>Stats are in-memory and reset on redeploy.</i>")
    )

@app.on_message(filters.command("mystats") & PRIVATE & authorized)
async def cmd_mystats(_, m: Message):
    touch_user(m)
    await m.reply(text_mystats(m.from_user.id), reply_markup=back_home_kb())

@app.on_message(filters.command("id") & PRIVATE)
async def cmd_id(_, m: Message):
    await m.reply(f"🆔 Your Telegram ID: {code(m.from_user.id)}")

# ── Owner / admin commands ──────────────────────────────────────────────────
def text_owner() -> str:
    s = store.stats
    up = fmt_time(time.time() - store.booted)
    c = store.count_by_status()
    return (
        header("Owner Panel", "👑") +
        f"⏱ Uptime: {b(up)}\n"
        f"👥 Users: {b(len(store.users))}  ·  ✅ Approved: {b(c.get(ACCESS_APPROVED, 0))}  ·  ⏳ Pending: {b(c.get(ACCESS_PENDING, 0))}\n"
        f"⌛ Expired: {b(c.get(ACCESS_EXPIRED, 0))}  ·  ❌ Rejected: {b(c.get(ACCESS_REJECTED, 0))}  ·  🚫 Banned: {b(c.get(ACCESS_BANNED, 0))}\n"
        f"🛡 Admins: {b(len(store.admins()) - (1 if store.owner_id else 0))}  ·  🔓 Public: {b('yes' if PUBLIC_MODE else 'no')}\n"
        f"💬 Tracked chats: {b(len(store.chats))}\n"
        f"📚 Jobs done: {b(s['jobs'])}  ·  ❌ Failed: {b(s['failed'])}  ·  🛑 Cancelled: {b(s['cancelled'])}\n"
        f"🧩 Parts: {b(s['parts'])}  ·  🔤 Chars: {b(fmt_int(s['chars']))}\n"
        f"▶️ Active: {b(ACTIVE.novel_name if ACTIVE else '—')}\n"
        f"⏳ Queue: {b(len(QUEUE))}  ·  📝 Pending wizards: {b(len(PENDING))}\n"
        f"🗄 Backup group: {b(BACKUP_GROUP_ID or 'disabled')}\n"
        f"{DIV}\n"
        f"{b('Access')}\n"
        "/pending – requests waiting\n"
        "/approve &lt;id&gt; [1w|1m|1y|forever|45d|2026-12-31]\n"
        "/extend &lt;id&gt; &lt;duration&gt; · /revoke &lt;id&gt;\n"
        "/reject &lt;id&gt; [reason] · /ban &lt;id&gt; [reason] · /unban &lt;id&gt;\n"
        "/users · /userinfo &lt;id&gt; · /audit\n"
        f"{b('Owner')}\n"
        "/admins · /addadmin &lt;id&gt; · /deladmin &lt;id&gt;\n"
        "/broadcast &lt;text&gt; (or reply) · /links\n"
        f"{DIV}\n"
        + (f"<i>🗄 MongoDB connected ({esc(MONGO_DB)}) — users, access\nand history are persistent.</i>\n"
           f"<i>💾 Storage: {store.db_size_mb:.1f} / {DB_BUDGET_MB} MB ({store.budget_info()['percent']}%)"
           f" · {fmt_int(store.db_job_docs)} job docs · TTL {DB_JOB_TTL_DAYS} d</i>"
           if store.connected else
           "<i>⚠️ No database: approvals given at runtime are lost on\n"
           "restart. Put permanent IDs in AUTHORIZED_USERS / ADMIN_USERS env.</i>")
    )

@app.on_message(filters.command("stats") & PRIVATE & owner_only)
async def cmd_stats(_, m: Message):
    await m.reply(text_owner(), reply_markup=back_home_kb())

def _parse_target(m: Message, _usage: str = "") -> Optional[int]:
    """`/cmd <id>` or `/cmd` as a reply to a forwarded message → user id."""
    if len(m.command) > 1 and m.command[1].lstrip("-").isdigit():
        return int(m.command[1])
    if m.reply_to_message and m.reply_to_message.forward_from:
        return m.reply_to_message.forward_from.id
    return None

def text_pending() -> str:
    rows = store.pending_users()
    lines = [header(f"Access Requests ({len(rows)})", "⏳")]
    if not rows:
        lines.append("<i>No pending requests. 🎉</i>")
    for uid, u in rows[:25]:
        a = u["access"]
        note = f"\n   📝 <i>{esc(a['note'][:80])}</i>" if a.get("note") else ""
        lines.append(f"• {user_line(uid)}\n   📨 {fmt_datetime(a.get('requested_at', 0))} · #{a.get('requests', 1)}{note}")
    if len(rows) > 25:
        lines.append(f"… and {len(rows) - 25} more")
    return "\n".join(lines)

def pending_kb() -> InlineKeyboardMarkup:
    rows = []
    for uid, u in store.pending_users()[:8]:
        name = (u.get("name") or "User")[:18]
        rows.append([InlineKeyboardButton(f"👤 {name}", callback_data=f"ui:{uid}"),
                     InlineKeyboardButton("1 M", callback_data=f"ap:{uid}:1m"),
                     InlineKeyboardButton("1 Y", callback_data=f"ap:{uid}:1y"),
                     InlineKeyboardButton("♾", callback_data=f"ap:{uid}:forever"),
                     InlineKeyboardButton("❌", callback_data=f"rj:{uid}")])
    rows.append([InlineKeyboardButton("🔄 Refresh", callback_data="nav:pending"),
                 InlineKeyboardButton("🏠 Home", callback_data="nav:home")])
    return InlineKeyboardMarkup(rows)

@app.on_message(filters.command("pending") & PRIVATE & admin_only)
async def cmd_pending(_, m: Message):
    await m.reply(text_pending(), reply_markup=pending_kb(), disable_web_page_preview=True)

def text_users(page: int = 0, per_page: int = 30) -> str:
    users = sorted(store.users.items(), key=lambda kv: (
        {ACCESS_PENDING: 0, ACCESS_APPROVED: 1, ACCESS_EXPIRED: 2, ACCESS_NONE: 3,
         ACCESS_REJECTED: 4, ACCESS_BANNED: 5}.get(store.status(kv[0]), 9),
        -kv[1].get("last_seen", 0)))
    total = len(users)
    chunk = users[page * per_page:(page + 1) * per_page]
    c = store.count_by_status()
    lines = [header(f"Users ({total})", "👥"),
             f"✅ {c.get(ACCESS_APPROVED, 0)} · ⏳ {c.get(ACCESS_PENDING, 0)} · ⌛ {c.get(ACCESS_EXPIRED, 0)} · "
             f"🔒 {c.get(ACCESS_NONE, 0)} · ❌ {c.get(ACCESS_REJECTED, 0)} · 🚫 {c.get(ACCESS_BANNED, 0)}\n{DIV}"]
    for uid, u in chunk:
        st = store.status(uid)
        icon = "👑" if store.is_owner(uid) else "🛡" if u.get("role") == "admin" else ACCESS_LABELS.get(st, ("", "❓"))[1]
        a = u.get("access", {})
        tail = ""
        if st == ACCESS_APPROVED and not store.is_owner(uid):
            tail = " · ♾" if not a.get("expires") else f" · ⏳ {days_left(a['expires'])} d"
        lines.append(f"{icon} {esc(u.get('name', 'User'))} — {code(uid)} · {u.get('stats', {}).get('jobs', 0)} jobs{tail}")
    if not chunk:
        lines.append("<i>No users yet.</i>")
    if total > per_page:
        lines.append(f"\n<i>Page {page + 1} / {(total + per_page - 1) // per_page}</i>")
    lines.append("\n<i>/userinfo &lt;id&gt; opens a profile with action buttons.</i>")
    return "\n".join(lines)

def users_kb(page: int = 0, per_page: int = 30) -> InlineKeyboardMarkup:
    total = len(store.users)
    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton("◀️ Prev", callback_data=f"nav:users:{page - 1}"))
    if (page + 1) * per_page < total:
        nav.append(InlineKeyboardButton("Next ▶️", callback_data=f"nav:users:{page + 1}"))
    rows = [nav] if nav else []
    rows.append([InlineKeyboardButton("⏳ Requests", callback_data="nav:pending"),
                 InlineKeyboardButton("🏠 Home", callback_data="nav:home")])
    return InlineKeyboardMarkup(rows)

@app.on_message(filters.command("users") & PRIVATE & admin_only)
async def cmd_users(_, m: Message):
    await m.reply(text_users(), reply_markup=users_kb())

@app.on_message(filters.command("userinfo") & PRIVATE & admin_only)
async def cmd_userinfo(_, m: Message):
    uid = _parse_target(m, "userinfo")
    if not uid:
        return await m.reply("Usage: <code>/userinfo 123456789</code>")
    await m.reply(text_user_card(uid), reply_markup=user_card_kb(uid), disable_web_page_preview=True)

@app.on_message(filters.command(["approve", "adduser", "extend"]) & PRIVATE & admin_only)
async def cmd_approve(_, m: Message):
    cmd = m.command[0].lower()
    extend = cmd == "extend"
    uid = _parse_target(m, cmd)
    if not uid:
        return await m.reply(f"Usage: <code>/{cmd} 123456789 1m</code>\n"
                             "Durations: <code>1w</code> · <code>1m</code> · <code>3m</code> · <code>1y</code> · "
                             "<code>forever</code> · <code>45d</code> · <code>12h</code> · <code>2026-12-31</code>")
    spec = m.command[2] if len(m.command) > 2 else ("" if extend else DEFAULT_APPROVAL)
    if not spec:
        return await m.reply("Usage: <code>/extend 123456789 1m</code>")
    if parse_duration(spec) is None:
        return await m.reply(f"⚠️ I don't understand the duration {code(spec)}.\n"
                             "Try <code>1w</code>, <code>1m</code>, <code>1y</code>, <code>forever</code>, "
                             "<code>45d</code> or a date like <code>2026-12-31</code>.")
    if store.is_banned(uid):
        return await m.reply("🚫 This user is banned — /unban first.")
    a = await grant_and_notify(uid, spec, m.from_user.id, extend=extend)
    if a is None:
        return await m.reply("⚠️ Could not approve this user.")
    await m.reply((f"✅ {user_line(uid)} " + ("extended" if extend else "approved") + ".\n"
                   f"📦 Plan: {b(plan_label(a.get('plan', spec)))} · ⏳ "
                   f"{b('♾ Lifetime' if not a.get('expires') else expiry_label(a['expires']))}")
                  + ("" if store.connected else "\n<i>No database — add to AUTHORIZED_USERS env to keep after restarts.</i>"),
                  disable_web_page_preview=True)
    if not extend:
        await notify_admins(f"✅ {user_link(m.from_user.id, m.from_user.first_name)} approved {user_line(uid)} "
                            f"({plan_label(a.get('plan', spec))}).", exclude=m.from_user.id)

@app.on_message(filters.command(["reject", "revoke", "deluser", "ban", "unban"]) & PRIVATE & admin_only)
async def cmd_moderate(_, m: Message):
    cmd = m.command[0].lower()
    uid = _parse_target(m, cmd)
    if not uid:
        return await m.reply(f"Usage: <code>/{cmd} 123456789 [reason]</code>")
    if store.is_owner(uid):
        return await m.reply("👑 You cannot moderate the owner.")
    if store.role(uid) == "admin" and not store.is_owner(m.from_user.id):
        return await m.reply("🛡 Only the owner can moderate another admin.")
    reason = " ".join(m.command[2:]).strip() if len(m.command) > 2 else ""
    by = m.from_user.id
    if cmd == "reject":
        if not store.reject(uid, by, reason):
            return await m.reply("⚠️ Unknown user.")
        await m.reply(f"❌ Request of {user_line(uid)} rejected.", disable_web_page_preview=True)
        await safe_send(uid, header("Request Declined", "❌") +
                        (f"💬 {esc(reason)}\n\n" if reason else "") +
                        (f"<i>You may send a new request after {REJECT_COOLDOWN_H} h.</i>" if REJECT_COOLDOWN_H
                         else "<i>You may send a new request anytime.</i>"),
                        reply_markup=locked_keyboard(ACCESS_REJECTED))
    elif cmd in ("revoke", "deluser"):
        if not store.revoke(uid, by, reason):
            return await m.reply("⚠️ User has no access to revoke.")
        await clear_user_commands(uid)
        await m.reply(f"🔒 Access of {user_line(uid)} revoked.", disable_web_page_preview=True)
        await safe_send(uid, header("Access Revoked", "🔒") + "Your access has been removed." +
                        (f"\n💬 {esc(reason)}" if reason else "") +
                        "\n\n<i>You may request access again.</i>",
                        reply_markup=locked_keyboard(ACCESS_NONE))
    elif cmd == "ban":
        store.ban(uid, by, reason)
        await clear_user_commands(uid)
        # drop their running / queued jobs
        for j in [j for j in QUEUE if j.user_id == uid]:
            QUEUE.remove(j); j.status = "cancelled"; j.cleanup()
        if ACTIVE and ACTIVE.user_id == uid:
            ACTIVE.cancel.set()
        for jid, j in [(k, v) for k, v in PENDING.items() if v.user_id == uid]:
            PENDING.pop(jid, None); j.cleanup()
        await m.reply(f"🚫 {user_line(uid)} banned.", disable_web_page_preview=True)
        await safe_send(uid, header("Access Blocked", "🚫") + "You have been blocked from using this bot." +
                        (f"\n💬 {esc(reason)}" if reason else ""), reply_markup=ReplyKeyboardRemove())
    elif cmd == "unban":
        if not store.unban(uid, by):
            return await m.reply("⚠️ This user is not banned.")
        await m.reply(f"♻️ {user_line(uid)} unbanned — they may request access again.", disable_web_page_preview=True)
        await safe_send(uid, header("Unblocked", "♻️") + "You may request access again.",
                        reply_markup=locked_keyboard(ACCESS_NONE))

@app.on_message(filters.command(["admins", "addadmin", "deladmin"]) & PRIVATE & owner_only)
async def cmd_admins(_, m: Message):
    cmd = m.command[0].lower()
    if cmd == "admins":
        ids = [a for a in store.admins() if not store.is_owner(a)]
        lines = [header(f"Admins ({len(ids)})", "🛡"), f"👑 {user_line(store.owner_id)} — owner"]
        lines += [f"🛡 {user_line(a)}" for a in ids]
        if ADMIN_USERS:
            lines.append(f"\n<i>From env ADMIN_USERS: {', '.join(map(str, ADMIN_USERS))}</i>")
        lines.append("\n/addadmin &lt;id&gt; · /deladmin &lt;id&gt;\n<i>Admins can approve, extend, reject and ban users.</i>")
        return await m.reply("\n".join(lines), disable_web_page_preview=True)
    uid = _parse_target(m, cmd)
    if not uid:
        return await m.reply(f"Usage: <code>/{cmd} 123456789</code>")
    if cmd == "addadmin":
        if not store.set_admin(uid, True, m.from_user.id):
            return await m.reply("⚠️ Already an admin (or the owner).")
        await set_user_commands(uid, force=True)
        await m.reply(f"🛡 {user_line(uid)} is now an admin (lifetime access).", disable_web_page_preview=True)
        await safe_send(uid, header("You are an Admin", "🛡") +
                        "You can now approve access requests.\nSend /start to refresh your menu.",
                        reply_markup=reply_keyboard(uid))
    else:
        if uid in ADMIN_USERS:
            return await m.reply("⚠️ This admin is configured in ADMIN_USERS env — remove it there first.")
        if not store.set_admin(uid, False, m.from_user.id):
            return await m.reply("⚠️ Not an admin.")
        await set_user_commands(uid, force=True)
        await m.reply(f"👤 {user_line(uid)} is no longer an admin (keeps user access).", disable_web_page_preview=True)
        await safe_send(uid, "ℹ️ You are no longer an admin. Send /start to refresh your menu.",
                        reply_markup=reply_keyboard(uid))

AUDIT_ICONS = {"request": "🙋", "approve": "✅", "extend": "➕", "reject": "❌", "revoke": "🔒", "ban": "🚫",
               "unban": "♻️", "expired": "⌛", "promote": "🛡", "demote": "👤", "withdraw": "↩️", "reminder": "🔔"}

def text_audit(limit: int = 30) -> str:
    rows = store.recent_audit(limit)
    lines = [header("Audit Log", "📜")]
    if not rows:
        lines.append("<i>Nothing yet.</i>")
    for e in rows:
        icon = AUDIT_ICONS.get(e.get("action"), "•")
        actor = "system" if not e.get("actor") else esc(e.get("actor_name") or e["actor"])
        target = esc(e.get("target_name") or e.get("target", ""))
        extra = ""
        if e.get("plan"):
            extra = f" · {esc(plan_label(e['plan']))}"
        if e.get("reason"):
            extra += f" · <i>{esc(e['reason'])}</i>"
        lines.append(f"{icon} {time.strftime('%d %b %H:%M', time.localtime(e.get('ts', 0)))} · "
                     f"{esc(e.get('action', ''))} · {target} <i>by {actor}</i>{extra}")
    return "\n".join(lines)

@app.on_message(filters.command("audit") & PRIVATE & owner_only)
async def cmd_audit(_, m: Message):
    await m.reply(text_audit(), reply_markup=InlineKeyboardMarkup([[
        InlineKeyboardButton("🔄 Refresh", callback_data="nav:audit"),
        InlineKeyboardButton("🏠 Home", callback_data="nav:home")]]))

@app.on_message(filters.command("broadcast") & PRIVATE & owner_only)
async def cmd_broadcast(_, m: Message):
    src = m.reply_to_message
    text = m.text.split(None, 1)[1] if len(m.command) > 1 and m.command[0] == "broadcast" else None
    if not src and not text:
        return await m.reply("Usage: <code>/broadcast Hello everyone</code>\n"
                             "or reply to any message with <code>/broadcast</code>.")
    sent = failed = 0
    status = await m.reply("📣 Broadcasting…")
    for uid in store.approved_users():
        for _attempt in range(2):
            try:
                if src:
                    await src.copy(uid)
                else:
                    await app.send_message(uid, header("Announcement", "📣") + esc(text))
                sent += 1
                break
            except FloodWait as e:
                await asyncio.sleep(min(float(e.value) + 1, 60))
            except Exception:
                failed += 1
                break
        await asyncio.sleep(0.1)
    await status.edit_text(f"📣 Broadcast done.\n✅ Sent: {b(sent)}  ·  ❌ Failed: {b(failed)}")

@app.on_message(filters.command("links") & PRIVATE & owner_only)
async def cmd_links(client: Client, m: Message):
    status = await m.reply("🔄 Collecting invite links…")
    chats = dict(store.chats)
    if BACKUP_GROUP_ID and BACKUP_GROUP_ID not in chats:
        chats[BACKUP_GROUP_ID] = {"title": "Backup Group", "type": "supergroup"}
    lines = []
    for cid in chats:
        try:
            chat = await client.get_chat(cid)
            me = await client.get_chat_member(cid, "me")
            if me.status not in (enums.ChatMemberStatus.ADMINISTRATOR, enums.ChatMemberStatus.OWNER):
                continue
            link = chat.invite_link or await client.export_chat_invite_link(cid)
            lines.append(f"📌 {b(chat.title)}\n🔗 {esc(link)}")
        except Exception as e:
            log.debug("links(%s): %s", cid, e)
    if lines:
        await status.edit_text(header("Admin Invite Links", "🔗") + "\n\n".join(lines), disable_web_page_preview=True)
    else:
        await status.edit_text("❌ No chats found where the bot is admin.")

@app.on_chat_member_updated()
async def track_membership(_, upd: ChatMemberUpdated):
    try:
        new = upd.new_chat_member
        if not new or not new.user or not new.user.is_self:
            return
        if new.status in (enums.ChatMemberStatus.LEFT, enums.ChatMemberStatus.BANNED):
            store.untrack_chat(upd.chat.id)
        else:
            store.track_chat(upd.chat.id, upd.chat.title or "Chat", str(upd.chat.type).split(".")[-1].lower())
    except Exception:
        pass


# ═══════════════════════════════════════════════════════════════════════════
# 📥 DOCUMENT INTAKE  →  WIZARD
# ═══════════════════════════════════════════════════════════════════════════
def wizard_lang_text(job: Job) -> str:
    return header("Step 1 · Language", "🌐") + summary_block(job) + f"{DIV}\nChoose the target language:"

def wizard_lang_kb(job: Job) -> InlineKeyboardMarkup:
    kb = language_keyboard(f"jl:{job.job_id}:", job.lang, f"jx:{job.job_id}")
    rows = [[InlineKeyboardButton("⚡ Quick Start (use my defaults)", callback_data=f"jq:{job.job_id}")]]
    if MINI_APP_URL.startswith("https://"):
        rows.append([InlineKeyboardButton("📱 Configure in Mini App",
                                          web_app=WebAppInfo(url=f"{MINI_APP_URL}#job={job.job_id}"))])
    rows += kb.inline_keyboard[:-1]
    rows.append([InlineKeyboardButton("❌ Cancel", callback_data=f"jx:{job.job_id}")])
    return InlineKeyboardMarkup(rows)

def wizard_fmt_text(job: Job) -> str:
    return header("Step 2 · Format", "📄") + summary_block(job) + f"{DIV}\nChoose the output format:"

def wizard_fmt_kb(job: Job) -> InlineKeyboardMarkup:
    return format_keyboard(f"jf:{job.job_id}:", job.out_format, f"jb:{job.job_id}:lang")

def wizard_split_text(job: Job) -> str:
    return header("Step 3 · Split size", "✂️") + summary_block(job) + f"{DIV}\nMax size per delivered part:"

def wizard_split_kb(job: Job) -> InlineKeyboardMarkup:
    return split_keyboard(f"js:{job.job_id}:", job.split_kb, f"jb:{job.job_id}:fmt", f"jc:{job.job_id}")

def wizard_confirm_text(job: Job) -> str:
    if job.split_kb:
        est_parts = max(1, -(-job.file_size // (job.split_kb * 1024)))
    else:
        est_parts = 1
    ahead = len(QUEUE) + (1 if ACTIVE else 0)
    return (header("Step 4 · Confirm", "✅") + summary_block(job) + f"{DIV}\n"
            f"🧩 Estimated parts: {b('~' + str(est_parts))}\n"
            f"👥 Jobs ahead of you: {b(ahead)}\n\n"
            "Everything looks good?")

def wizard_confirm_kb(job: Job) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🚀 Start translation", callback_data=f"jgo:{job.job_id}")],
        [InlineKeyboardButton("◀️ Back", callback_data=f"jb:{job.job_id}:spl"),
         InlineKeyboardButton("❌ Cancel", callback_data=f"jx:{job.job_id}")],
    ])

@app.on_message(filters.document & PRIVATE & authorized)
async def handle_document(_, m: Message):
    touch_user(m)
    doc = m.document
    file_name = doc.file_name or "Document"
    if OUTPUT_NAME_RE.search(file_name):
        return  # user forwarded one of our own outputs back
    ext = os.path.splitext(file_name)[1].lower()
    if ext not in INPUT_EXTS:
        return await m.reply(header("Unsupported file", "❌") +
                             f"Allowed: {b(', '.join(sorted(INPUT_EXTS)))}")
    if doc.file_size and doc.file_size > MAX_INPUT_MB * 1024 * 1024:
        return await m.reply(f"❌ File too large. Limit is {b(f'{MAX_INPUT_MB} MB')}.")
    if SHUTTING_DOWN:
        return await m.reply("🔄 Server is restarting, please try again in a minute.")

    uid = m.from_user.id
    if await user_job_count_all(uid) >= MAX_JOBS_PER_USER:
        return await m.reply(f"⏳ You already have {b(MAX_JOBS_PER_USER)} job(s) running/queued.\n"
                             "Please wait for them to finish or use /cancel.")

    # discard any unfinished wizard for this chat
    for jid, j in [(k, v) for k, v in PENDING.items() if v.chat_id == m.chat.id]:
        PENDING.pop(jid, None); j.cleanup()
    USER_STATE.pop(m.chat.id, None)

    novel_name = safe_filename(os.path.splitext(file_name)[0])
    lang, fmt, split = user_prefs(uid)
    job = Job(job_id=uuid.uuid4().hex[:10], chat_id=m.chat.id, user_id=uid,
              user_name=m.from_user.first_name or "User", file_path="", ext=ext,
              novel_name=novel_name, file_size=doc.file_size or 0,
              lang=lang, out_format=fmt, split_kb=split)
    job.file_path = ""

    status = await m.reply(header("Downloading", "📥") + f"📘 {b(novel_name)}\n{progress_bar(0)} 0%")
    live = LiveMessage(status, min_interval=3.0)

    async def dl_progress(current: int, total: int):
        r = current / total if total else 0
        await live.update(header("Downloading", "📥") + f"📘 {b(novel_name)}\n{progress_bar(r)} {r * 100:.0f}%\n"
                          f"{fmt_size(current)} / {fmt_size(total)}")

    try:
        if not QUEUE_REPO or not QUEUE_REPO.files:
            raise RuntimeError("MongoDB GridFS is required for document uploads")
        stream = await m.download(in_memory=True, progress=dl_progress)
        if not stream:
            raise RuntimeError("download returned no file")
        content = stream.getvalue() if hasattr(stream, "getvalue") else bytes(stream)
        job.gridfs_id = await QUEUE_REPO.upload_bytes(file_name, content, {
            "job_id": job.job_id, "user_id": uid, "content_type": ext,
        }) or ""
        if not job.gridfs_id:
            raise RuntimeError("GridFS upload failed")
    except Exception as e:
        log.warning("download failed: %s", e)
        job.cleanup()
        return await live.update(header("Download failed", "❌") + "Please send the file again.", back_home_kb(), force=True)

    job.msg = status
    PENDING[job.job_id] = job
    await live.update(wizard_lang_text(job), wizard_lang_kb(job), force=True)


# ═══════════════════════════════════════════════════════════════════════════
# 🔘 CALLBACKS
# ═══════════════════════════════════════════════════════════════════════════
async def _edit(q: CallbackQuery, text: str, kb: Optional[InlineKeyboardMarkup] = None):
    try:
        await q.message.edit_text(text, reply_markup=kb, disable_web_page_preview=True)
    except MessageNotModified:
        pass
    except FloodWait as e:
        await asyncio.sleep(min(float(e.value) + 1, 30))
        try:
            await q.message.edit_text(text, reply_markup=kb, disable_web_page_preview=True)
        except Exception:
            pass
    except Exception as e:
        log.debug("callback edit failed: %s", e)

async def _answer(q: CallbackQuery, text: Optional[str] = None, alert: bool = False):
    try:
        await q.answer(text, show_alert=alert)
    except Exception:
        pass

def _pending(q: CallbackQuery, job_id: str) -> Optional[Job]:
    job = PENDING.get(job_id)
    if not job or job.chat_id != q.message.chat.id or job.user_id != q.from_user.id:
        return None
    return job

@app.on_callback_query(authorized_cb)
async def callbacks(_, q: CallbackQuery):
    if not q.message:
        return await _answer(q)
    data = q.data or ""
    uid = q.from_user.id
    chat_id = q.message.chat.id
    parts = data.split(":")
    kind = parts[0]
    arg1 = parts[1] if len(parts) > 1 else ""
    arg2 = parts[2] if len(parts) > 2 else ""
    store.ensure_user(uid, q.from_user.first_name or "User", q.from_user.username or "")

    # ── Admin: approval actions on request / user cards ─────────────────────────
    if kind in ("ap", "apc", "rj", "bn", "ub", "rv", "ex", "ui"):
        if not store.is_admin(uid):
            return await _answer(q, "🛡 Admins only.", alert=True)
        if not arg1.lstrip("-").isdigit():
            return await _answer(q)
        target = int(arg1)
        if store.is_owner(target) and kind != "ui":
            return await _answer(q, "👑 The owner cannot be modified.", alert=True)
        if store.role(target) == "admin" and not store.is_owner(uid) and kind in ("rj", "bn", "rv"):
            return await _answer(q, "🛡 Only the owner can moderate an admin.", alert=True)
        if kind == "ui":
            await _edit(q, text_user_card(target), user_card_kb(target))
            return await _answer(q)
        if kind in ("ap", "ex"):
            if store.is_banned(target):
                return await _answer(q, "🚫 User is banned — unban first.", alert=True)
            a = await grant_and_notify(target, arg2, uid, extend=(kind == "ex"))
            if a is None:
                return await _answer(q, "⚠️ Could not approve.", alert=True)
            verb = "extended" if kind == "ex" else "approved"
            await _edit(q, header("Approved" if kind == "ap" else "Extended", "✅") +
                        f"👤 {user_line(target)}\n📦 Plan: {b(plan_label(a.get('plan', arg2)))}\n"
                        f"⏳ {b('♾ Lifetime' if not a.get('expires') else expiry_label(a['expires']))}\n"
                        f"<i>by {esc(q.from_user.first_name or 'admin')}</i>", user_card_kb(target))
            await notify_admins(f"✅ {user_link(uid, q.from_user.first_name)} {verb} {user_line(target)} "
                                f"({plan_label(a.get('plan', arg2))}).", exclude=uid)
            return await _answer(q, f"✅ {verb.title()}")
        if kind == "apc":
            USER_STATE[chat_id] = {"state": "await_duration", "target": target}
            await _edit(q, header("Custom duration", "✏️") + f"👤 {user_line(target)}\n\n"
                        "Type the access duration, e.g.\n"
                        "<code>45d</code> · <code>2w</code> · <code>3m</code> · <code>2y</code> · <code>12h</code> · "
                        "<code>forever</code> · <code>2026-12-31</code>",
                        InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Back", callback_data=f"ui:{target}")]]))
            return await _answer(q)
        if kind == "rj":
            USER_STATE[chat_id] = {"state": "await_reason", "target": target, "action": "reject"}
            await _edit(q, header("Reject request", "❌") + f"👤 {user_line(target)}\n\n"
                        "Type a short reason for the user, or tap <b>Skip</b>.",
                        InlineKeyboardMarkup([[InlineKeyboardButton("⏭ Skip reason", callback_data=f"rjx:{target}"),
                                               InlineKeyboardButton("◀️ Back", callback_data=f"ui:{target}")]]))
            return await _answer(q)
        if kind == "bn":
            USER_STATE[chat_id] = {"state": "await_reason", "target": target, "action": "ban"}
            await _edit(q, header("Ban user", "🚫") + f"👤 {user_line(target)}\n\n"
                        "Type a reason, or tap <b>Ban now</b>.",
                        InlineKeyboardMarkup([[InlineKeyboardButton("🚫 Ban now", callback_data=f"bnx:{target}"),
                                               InlineKeyboardButton("◀️ Back", callback_data=f"ui:{target}")]]))
            return await _answer(q)
        if kind == "ub":
            ok = store.unban(target, uid)
            if ok:
                await safe_send(target, header("Unblocked", "♻️") + "You may request access again.",
                                reply_markup=locked_keyboard(ACCESS_NONE))
            await _edit(q, text_user_card(target), user_card_kb(target))
            return await _answer(q, "♻️ Unbanned" if ok else "Not banned")
        if kind == "rv":
            ok = store.revoke(target, uid)
            if ok:
                await clear_user_commands(target)
                await safe_send(target, header("Access Revoked", "🔒") + "Your access has been removed.\n\n"
                                "<i>You may request access again.</i>", reply_markup=locked_keyboard(ACCESS_NONE))
            await _edit(q, text_user_card(target), user_card_kb(target))
            return await _answer(q, "🔒 Revoked" if ok else "Nothing to revoke")

    if kind in ("rjx", "bnx"):                    # reject / ban without a reason
        if not store.is_admin(uid) or not arg1.lstrip("-").isdigit():
            return await _answer(q, "🛡 Admins only.", alert=True)
        USER_STATE.pop(chat_id, None)
        await finish_moderation(int(arg1), uid, "reject" if kind == "rjx" else "ban", "", q=q)
        return await _answer(q)

    if kind == "nav":
        USER_STATE.pop(chat_id, None)
        if arg1 in ("pending", "users", "audit") and not store.is_admin(uid):
            return await _answer(q, "🛡 Admins only.", alert=True)
        if arg1 == "pending":
            await _edit(q, text_pending(), pending_kb())
        elif arg1 == "users":
            page = int(arg2) if arg2.isdigit() else 0
            await _edit(q, text_users(page), users_kb(page))
        elif arg1 == "audit" and store.is_owner(uid):
            await _edit(q, text_audit(), InlineKeyboardMarkup([[
                InlineKeyboardButton("🔄 Refresh", callback_data="nav:audit"),
                InlineKeyboardButton("🏠 Home", callback_data="nav:home")]]))
        elif arg1 == "access":
            await _edit(q, text_access(uid), back_home_kb())
        elif arg1 == "home":
            await _edit(q, text_home(uid, q.from_user.first_name or "there"), home_keyboard(uid))
        elif arg1 == "help":
            await _edit(q, text_help(), back_home_kb())
        elif arg1 == "settings":
            await _edit(q, text_settings(uid), settings_keyboard())
        elif arg1 == "queue":
            await _edit(q, text_queue(uid), InlineKeyboardMarkup([[
                InlineKeyboardButton("🔄 Refresh", callback_data="nav:queue"),
                InlineKeyboardButton("🏠 Home", callback_data="nav:home")]]))
        elif arg1 == "mystats":
            await _edit(q, text_mystats(uid), back_home_kb())
        elif arg1 == "owner" and store.is_owner(uid):
            await _edit(q, text_owner(), InlineKeyboardMarkup([[
                InlineKeyboardButton("🔄 Refresh", callback_data="nav:owner"),
                InlineKeyboardButton("🏠 Home", callback_data="nav:home")]]))
        return await _answer(q)

    if kind == "st":
        lang, fmt, split = user_prefs(uid)
        if arg1 == "lang":
            await _edit(q, header("Default language", "🌐") + "Pick your default target language:",
                        language_keyboard("sl:", lang, "nav:settings"))
        elif arg1 == "fmt":
            await _edit(q, header("Default format", "📄") + "Pick your default output format:",
                        format_keyboard("sf:", fmt, "nav:settings"))
        elif arg1 == "spl":
            await _edit(q, header("Default split size", "✂️") + "Pick your default part size:",
                        split_keyboard("ss:", split, "nav:settings", None))
        return await _answer(q)

    if kind in ("sl", "sf", "ss"):
        if kind == "sl" and arg1 in LANGUAGES:
            store.set_pref(uid, "lang", arg1)
        elif kind == "sf" and arg1 in OUTPUT_FORMATS:
            store.set_pref(uid, "fmt", arg1)
        elif kind == "ss" and arg1.isdigit():
            store.set_pref(uid, "split", int(arg1))
        await _edit(q, text_settings(uid), settings_keyboard())
        return await _answer(q, "✅ Saved")

    if kind == "qx":
        if ACTIVE and ACTIVE.job_id == arg1 and ACTIVE.user_id == uid:
            ACTIVE.cancel.set()
            return await _answer(q, "🛑 Stopping…")
        for j in list(QUEUE):
            if j.job_id == arg1 and j.user_id == uid:
                QUEUE.remove(j); j.status = "cancelled"; j.cleanup()
                await _edit(q, header("Cancelled", "🛑") + summary_block(j), back_home_kb())
                await notify_positions()
                return await _answer(q, "Removed from queue")
        if await cancel_persisted(arg1, uid):
            doc = await QUEUE_REPO.get_job(arg1) if QUEUE_REPO else None
            if doc and doc.get("worker_id"):
                return await _answer(q, "🛑 Stopping… the worker will confirm shortly.")
            await _edit(q, header("Cancelled", "🛑") +
                        (f"📘 File: {b(doc.get('name', 'Document'))}\n" if doc else "") +
                        "Removed from the queue.", back_home_kb())
            return await _answer(q, "Removed from queue")
        return await _answer(q, "Job already finished.")

    if kind in ("jl", "jf", "js", "jc", "jx", "jq", "jgo", "jb"):
        job = _pending(q, arg1)
        if not job:
            await _edit(q, header("Session expired", "⌛") + "Please send the file again.", back_home_kb())
            return await _answer(q, "Session expired", alert=True)

        if kind == "jx":
            PENDING.pop(arg1, None); job.cleanup(); USER_STATE.pop(chat_id, None)
            await _edit(q, header("Cancelled", "❌") + f"📘 {b(job.novel_name)} was discarded.", back_home_kb())
            return await _answer(q)

        if kind == "jl":
            if arg2 in LANGUAGES:
                job.lang = arg2
            await _edit(q, wizard_fmt_text(job), wizard_fmt_kb(job))
        elif kind == "jf":
            if arg2 in OUTPUT_FORMATS:
                job.out_format = arg2
            await _edit(q, wizard_split_text(job), wizard_split_kb(job))
        elif kind == "js":
            if arg2.isdigit():
                val = int(arg2)
                if val == 0 or MIN_SPLIT_KB <= val <= MAX_SPLIT_KB:
                    job.split_kb = val
            USER_STATE.pop(chat_id, None)
            await _edit(q, wizard_confirm_text(job), wizard_confirm_kb(job))
        elif kind == "jc":
            USER_STATE[chat_id] = {"state": "await_size", "job_id": arg1}
            await _edit(q, header("Custom split size", "✏️") + summary_block(job) + f"{DIV}\n"
                        "Type the size per part, e.g.\n"
                        "<code>750</code> (KB) · <code>2 MB</code> · <code>900kb</code>\n"
                        f"<i>Allowed: {MIN_SPLIT_KB} KB – {MAX_SPLIT_KB // 1024} MB</i>",
                        InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Back", callback_data=f"jb:{arg1}:spl")]]))
        elif kind == "jb":
            USER_STATE.pop(chat_id, None)
            if arg2 == "lang":
                await _edit(q, wizard_lang_text(job), wizard_lang_kb(job))
            elif arg2 == "fmt":
                await _edit(q, wizard_fmt_text(job), wizard_fmt_kb(job))
            else:
                await _edit(q, wizard_split_text(job), wizard_split_kb(job))
        elif kind in ("jq", "jgo"):
            if user_job_count(uid) >= MAX_JOBS_PER_USER:
                return await _answer(q, f"⏳ Max {MAX_JOBS_PER_USER} jobs at once. Wait or /cancel.", alert=True)
            if SHUTTING_DOWN:
                return await _answer(q, "🔄 Server restarting, try again shortly.", alert=True)
            PENDING.pop(arg1, None)
            USER_STATE.pop(chat_id, None)
            if kind == "jq":
                job.lang, job.out_format, job.split_kb = user_prefs(uid)
            await enqueue(job, q.message)
        return await _answer(q)

    await _answer(q)


# ═══════════════════════════════════════════════════════════════════════════
# ⌨️ FREE TEXT  (custom size input + hint)
# ═══════════════════════════════════════════════════════════════════════════
def parse_size_kb(text: str) -> Optional[int]:
    t = text.strip().lower().replace(" ", "").replace(",", ".")
    mt = re.fullmatch(r"(\d+(?:\.\d+)?)(mb|m|kb|k)?", t)
    if not mt:
        return None
    val = float(mt.group(1))
    unit = mt.group(2) or "kb"
    return int(val * 1024) if unit.startswith("m") else int(val)

async def finish_moderation(target: int, by: int, action: str, reason: str,
                            q: Optional[CallbackQuery] = None, m: Optional[Message] = None) -> None:
    """Shared tail of reject / ban started from inline buttons."""
    if store.is_owner(target):
        return
    if action == "reject":
        ok = store.reject(target, by, reason)
        if ok:
            await safe_send(target, header("Request Declined", "❌") +
                            (f"💬 {esc(reason)}\n\n" if reason else "") +
                            (f"<i>You may send a new request after {REJECT_COOLDOWN_H} h.</i>" if REJECT_COOLDOWN_H
                             else "<i>You may send a new request anytime.</i>"),
                            reply_markup=locked_keyboard(ACCESS_REJECTED))
        title = "Rejected" if ok else "Nothing to reject"
    else:
        store.ban(target, by, reason)
        await clear_user_commands(target)
        for j in [j for j in QUEUE if j.user_id == target]:
            QUEUE.remove(j); j.status = "cancelled"; j.cleanup()
        if ACTIVE and ACTIVE.user_id == target:
            ACTIVE.cancel.set()
        for jid, j in [(k, v) for k, v in PENDING.items() if v.user_id == target]:
            PENDING.pop(jid, None); j.cleanup()
        await safe_send(target, header("Access Blocked", "🚫") + "You have been blocked from using this bot." +
                        (f"\n💬 {esc(reason)}" if reason else ""), reply_markup=ReplyKeyboardRemove())
        title = "Banned"
    text = header(title, "❌" if action == "reject" else "🚫") + f"👤 {user_line(target)}" + \
        (f"\n💬 <i>{esc(reason)}</i>" if reason else "")
    if q is not None:
        await _edit(q, text, user_card_kb(target))
    elif m is not None:
        await m.reply(text, reply_markup=user_card_kb(target), disable_web_page_preview=True)
    actor = (store.user(by) or {}).get("name", "admin")
    await notify_admins(f"{'❌' if action == 'reject' else '🚫'} {user_link(by, actor)} {action}ed {user_line(target)}.",
                        exclude=by)

@app.on_message(filters.text & PRIVATE & authorized & ~filters.via_bot)
async def handle_text(client: Client, m: Message):
    if not m.text or m.text.startswith("/"):
        return  # unknown command → ignore silently
    chat_id = m.chat.id
    text_lower = m.text.strip().lower()

    label = m.text.strip()

    # ── Admin free-text states (custom duration / reject-ban reason) ──────────
    state = USER_STATE.get(chat_id)
    if state and state.get("state") in ("await_duration", "await_reason") and store.is_admin(m.from_user.id):
        if label in MENU_BUTTONS or label in ALL_BUTTONS:
            USER_STATE.pop(chat_id, None)        # admin tapped a menu button → abort the state
        else:
            USER_STATE.pop(chat_id, None)
            target = int(state["target"])
            if state["state"] == "await_duration":
                if parse_duration(label) is None:
                    USER_STATE[chat_id] = state
                    return await m.reply("⚠️ I don't understand that duration. Try <code>45d</code>, <code>3m</code>, "
                                         "<code>1y</code>, <code>forever</code> or <code>2026-12-31</code>.")
                if store.is_banned(target):
                    return await m.reply("🚫 User is banned — /unban first.")
                a = await grant_and_notify(target, label, m.from_user.id)
                if a is None:
                    return await m.reply("⚠️ Could not approve.")
                await notify_admins(f"✅ {user_link(m.from_user.id, m.from_user.first_name)} approved {user_line(target)} "
                                    f"({plan_label(a.get('plan', label))}).", exclude=m.from_user.id)
                return await m.reply(header("Approved", "✅") + f"👤 {user_line(target)}\n"
                                     f"⏳ {b('♾ Lifetime' if not a.get('expires') else expiry_label(a['expires']))}",
                                     reply_markup=user_card_kb(target), disable_web_page_preview=True)
            return await finish_moderation(target, m.from_user.id, state.get("action", "reject"), label[:200], m=m)

    # ── Menu buttons → switch the bottom keyboard page ─────────────────────
    if label in MENU_BUTTONS:
        USER_STATE.pop(chat_id, None)
        return await show_page(m, MENU_BUTTONS[label])
    if label == BTN_SETTINGS:
        # opens the settings sub-page *and* shows the inline settings card
        KB_PAGE[chat_id] = "settings"
        touch_user(m)
        return await m.reply(text_settings(m.from_user.id), reply_markup=reply_keyboard(m.from_user.id, "settings"))

    # ── Leaf buttons → dispatch to the matching command ────────────────────
    btn_cmd = ALL_BUTTONS.get(label)
    if btn_cmd:
        if btn_cmd in OWNER_BUTTONS.values() and not store.is_owner(m.from_user.id):
            # user somehow pressed an owner button (stale keyboard) → refresh their keyboard
            return await show_page(m, "main", "🔒 Owner only — menu refreshed.")
        if btn_cmd in ADMIN_BUTTONS.values() and not store.is_admin(m.from_user.id):
            return await show_page(m, "main", "🔒 Admins only — menu refreshed.")
        m.command = [btn_cmd]            # so handlers see it as a real command
        handler = {
            "start": cmd_start, "settings": cmd_settings, "queue": cmd_queue,
            "mystats": cmd_mystats, "cancel": cmd_cancel, "help": cmd_help, "id": cmd_id,
            "access": cmd_access,
            "setlang": cmd_set_pref, "setformat": cmd_set_pref, "setsplit": cmd_set_pref, "app": cmd_app,
            "stats": cmd_stats, "users": cmd_users, "pending": cmd_pending,
            "broadcast": cmd_broadcast, "links": cmd_links,
        }[btn_cmd]
        return await handler(client, m)

    if text_lower == "give me link":
        if store.is_owner(m.from_user.id):
            return await cmd_links(client, m)
        return

    state = USER_STATE.get(chat_id)
    if not state or state.get("state") != "await_size":
        return await m.reply("📎 Send me a document to translate,\nor use /help for instructions.")

    size_kb = parse_size_kb(m.text)
    if size_kb is None:
        return await m.reply("⚠️ Invalid format. Try <code>500</code> or <code>2 MB</code>.")
    if not MIN_SPLIT_KB <= size_kb <= MAX_SPLIT_KB:
        return await m.reply(f"⚠️ Size must be between {MIN_SPLIT_KB} KB and {MAX_SPLIT_KB // 1024} MB.")

    job = PENDING.get(state["job_id"])
    USER_STATE.pop(chat_id, None)
    if not job or job.chat_id != chat_id:
        return await m.reply("⌛ Session expired. Please send the file again.")
    job.split_kb = size_kb
    await m.reply(wizard_confirm_text(job), reply_markup=wizard_confirm_kb(job))


# ═══════════════════════════════════════════════════════════════════════════
# 📱 MINI APP  —  Telegram Web App served from /app  +  JSON API under /api
# ═══════════════════════════════════════════════════════════════════════════
BOT_USERNAME = ""

def verify_init_data(init_data: str) -> Optional[dict]:
    """Validate Telegram WebApp initData (HMAC-SHA256) → parsed dict or None."""
    if not init_data:
        return None
    try:
        pairs = dict(parse_qsl(init_data, keep_blank_values=True))
        their_hash = pairs.pop("hash", "")
        if not their_hash:
            return None
        check = "\n".join(f"{k}={v}" for k, v in sorted(pairs.items()))
        secret = hmac.new(b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256).digest()
        calc = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(calc, their_hash):
            return None
        auth_date = int(pairs.get("auth_date", "0") or 0)
        if auth_date and time.time() - auth_date > INIT_DATA_MAX_AGE:
            return None
        if "user" in pairs:
            pairs["user"] = json.loads(pairs["user"])
        return pairs
    except Exception as e:
        log.debug("initData verification failed: %s", e)
        return None

def _api_user(request: web.Request) -> Optional[dict]:
    """Resolve the Telegram user behind an API call (header: Authorization: tma <initData>)."""
    auth = request.headers.get("Authorization", "")
    init_data = auth[4:].strip() if auth.lower().startswith("tma ") else ""
    data = verify_init_data(init_data)
    if data and isinstance(data.get("user"), dict) and data["user"].get("id"):
        u = data["user"]
        return {"id": int(u["id"]), "name": (u.get("first_name") or "User").strip(),
                "username": u.get("username") or "", "photo": u.get("photo_url") or "",
                "lang_code": u.get("language_code") or ""}
    if MINIAPP_DEV_USER and not init_data and not RENDER_EXTERNAL_URL:
        return {"id": MINIAPP_DEV_USER, "name": "Dev User", "username": "dev", "photo": "", "lang_code": "en"}
    return None

def _json(data: Any, status: int = 200) -> web.Response:
    return web.json_response(data, status=status, headers={"Cache-Control": "no-store"})

def _err(msg: str, http_status: int = 400, **extra) -> web.Response:
    """Error envelope. `extra` may carry an access `status` field — hence the `http_status` name."""
    return _json({"ok": False, "error": msg, **extra}, http_status)

def _job_public(j: Job, uid: int, position: int = 0) -> dict:
    eng = j.progress.get("engine")
    done = eng.done if eng else 0
    total = eng.total if eng else 0
    part = j.progress.get("part", 0)
    parts = j.progress.get("parts", 0)
    ratio = 0.0
    if j.status == "running" and parts:
        ratio = ((part - 1) + (done / total if total else 0)) / parts
    return {
        "job_id": j.job_id, "name": j.novel_name, "size": j.file_size, "ext": j.ext,
        "lang": j.lang, "fmt": j.out_format, "split": j.split_kb, "status": j.status,
        "mine": j.user_id == uid, "user": j.user_name if store.is_owner(uid) or j.user_id == uid else "",
        "position": position, "created": int(j.created),
        "progress": {"ratio": round(ratio, 4), "part": part, "parts": parts,
                     "chunks_done": done, "chunks_total": total,
                     "speed": round(eng.speed, 2) if eng else 0, "eta": int(eng.eta) if eng else 0,
                     "elapsed": int(time.time() - eng.started) if eng else 0,
                     "phase": j.progress.get("phase", "")},
    }

def _job_public_doc(job: Optional[dict], uid: int, position: int = 0) -> Optional[dict]:
    """Serialize a MongoDB jobs_queue document for the Mini App."""
    if not job:
        return None
    progress = job.get("progress") or {}
    return {
        "job_id": str(job.get("job_id", job.get("_id", ""))),
        "name": job.get("name", "Document"), "size": int(job.get("size", 0) or 0),
        "ext": job.get("ext", ".txt"), "lang": job.get("lang", DEFAULT_LANG),
        "fmt": job.get("fmt", DEFAULT_FORMAT), "split": int(job.get("split_kb", 0) or 0),
        "status": job.get("status", "queued"), "mine": int(job.get("user_id", 0)) == uid,
        "user": job.get("user_name", "") if store.is_owner(uid) else "", "position": position,
        "created": int(job.get("created_at", time.time())), "progress": progress,
    }

def _config_public() -> dict:
    return {
        "bot": BOT_NAME, "username": BOT_USERNAME, "public_mode": PUBLIC_MODE,
        "db": store.connected, "db_name": MONGO_DB if store.connected else "",
        "languages": [{"code": c, "name": n, "flag": f} for c, (n, f) in LANGUAGES.items()],
        "formats": [{"code": c, "label": lbl} for c, lbl in OUTPUT_FORMATS.items()],
        "split_presets": [{"kb": kb, "label": lbl} for kb, lbl in SPLIT_PRESETS],
        "split_range": [MIN_SPLIT_KB, MAX_SPLIT_KB], "max_input_mb": MAX_INPUT_MB,
        "max_jobs": MAX_JOBS_PER_USER, "input_exts": sorted(INPUT_EXTS),
        "plans": [{"code": c, "label": lbl} for c, (lbl, _spec) in PLAN_PRESETS.items()],
        "reject_cooldown_h": REJECT_COOLDOWN_H, "reminder_days": EXPIRY_REMINDER_DAYS,
        "version": VERSION,
    }

def _access_public(uid: int) -> dict:
    """JSON view of a user's access block (for /api/me and admin lists)."""
    a = store.access(uid)
    status = store.status(uid)
    label, icon = ACCESS_LABELS.get(status, ("Unknown", "❓"))
    owner = store.is_owner(uid)
    expires = 0 if owner else int(a.get("expires") or 0)
    cooldown = 0
    if status == ACCESS_REJECTED and REJECT_COOLDOWN_H:
        cooldown = max(0, int(a.get("updated", 0) + REJECT_COOLDOWN_H * 3600 - time.time()))
    return {
        "status": status, "label": label, "icon": icon,
        "lifetime": status == ACCESS_APPROVED and not expires,
        "expires": expires, "days_left": days_left(expires) if expires else None,
        "expiry_label": ("♾ Lifetime" if status == ACCESS_APPROVED and not expires
                         else expiry_label(expires) if expires else ""),
        "plan": a.get("plan", ""), "plan_label": plan_label(a.get("plan", "")) if a.get("plan") else "",
        "approved_by": a.get("approved_by", 0), "approved_at": a.get("approved_at", 0),
        "requested_at": a.get("requested_at", 0), "requests": a.get("requests", 0),
        "note": a.get("note", ""), "reason": a.get("reason", ""), "updated": a.get("updated", 0),
        "cooldown": cooldown, "can_request": status in (ACCESS_NONE, ACCESS_EXPIRED, ACCESS_REJECTED) and not cooldown,
        "history": (a.get("history") or [])[-8:][::-1],
    }

def _me_payload(tg: dict) -> dict:
    uid = tg["id"]
    u = store.user(uid) or {}
    lang, fmt, split = user_prefs(uid)
    return {
        "id": uid, "name": u.get("name") or tg["name"], "username": tg.get("username", "") or u.get("username", ""),
        "photo": tg.get("photo", ""), "role": store.role(uid),
        "owner": store.is_owner(uid), "admin": store.is_admin(uid), "joined": u.get("joined", 0),
        "authorized": store.is_authorized(uid),
        "access": _access_public(uid),
        "prefs": {"lang": lang, "fmt": fmt, "split": split},
        "stats": {**{"jobs": 0, "parts": 0, "chars": 0}, **(u.get("stats") or {})},
        "active_jobs": user_job_count(uid),
    }

def _require(request: web.Request) -> Tuple[Optional[dict], Optional[web.Response]]:
    """Authenticated *and* approved user, else an error response."""
    tg = _api_user(request)
    if not tg:
        return None, _err("Unauthorized — open this page from Telegram.", 401, code="auth")
    store.ensure_user(tg["id"], tg["name"], tg.get("username", ""))
    if not store.is_authorized(tg["id"]):
        status = store.status(tg["id"])
        return tg, _err("This bot is private — request access from the owner.", 403,
                        code="locked", status=status, access_status=status, id=tg["id"],
                        me=_me_payload(tg), config=_config_public())
    store.touch(tg["id"])
    return tg, None

def _require_admin(request: web.Request, owner: bool = False) -> Tuple[Optional[dict], Optional[web.Response]]:
    tg, err = _require(request)
    if err:
        return tg, err
    if owner and not store.is_owner(tg["id"]):
        return tg, _err("Owner only", 403)
    if not owner and not store.is_admin(tg["id"]):
        return tg, _err("Admins only", 403)
    return tg, None

async def _body(request: web.Request) -> dict:
    try:
        data = await request.json()
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}

async def api_me(request: web.Request) -> web.Response:
    tg, err = _require(request)
    if err:
        return err
    return _json({"ok": True, "me": _me_payload(tg), "config": _config_public()})

async def api_access(request: web.Request) -> web.Response:
    """Own access card — works for everyone who is signed in (also locked users)."""
    tg = _api_user(request)
    if not tg:
        return _err("Unauthorized", 401, code="auth")
    store.ensure_user(tg["id"], tg["name"], tg.get("username", ""))
    return _json({"ok": True, "me": _me_payload(tg), "config": _config_public()})

async def api_request_access(request: web.Request) -> web.Response:
    """Locked user asks for approval from the Mini App (optional note)."""
    tg = _api_user(request)
    if not tg:
        return _err("Unauthorized", 401, code="auth")
    uid = tg["id"]
    if store.is_banned(uid):
        return _err("Access blocked.", 403, code="banned")
    body = await _body(request)
    note = str(body.get("note", "") or "").strip()[:200]
    if store.is_authorized(uid):
        return _json({"ok": True, "already": True, "me": _me_payload(tg), "config": _config_public()})
    ok, why = await send_access_request(uid, tg["name"], tg.get("username", ""), note)
    if not ok and why == "pending" and note:
        a = store.access(uid)
        a["note"] = note
        store._save_user(uid)
        why = "noted"
    if not ok and why == "cooldown":
        a = store.access(uid)
        wait = a.get("updated", 0) + REJECT_COOLDOWN_H * 3600 - time.time()
        return _err(f"Your last request was declined. Try again in {fmt_time(wait)}.", 429,
                    code="cooldown", me=_me_payload(tg))
    if ok:
        asyncio.create_task(safe_send(uid, header("Request Sent", "📨") +
                                      "Your access request has been sent to the owner.\n"
                                      "You will be notified here once it is decided.",
                                      reply_markup=locked_keyboard(ACCESS_PENDING)))
    return _json({"ok": True, "sent": ok, "why": why, "me": _me_payload(tg), "config": _config_public()})

async def api_withdraw_request(request: web.Request) -> web.Response:
    tg = _api_user(request)
    if not tg:
        return _err("Unauthorized", 401, code="auth")
    uid = tg["id"]
    a = store.access(uid)
    if a["status"] == ACCESS_PENDING:
        a.update({"status": ACCESS_NONE, "updated": int(time.time())})
        store._save_user(uid)
        store.audit("withdraw", uid, uid)
    return _json({"ok": True, "me": _me_payload(tg)})

async def api_settings(request: web.Request) -> web.Response:
    tg, err = _require(request)
    if err:
        return err
    body = await _body(request)
    if not body:
        return _err("Invalid JSON")
    uid = tg["id"]
    changed = []
    if "lang" in body:
        if body["lang"] not in LANGUAGES:
            return _err("Unknown language")
        store.set_pref(uid, "lang", body["lang"]); changed.append("lang")
    if "fmt" in body:
        if body["fmt"] not in OUTPUT_FORMATS:
            return _err("Unknown format")
        store.set_pref(uid, "fmt", body["fmt"]); changed.append("fmt")
    if "split" in body:
        try:
            val = int(body["split"])
        except (TypeError, ValueError):
            return _err("Split must be a number (KB)")
        if val != 0 and not MIN_SPLIT_KB <= val <= MAX_SPLIT_KB:
            return _err(f"Split must be 0 or between {MIN_SPLIT_KB} KB and {MAX_SPLIT_KB // 1024} MB")
        store.set_pref(uid, "split", val); changed.append("split")
    return _json({"ok": True, "changed": changed, "me": _me_payload(tg)})

async def api_jobs(request: web.Request) -> web.Response:
    tg, err = _require(request)
    if err:
        return err
    uid = tg["id"]
    owner = store.is_owner(uid)
    active = _job_public(ACTIVE, uid) if ACTIVE else None
    queue = [_job_public(j, uid, i) for i, j in enumerate(QUEUE, 1)]
    persisted = await QUEUE_REPO.list_user_jobs(uid, owner=owner) if QUEUE_REPO else []
    if persisted:
        active_doc = next((j for j in persisted if j.get("status") == "running"), None)
        active = _job_public_doc(active_doc, uid) if active_doc else active
        queue = [_job_public_doc(j, uid, index) for index, j in enumerate(
            [j for j in persisted if j.get("status") == "queued"], 1)]
    pending = [_job_public(j, uid) for j in PENDING.values() if j.user_id == uid]
    history = store.user_history(uid, 25)
    return _json({"ok": True, "active": active, "queue": queue, "pending": pending, "history": history,
                  "queue_len": len(queue), "global": store.stats if owner else None,
                  "server_time": int(time.time()), "shutting_down": SHUTTING_DOWN})

async def api_job_start(request: web.Request) -> web.Response:
    """Finish the wizard from the Mini App: pick options for a pending upload and enqueue it."""
    tg, err = _require(request)
    if err:
        return err
    body = await _body(request)
    if not body:
        return _err("Invalid JSON")
    uid = tg["id"]
    job = PENDING.get(str(body.get("job_id", "")))
    if not job or job.user_id != uid:
        return _err("Upload not found or expired — send the file again.", 404)
    if SHUTTING_DOWN:
        return _err("Server is restarting, try again shortly.", 503)
    if user_job_count(uid) >= MAX_JOBS_PER_USER:
        return _err(f"Max {MAX_JOBS_PER_USER} jobs at once. Wait or cancel one.", 429)
    lang = body.get("lang", job.lang); fmt = body.get("fmt", job.out_format)
    try:
        split = int(body.get("split", job.split_kb))
    except (TypeError, ValueError):
        return _err("Split must be a number (KB)")
    if lang not in LANGUAGES or fmt not in OUTPUT_FORMATS:
        return _err("Unknown language or format")
    if split != 0 and not MIN_SPLIT_KB <= split <= MAX_SPLIT_KB:
        return _err("Split size out of range")
    job.lang, job.out_format, job.split_kb = lang, fmt, split
    PENDING.pop(job.job_id, None)
    USER_STATE.pop(job.chat_id, None)
    if job.msg is not None:
        await enqueue(job, job.msg)
    else:
        msg = await safe_send(job.chat_id, header("Queued", "⏳") + summary_block(job))
        await enqueue(job, msg)
    return _json({"ok": True, "job": _job_public(job, uid, queue_position(job))})

async def api_job_cancel(request: web.Request) -> web.Response:
    tg, err = _require(request)
    if err:
        return err
    body = await _body(request)
    uid = tg["id"]
    target = str(body.get("job_id", "") or "")
    owner = store.is_owner(uid)
    n = 0
    if QUEUE_REPO and target:
        n += int(await cancel_persisted(target, uid))
    elif QUEUE_REPO:
        n += await cancel_all_persisted(uid)
    if ACTIVE and (not target or ACTIVE.job_id == target) and (ACTIVE.user_id == uid or owner):
        ACTIVE.cancel.set(); n += 1
    for j in [j for j in QUEUE if (not target or j.job_id == target) and (j.user_id == uid or owner)]:
        QUEUE.remove(j); j.status = "cancelled"; j.cleanup(); n += 1
        await LiveMessage(j.msg).update(header("Cancelled", "🛑") + summary_block(j), back_home_kb(), force=True)
    for jid, j in [(k, v) for k, v in PENDING.items() if (not target or k == target) and v.user_id == uid]:
        PENDING.pop(jid, None); j.cleanup(); USER_STATE.pop(j.chat_id, None); n += 1
        await LiveMessage(j.msg).update(header("Cancelled", "❌") + f"📘 {b(j.novel_name)} was discarded.",
                                        back_home_kb(), force=True)
    if n:
        asyncio.create_task(notify_positions())
    return _json({"ok": True, "cancelled": n})

# ── admin API (owner + admins) ─────────────────────────────────────────────
def _user_public(uid: int, u: dict) -> dict:
    return {"id": uid, "name": u.get("name", "User"), "username": u.get("username", ""),
            "role": store.role(uid), "joined": u.get("joined", 0), "last_seen": u.get("last_seen", 0),
            "stats": {**{"jobs": 0, "parts": 0, "chars": 0}, **(u.get("stats") or {})},
            "lang": u.get("lang", DEFAULT_LANG), "access": _access_public(uid),
            "env": uid in AUTHORIZED_USERS or uid in ADMIN_USERS or uid == OWNER_ID}

_STATUS_ORDER = {ACCESS_PENDING: 0, ACCESS_APPROVED: 1, ACCESS_EXPIRED: 2, ACCESS_NONE: 3,
                 ACCESS_REJECTED: 4, ACCESS_BANNED: 5}

async def api_admin_overview(request: web.Request) -> web.Response:
    tg, err = _require_admin(request)
    if err:
        return err
    owner = store.is_owner(tg["id"])
    users = [_user_public(uid, u) for uid, u in store.users.items()]
    users.sort(key=lambda x: (_STATUS_ORDER.get(x["access"]["status"], 9), -x["last_seen"]))
    workers = await QUEUE_REPO.workers() if QUEUE_REPO else []
    queue_len = len(QUEUE)
    active_doc = None
    if QUEUE_REPO:
        try:
            rows = await QUEUE_REPO.list_user_jobs(tg["id"], owner=True, limit=100)
            queue_len += sum(1 for r in rows if r.get("status") == "queued")
            active_doc = next((r for r in rows if r.get("status") == "running"), None)
        except Exception as e:
            log.debug("overview queue: %s", e)
    return _json({"ok": True, "stats": store.stats, "uptime": int(time.time() - store.booted),
                  "version": VERSION, "embedded_worker": EMBEDDED.status() if EMBEDDED else None,
                  "users": users, "counts": store.count_by_status(),
                  "pending": [_user_public(uid, u) for uid, u in store.pending_users()],
                  "admins": [a for a in store.admins() if not store.is_owner(a)], "owner_id": store.owner_id,
                  "chats": len(store.chats), "recent": store.recent_history(30) if owner else [],
                  "audit": store.recent_audit(40) if owner else [],
                  "active": (_job_public(ACTIVE, tg["id"]) if ACTIVE
                             else _job_public_doc(active_doc, tg["id"]) if active_doc else None),
                  "queue_len": queue_len, "pending_len": len(PENDING),
                  "db": store.connected, "db_budget": store.budget_info() if store.connected else None,
                  "public_mode": PUBLIC_MODE, "backup_group": BACKUP_GROUP_ID,
                  "env_admins": ADMIN_USERS, "env_users": AUTHORIZED_USERS,
                  "workers": workers, "worker_count": len(workers),
                  "workers_online": sum(1 for worker in workers if worker.get("status") != "offline")})

async def api_admin_user(request: web.Request) -> web.Response:
    tg, err = _require_admin(request)
    if err:
        return err
    try:
        uid = int(request.match_info["id"])
    except (KeyError, ValueError):
        return _err("Bad id")
    u = store.user(uid)
    if not u:
        return _err("Unknown user", 404)
    return _json({"ok": True, "user": _user_public(uid, u), "history": store.user_history(uid, 15)})

async def api_admin_users(request: web.Request) -> web.Response:
    """Approval actions from the Mini App:
       {action: approve|extend|reject|revoke|ban|unban|promote|demote, id, duration?, reason?}"""
    tg, err = _require_admin(request)
    if err:
        return err
    by = tg["id"]
    body = await _body(request)
    action = str(body.get("action", "") or "").lower()
    try:
        uid = int(body.get("id"))
    except (TypeError, ValueError):
        return _err("Expected {action, id: <int>}")
    if action in ("add", "remove"):                      # v4 names still accepted
        action = "approve" if action == "add" else "revoke"
    if store.is_owner(uid):
        return _err("The owner cannot be modified", 403)
    if store.role(uid) == "admin" and not store.is_owner(by) and action in ("reject", "revoke", "ban", "demote"):
        return _err("Only the owner can moderate an admin", 403)
    reason = str(body.get("reason", "") or "").strip()[:200]
    actor_name = (store.user(by) or {}).get("name") or tg["name"]

    if action in ("approve", "extend"):
        spec = str(body.get("duration", "") or DEFAULT_APPROVAL).strip()
        if parse_duration(spec) is None:
            return _err("Unknown duration — try 1w, 1m, 1y, forever, 45d or 2026-12-31")
        if store.is_banned(uid):
            return _err("User is banned — unban first", 409)
        a = await grant_and_notify(uid, spec, by, extend=(action == "extend"))
        if a is None:
            return _err("Could not approve", 409)
        asyncio.create_task(notify_admins(f"✅ {user_link(by, actor_name)} {'extended' if action == 'extend' else 'approved'} "
                                          f"{user_line(uid)} ({plan_label(a.get('plan', spec))}).", exclude=by))
    elif action == "reject":
        if not store.reject(uid, by, reason):
            return _err("Unknown user", 404)
        asyncio.create_task(safe_send(uid, header("Request Declined", "❌") +
                                      (f"💬 {esc(reason)}\n\n" if reason else "") +
                                      (f"<i>You may send a new request after {REJECT_COOLDOWN_H} h.</i>" if REJECT_COOLDOWN_H
                                       else "<i>You may send a new request anytime.</i>"),
                                      reply_markup=locked_keyboard(ACCESS_REJECTED)))
    elif action == "revoke":
        if not store.revoke(uid, by, reason):
            return _err("User has no access to revoke", 404)
        asyncio.create_task(clear_user_commands(uid))
        asyncio.create_task(safe_send(uid, header("Access Revoked", "🔒") + "Your access has been removed." +
                                      (f"\n💬 {esc(reason)}" if reason else "") + "\n\n<i>You may request access again.</i>",
                                      reply_markup=locked_keyboard(ACCESS_NONE)))
    elif action == "ban":
        store.ban(uid, by, reason)
        asyncio.create_task(clear_user_commands(uid))
        for j in [j for j in QUEUE if j.user_id == uid]:
            QUEUE.remove(j); j.status = "cancelled"; j.cleanup()
        if ACTIVE and ACTIVE.user_id == uid:
            ACTIVE.cancel.set()
        for jid, j in [(k, v) for k, v in PENDING.items() if v.user_id == uid]:
            PENDING.pop(jid, None); j.cleanup()
        asyncio.create_task(safe_send(uid, header("Access Blocked", "🚫") + "You have been blocked from using this bot." +
                                      (f"\n💬 {esc(reason)}" if reason else ""), reply_markup=ReplyKeyboardRemove()))
    elif action == "unban":
        if not store.unban(uid, by):
            return _err("User is not banned", 409)
        asyncio.create_task(safe_send(uid, header("Unblocked", "♻️") + "You may request access again.",
                                      reply_markup=locked_keyboard(ACCESS_NONE)))
    elif action in ("promote", "demote"):
        if not store.is_owner(by):
            return _err("Owner only", 403)
        if action == "demote" and uid in ADMIN_USERS:
            return _err("This admin is configured in ADMIN_USERS env — remove it there first", 409)
        if not store.set_admin(uid, action == "promote", by):
            return _err("Already an admin" if action == "promote" else "Not an admin", 409)
        asyncio.create_task(set_user_commands(uid, force=True))
        asyncio.create_task(safe_send(uid, header("You are an Admin", "🛡") + "You can now approve access requests.\n"
                                      "Send /start to refresh your menu." if action == "promote"
                                      else "ℹ️ You are no longer an admin. Send /start to refresh your menu.",
                                      reply_markup=reply_keyboard(uid)))
    else:
        return _err("Unknown action")
    u = store.user(uid) or store.ensure_user(uid, "User")
    return _json({"ok": True, "user": _user_public(uid, u), "counts": store.count_by_status()})

async def api_admin_broadcast(request: web.Request) -> web.Response:
    tg, err = _require_admin(request, owner=True)
    if err:
        return err
    body = await _body(request)
    text = str(body.get("text", "") or "").strip()
    if not text or len(text) > 3500:
        return _err("Text required (max 3500 chars)")
    recipients = store.approved_users()

    async def _run():
        sent = failed = 0
        for uid in recipients:
            try:
                await app.send_message(uid, header("Announcement", "📣") + esc(text))
                sent += 1
            except FloodWait as e:
                await asyncio.sleep(min(float(e.value) + 1, 60))
            except Exception:
                failed += 1
            await asyncio.sleep(0.1)
        await safe_send(tg["id"], f"📣 Broadcast done.\n✅ Sent: {b(sent)}  ·  ❌ Failed: {b(failed)}")

    asyncio.create_task(_run())
    return _json({"ok": True, "recipients": len(recipients)})

# ── static Mini App files ──────────────────────────────────────────────────
_STATIC_TYPES = {".html": "text/html", ".css": "text/css", ".js": "application/javascript",
                 ".svg": "image/svg+xml", ".png": "image/png", ".ico": "image/x-icon",
                 ".json": "application/json", ".webmanifest": "application/manifest+json"}

_ASSET_V_CACHE: Dict[str, str] = {}


def _asset_version() -> str:
    """Short fingerprint of the Mini App bundle (app.js + style.css).

    ``index.html`` references its assets as ``/app/app.js?v=__ASSET_V__``; the
    placeholder is replaced with this value when the page is served. Because
    the version changes whenever the files change, Telegram's aggressive
    WebView cache can never pair a new ``index.html`` with a stale ``app.js``
    (the classic "stuck on the skeleton after a deploy" bug), while unchanged
    assets stay cacheable for a long time.
    """
    sig = []
    for name in ("app.js", "style.css"):
        try:
            st = os.stat(os.path.join(MINI_APP_DIR, name))
            sig.append(f"{name}:{st.st_mtime_ns}:{st.st_size}")
        except OSError:
            sig.append(f"{name}:missing")
    key = "|".join(sig)
    if key not in _ASSET_V_CACHE:
        _ASSET_V_CACHE.clear()          # only ever one live entry
        h = hashlib.sha1()
        for name in ("app.js", "style.css"):
            try:
                with open(os.path.join(MINI_APP_DIR, name), "rb") as fh:
                    h.update(fh.read())
            except OSError:
                h.update(b"missing")
        h.update(VERSION.encode())
        _ASSET_V_CACHE[key] = h.hexdigest()[:10]
    return _ASSET_V_CACHE[key]


async def miniapp_file(request: web.Request) -> web.Response:
    rel = request.match_info.get("path", "") or "index.html"
    rel = os.path.normpath(rel).lstrip(os.sep).replace("\\", "/")
    if rel.startswith("..") or "/.." in rel:
        raise web.HTTPNotFound()
    path = os.path.join(MINI_APP_DIR, rel)
    if not os.path.isfile(path):
        if os.path.isfile(os.path.join(MINI_APP_DIR, "index.html")):
            path = os.path.join(MINI_APP_DIR, "index.html")
        else:
            raise web.HTTPNotFound()
    ext = os.path.splitext(path)[1].lower()
    ctype = _STATIC_TYPES.get(ext, "application/octet-stream")
    if ext == ".html":
        # The shell is tiny and must always be fresh so it points at the current asset version.
        with open(path, "r", encoding="utf-8") as fh:
            body = fh.read().replace("__ASSET_V__", _asset_version())
        return web.Response(text=body, content_type="text/html", charset="utf-8",
                            headers={"Cache-Control": "no-store, max-age=0", "X-Asset-Version": _asset_version()})
    # Versioned assets (…?v=<hash>) are immutable; bare requests get a short TTL + revalidation.
    if request.query.get("v"):
        cache = "public, max-age=31536000, immutable"
    else:
        cache = "public, max-age=300, must-revalidate"
    text_like = ctype.startswith("text/") or "javascript" in ctype or "json" in ctype
    return web.FileResponse(path, headers={"Cache-Control": cache,
                                           "Content-Type": f"{ctype}; charset=utf-8" if text_like else ctype})


# ═══════════════════════════════════════════════════════════════════════════
# 🌍 HEALTH SERVER  (Render web services must bind $PORT)
# ═══════════════════════════════════════════════════════════════════════════
async def health(_request: web.Request) -> web.Response:
    workers = []
    gridfs = {"files": 0, "mb": 0.0}
    if QUEUE_REPO:
        try:
            workers = await QUEUE_REPO.workers()
            gridfs = await QUEUE_REPO.storage_stats()
        except Exception:
            workers = []
    return web.json_response({
        "status": "ok" if not SHUTTING_DOWN else "shutting_down",
        "role": "master+worker" if EMBEDDED else "master",
        "polling": True,
        "embedded_worker": EMBEDDED.status() if EMBEDDED else None,
        "workers": len(workers),
        "workers_busy": sum(1 for w in workers if w.get("status") == "running"),
        "bot": BOT_USERNAME,
        "version": VERSION,
        "uptime_sec": int(time.time() - store.booted),
        "db": "mongodb" if store.connected else "memory",
        "db_size_mb": store.db_size_mb if store.connected else 0,
        "db_budget_mb": DB_BUDGET_MB if store.connected else 0,
        "gridfs": gridfs,
        "mini_app": bool(MINI_APP_URL),
        "users": len(store.users),
        "access": store.count_by_status(),
        "active": ACTIVE.novel_name if ACTIVE else None,
        "queue": len(QUEUE),
        "pending": len(PENDING),
        "stats": store.stats,
    })

async def index(_request: web.Request) -> web.Response:
    body = (f"<!doctype html><meta charset='utf-8'><title>{html.escape(BOT_NAME)}</title>"
            f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<link rel='preconnect' href='https://fonts.googleapis.com'>"
            f"<link rel='stylesheet' href='https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;600;800&display=swap'>"
            f"<body style=\"font-family:'Plus Jakarta Sans',system-ui,sans-serif;-webkit-font-smoothing:antialiased;"
            f"padding:2rem;background:#0f172a;color:#e2e8f0;line-height:1.5\">"
            f"<h1>📚 {html.escape(BOT_NAME)}</h1>"
            f"<p>Telegram bot is <b style='color:#4ade80'>online</b>"
            + (f" as <a style='color:#93c5fd' href='https://t.me/{html.escape(BOT_USERNAME)}'>@{html.escape(BOT_USERNAME)}</a>" if BOT_USERNAME else "")
            + f".</p><p>Storage: <b>{'MongoDB' if store.connected else 'in-memory'}</b></p>"
            f"<p><a style='color:#93c5fd' href='/health'>/health</a> · "
            f"<a style='color:#93c5fd' href='/app'>/app</a> (Mini App — open from Telegram)</p></body>")
    return web.Response(text=body, content_type="text/html")

async def start_health_server() -> web.AppRunner:
    web_app = web.Application(client_max_size=256 * 1024)
    web_app.add_routes([
        web.get("/", index),                      # aiohttp registers HEAD automatically
        web.get("/health", health),
        # Mini App
        web.get("/app", miniapp_file), web.get("/app/", miniapp_file), web.get("/app/{path:.*}", miniapp_file),
        web.get("/api/me", api_me),
        web.get("/api/access", api_access),
        web.post("/api/access/request", api_request_access),
        web.post("/api/access/withdraw", api_withdraw_request),
        web.post("/api/settings", api_settings),
        web.get("/api/jobs", api_jobs),
        web.post("/api/jobs/start", api_job_start),
        web.post("/api/jobs/cancel", api_job_cancel),
        web.get("/api/admin/overview", api_admin_overview),
        web.get("/api/admin/user/{id}", api_admin_user),
        web.post("/api/admin/users", api_admin_users),
        web.post("/api/admin/broadcast", api_admin_broadcast),
    ])
    runner = web.AppRunner(web_app, access_log=None)
    await runner.setup()
    site = web.TCPSite(runner, "0.0.0.0", PORT)
    await site.start()
    log.info("HTTP server listening on 0.0.0.0:%d (mini app dir: %s)", PORT, MINI_APP_DIR)
    return runner

async def keep_alive():
    """Render free tier sleeps after ~15 min without inbound traffic.
    Pinging our own public URL keeps the instance awake."""
    if not KEEP_ALIVE or not RENDER_EXTERNAL_URL:
        log.info("Keep-alive disabled (KEEP_ALIVE=%s, RENDER_EXTERNAL_URL=%r)", KEEP_ALIVE, RENDER_EXTERNAL_URL)
        return
    url = f"{RENDER_EXTERNAL_URL}/health"
    log.info("Keep-alive pinging %s every %ds", url, KEEP_ALIVE_INTERVAL)
    while not SHUTTING_DOWN:
        await asyncio.sleep(KEEP_ALIVE_INTERVAL)
        try:
            sess = await get_http()
            async with sess.get(url, timeout=aiohttp.ClientTimeout(total=20)) as r:
                log.debug("keep-alive → %s", r.status)
        except Exception as e:
            log.debug("keep-alive failed: %s", e)


# ═══════════════════════════════════════════════════════════════════════════
# ▶️ ENTRYPOINT
# ═══════════════════════════════════════════════════════════════════════════
async def shutdown_jobs():
    """Tell every affected user that a restart interrupted them."""
    global SHUTTING_DOWN
    SHUTTING_DOWN = True
    if ACTIVE:
        ACTIVE.cancel.set()
    for j in list(QUEUE):
        QUEUE.remove(j)
        j.cleanup()
        await LiveMessage(j.msg).update(
            header("Restarting", "🔄") + summary_block(j) + f"{DIV}\n"
            "The server is being redeployed. Please resend the file in a minute.",
            back_home_kb(), force=True)
    for jid, j in list(PENDING.items()):
        PENDING.pop(jid, None)
        j.cleanup()
    if QUEUE_WAKE:
        QUEUE_WAKE.set()

async def main():
    global QUEUE_WAKE, BOT_USERNAME, QUEUE_REPO, EMBEDDED
    QUEUE_WAKE = asyncio.Event()
    runner = await start_health_server()          # bind the port FIRST → Render sees us healthy
    await store.connect()                         # MongoDB → users, GridFS and shared jobs_queue
    QUEUE_REPO = job_repo_from_store(store)
    if QUEUE_REPO:
        try:
            await QUEUE_REPO.ensure_indexes()
        except Exception as e:
            log.debug("queue indexes: %s", e)

    for attempt in range(1, 6):
        try:
            await app.start()
            break
        except FloodWait as e:
            log.warning("FloodWait on start: sleeping %ss", e.value)
            await asyncio.sleep(float(e.value) + 1)
        except Exception as e:
            fatal = any(k in str(e) for k in ("API_ID_INVALID", "ACCESS_TOKEN_INVALID", "ACCESS_TOKEN_EXPIRED",
                                               "API_ID_PUBLISHED_FLOOD", "USER_DEACTIVATED"))
            log.error("Telegram start failed (attempt %d/5): %s", attempt, e)
            if fatal:
                log.critical("Credentials rejected by Telegram — check API_ID / API_HASH / BOT_TOKEN.")
            if attempt == 5 or fatal:
                await runner.cleanup()
                sys.exit(1)
            await asyncio.sleep(5 * attempt)

    me = await app.get_me()
    BOT_USERNAME = me.username or ""
    # Command menu: default scope = locked/user list; owner + pre-authorised users get their own list
    await set_global_commands()
    for uid in list(store.users.keys())[:200]:
        await set_user_commands(uid, force=True)
    # Chat "menu" button (next to the attach icon) opens the Mini App
    if MINI_APP_URL.startswith("https://"):
        try:
            await app.set_chat_menu_button(menu_button=MenuButtonWebApp("📱 App", WebAppInfo(url=MINI_APP_URL)))
        except Exception as e:
            log.debug("set_chat_menu_button: %s", e)
    tasks = [asyncio.create_task(janitor(), name="janitor"),
             asyncio.create_task(keep_alive(), name="keep_alive")]
    # Embedded worker: the master translates too, so one free Render service is enough.
    if EMBEDDED_WORKER and QUEUE_REPO:
        from worker import TranslationWorker
        EMBEDDED = TranslationWorker(client=app, db=QUEUE_REPO, embedded=True)
        tasks.append(asyncio.create_task(EMBEDDED.serve(), name="embedded_worker"))
        log.info("Embedded worker %s started (set EMBEDDED_WORKER=0 for a polling-only master)", EMBEDDED.node_id)
    elif EMBEDDED_WORKER:
        log.warning("EMBEDDED_WORKER=1 but MongoDB is unavailable → no translation worker is running!")
    else:
        log.info("Embedded worker disabled — deploy at least one worker.py service")
    workers_now = await QUEUE_REPO.workers() if QUEUE_REPO else []
    log.info("%s v%s online as @%s | owner=%s | users=%d | public=%s | pdf=%s | db=%s | app=%s | port=%d | workers=%d",
             BOT_NAME, VERSION, me.username, store.owner_id or "none", len(store.users), PUBLIC_MODE, HAS_PDF,
             "mongodb" if store.connected else "memory", MINI_APP_URL or "-", PORT, len(workers_now))
    if store.owner_id:
        await safe_send(store.owner_id, header("Bot Online", "🟢") +
                        f"@{me.username} is running · v{VERSION}\n"
                        f"👥 Users: {b(len(store.users))}\n"
                        f"📄 PDF support: {b('yes' if HAS_PDF else 'no')}\n"
                        f"🗄 Storage: {b('MongoDB · ' + MONGO_DB if store.connected else 'in-memory (no database)')}\n"
                        f"⚙️ Worker: {b('embedded · ' + EMBEDDED.node_id if EMBEDDED else 'external only')}"
                        + (f" · {len(workers_now)} online" if workers_now else "") + "\n"
                        f"📱 Mini App: {b('enabled' if MINI_APP_URL.startswith('https://') else 'not configured')}")
    else:
        log.warning("No OWNER_ID configured — nobody can approve access requests! Set OWNER_ID in the environment.")

    await idle()                                  # blocks until SIGINT/SIGTERM

    log.info("Shutting down…")
    await shutdown_jobs()
    if EMBEDDED:
        EMBEDDED.stop()                           # loops exit; a running job is handed back to the queue
    for t in tasks:
        t.cancel()
    await asyncio.gather(*tasks, return_exceptions=True)
    await store.flush()
    if HTTP and not HTTP.closed:
        await HTTP.close()
    try:
        await app.stop()
    except Exception:
        pass
    await runner.cleanup()
    shutil.rmtree(BASE_DIR, ignore_errors=True)
    log.info("Bot stopped cleanly.")

if __name__ == "__main__":
    try:
        app.run(main())
    except KeyboardInterrupt:
        pass
